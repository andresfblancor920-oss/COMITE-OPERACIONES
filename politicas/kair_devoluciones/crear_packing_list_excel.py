"""Genera la plantilla editable de la Lista Detallada de Empaque (Packing List) KAIR.

Uso: python politicas/kair_devoluciones/crear_packing_list_excel.py
Crea Plantilla_Packing_List_KAIR.xlsx junto a este script (lo sobrescribe).
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

NAVY, SAGE, SAGE_LIGHT, INPUT = "0E2F56", "8FAF8C", "DDE8DA", "FFF6D5"
thin = Side(style="thin", color="C5D0DC")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CATEGORIAS = ["Kingroof", "Kingfrigo", "Accesorios", "Fijaciones", "Puertas", "Ventanas", "Otros"]
FILAS_EJEMPLO = [
    ("Kingroof", "Panel de cubierta Kingroof"),
    ("Kingroof", "Remates / caballete de cubierta"),
    ("Kingfrigo", "Panel de muro Kingfrigo"),
    ("Kingfrigo", "Panel divisorio interior"),
    ("Accesorios", "Perfiles, ángulos y remates"),
    ("Accesorios", "Sellantes y cintas"),
    ("Fijaciones", "Tornillería autoperforante"),
    ("Fijaciones", "Anclajes / remaches"),
    ("Puertas", "Puerta con marco y herrajes"),
    ("Ventanas", "Ventana con marco y vidrio"),
]
N_FILAS = 30


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Packing List"
    ws.sheet_view.showGridLines = False

    anchos = [6, 14, 14, 36, 7, 8, 20, 10, 11, 9, 11, 26]
    for i, w in enumerate(anchos, start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    ws.merge_cells("A1:L1")
    ws["A1"] = "KAIR by Kingspan · LISTA DETALLADA DE EMPAQUE (PACKING LIST)"
    ws["A1"].font = Font(bold=True, size=15, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    ws.merge_cells("A2:L2")
    ws["A2"] = "Anexo B · Política de Devolución de Producto KAIR - Versión 1 - 29-09-2026 · Celdas amarillas: diligenciar"
    ws["A2"].font = Font(italic=True, size=9, color="50627A")

    # Encabezado de la remisión (3 bloques de etiqueta + valor por fila)
    campos = [
        ("N° de remisión", "Fecha de despacho", "Pedido / Orden de compra"),
        ("Cliente", "Proyecto / Obra", "Dirección de entrega"),
        ("Módulo", "Asesor comercial", "N° de bultos / paquetes"),
    ]
    bloques = [("A", "B", "D"), ("E", "G", "H"), ("I", "K", "L")]  # etiqueta, valor desde, valor hasta
    for r, fila in enumerate(campos, start=4):
        for (lab_col, v1, v2), etiqueta in zip(bloques, fila):
            ws[f"{lab_col}{r}"] = etiqueta
            ws[f"{lab_col}{r}"].font = Font(bold=True, size=9, color=NAVY)
            if lab_col == "E":
                ws.merge_cells(f"E{r}:F{r}")
            if lab_col == "I":
                ws.merge_cells(f"I{r}:J{r}")
            ws.merge_cells(f"{v1}{r}:{v2}{r}")
            ws[f"{v1}{r}"].fill = PatternFill("solid", fgColor=INPUT)
            ws[f"{v1}{r}"].border = BORDER
        ws.row_dimensions[r].height = 20
    ws["B6"] = "Módulo Familiar KAIR"

    # Tabla de productos
    h = 8
    cabeceras = ["Ítem", "Categoría", "Código / Ref.", "Descripción", "Und.", "Cant.",
                 "Dimensiones (L×A×E)", "Vol. (m³)", "Paquete N°", "Foto N°", "Recibido (Sí/No)", "Observaciones"]
    for c, t in enumerate(cabeceras, start=1):
        cell = ws.cell(row=h, column=c, value=t)
        cell.font = Font(bold=True, color="FFFFFF", size=9)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[h].height = 30

    dv_cat = DataValidation(type="list", formula1='"' + ",".join(CATEGORIAS) + '"', allow_blank=True)
    dv_rec = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True)
    ws.add_data_validation(dv_cat)
    ws.add_data_validation(dv_rec)

    primera = h + 1
    ultima = h + N_FILAS
    for i in range(N_FILAS):
        r = primera + i
        ws.cell(row=r, column=1, value=i + 1).alignment = Alignment(horizontal="center")
        if i < len(FILAS_EJEMPLO):
            ws.cell(row=r, column=2, value=FILAS_EJEMPLO[i][0])
            ws.cell(row=r, column=4, value=FILAS_EJEMPLO[i][1])
        for c in range(1, 13):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.font = Font(size=9)
            if c > 1:
                cell.fill = PatternFill("solid", fgColor=INPUT)
        ws.cell(row=r, column=8).number_format = "0.000"
        dv_cat.add(f"B{r}")
        dv_rec.add(f"K{r}")

    t = ultima + 1
    ws.merge_cells(f"A{t}:E{t}")
    ws[f"A{t}"] = "TOTALES"
    ws[f"A{t}"].alignment = Alignment(horizontal="right")
    ws[f"F{t}"] = f"=SUM(F{primera}:F{ultima})"
    ws[f"H{t}"] = f"=SUM(H{primera}:H{ultima})"
    ws[f"H{t}"].number_format = "0.000"
    ws[f"K{t}"] = f'=COUNTIF(K{primera}:K{ultima},"Sí")&" de "&COUNTA(D{primera}:D{ultima})&" ítems recibidos"'
    ws.merge_cells(f"K{t}:L{t}")
    for c in range(1, 13):
        cell = ws.cell(row=t, column=c)
        cell.fill = PatternFill("solid", fgColor=SAGE_LIGHT)
        cell.font = Font(bold=True, size=9, color=NAVY)
        cell.border = BORDER

    # Transporte (también se comparte por correo)
    r0 = t + 2
    ws.merge_cells(f"A{r0}:L{r0}")
    ws[f"A{r0}"] = "INFORMACIÓN DE TRANSPORTE (también enviada por correo electrónico)"
    ws[f"A{r0}"].font = Font(bold=True, color="FFFFFF")
    ws[f"A{r0}"].fill = PatternFill("solid", fgColor=SAGE)
    transporte = [("Transportadora", "Tipo de vehículo", "Placa"),
                  ("Conductor", "Documento de identidad", "Teléfono")]
    for k, fila in enumerate(transporte, start=r0 + 1):
        for (lab_col, v1, v2), etiqueta in zip(bloques, fila):
            ws[f"{lab_col}{k}"] = etiqueta
            ws[f"{lab_col}{k}"].font = Font(bold=True, size=9, color=NAVY)
            if lab_col == "E":
                ws.merge_cells(f"E{k}:F{k}")
            if lab_col == "I":
                ws.merge_cells(f"I{k}:J{k}")
            ws.merge_cells(f"{v1}{k}:{v2}{k}")
            ws[f"{v1}{k}"].fill = PatternFill("solid", fgColor=INPUT)
            ws[f"{v1}{k}"].border = BORDER
        ws.row_dimensions[k].height = 20

    # Registro fotográfico y observaciones
    r1 = r0 + 4
    ws[f"A{r1}"] = "Registro fotográfico adjunto"
    ws[f"A{r1}"].font = Font(bold=True, size=9, color=NAVY)
    ws.merge_cells(f"A{r1}:C{r1}")
    ws[f"D{r1}"].fill = PatternFill("solid", fgColor=INPUT)
    ws[f"D{r1}"].border = BORDER
    dv_rec.add(f"D{r1}")
    ws[f"E{r1}"] = "N° de fotos"
    ws[f"E{r1}"].font = Font(bold=True, size=9, color=NAVY)
    ws.merge_cells(f"E{r1}:F{r1}")
    ws[f"G{r1}"].fill = PatternFill("solid", fgColor=INPUT)
    ws[f"G{r1}"].border = BORDER
    ws[f"A{r1 + 1}"] = "Novedades en la entrega"
    ws[f"A{r1 + 1}"].font = Font(bold=True, size=9, color=NAVY)
    ws.merge_cells(f"A{r1 + 1}:C{r1 + 1}")
    ws.merge_cells(f"D{r1 + 1}:L{r1 + 2}")
    ws[f"D{r1 + 1}"].fill = PatternFill("solid", fgColor=INPUT)
    ws[f"D{r1 + 1}"].alignment = Alignment(wrap_text=True, vertical="top")

    # Firmas
    r2 = r1 + 5
    for col_ini, col_fin, texto in [("A", "D", "Despacha · KAIR by Kingspan"),
                                     ("E", "H", "Transportador"),
                                     ("I", "L", "Recibe conforme · Cliente")]:
        ws.merge_cells(f"{col_ini}{r2}:{col_fin}{r2}")
        ws[f"{col_ini}{r2}"] = texto
        ws[f"{col_ini}{r2}"].alignment = Alignment(horizontal="center")
        ws[f"{col_ini}{r2}"].font = Font(bold=True, size=9, color=NAVY)
        for c in range(ord(col_ini) - 64, ord(col_fin) - 64 + 1):
            ws.cell(row=r2, column=c).border = Border(top=Side(style="thin", color=NAVY))

    r3 = r2 + 2
    ws.merge_cells(f"A{r3}:L{r3 + 1}")
    ws[f"A{r3}"] = ("Verifique el material contra esta lista al momento de la entrega (numerales 1.3 y 5.4). "
                    "Reporte faltantes o productos incorrectos a devolucioneskair@kingspan.com dentro de los "
                    "5 días hábiles siguientes a la recepción (numerales 1.1 y 4.1).")
    ws[f"A{r3}"].alignment = Alignment(wrap_text=True, vertical="center")
    ws[f"A{r3}"].font = Font(size=9, color=NAVY)
    ws[f"A{r3}"].fill = PatternFill("solid", fgColor=SAGE_LIGHT)

    ws.freeze_panes = f"A{primera}"
    ws.print_title_rows = f"{h}:{h}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    salida = Path(__file__).with_name("Plantilla_Packing_List_KAIR.xlsx")
    wb.save(salida)
    print(f"Plantilla generada: {salida.name}")


if __name__ == "__main__":
    main()
