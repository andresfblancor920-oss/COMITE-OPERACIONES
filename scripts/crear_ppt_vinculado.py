"""
Crea la versión VINCULADA del comité de Ventas / Revenue / Order Intake:

    vinculado/Base_Comite_Ventas.xlsx       ← el Excel base que se diligencia cada semana
    vinculado/Comite_Ventas_Vinculado.pptx  ← 3 diapositivas vinculadas al Excel base
    vinculado/Actualizar_Rutas.bat / .ps1   ← apunta los vínculos al Excel de la misma carpeta
    vinculado/LEEME.txt                     ← instrucciones de uso

En el PowerPoint, encabezados + KPIs, tablas y alertas son rangos de Excel
vinculados (Pegado especial > Pegar vínculo) y los gráficos son gráficos
nativos vinculados al mismo libro. Todo el cálculo (subtotales, % de
cumplimiento, semáforos, textos de alerta) vive en fórmulas del Excel base,
así que para actualizar la presentación basta con diligenciar el Excel,
guardarlo y en PowerPoint "Actualizar vínculos". No hay macros.

    python scripts/crear_ppt_vinculado.py [carpeta_salida]

Requiere LibreOffice (solo para generar las vistas previas iniciales de los
rangos vinculados; PowerPoint las reemplaza en la primera actualización).
"""

import copy
import functools
import io
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import pymupdf
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import Rule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.text import PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import generar_ppt_ventas as G  # noqa: E402  (estilos de gráficos compartidos)

ROOT = Path(__file__).resolve().parent.parent
XLSX_NAME = "Base_Comite_Ventas.xlsx"
PPTX_NAME = "Comite_Ventas_Vinculado.pptx"
RUTA_DEFECTO = r"C:\Comite_Ventas"

# Paleta (hex para Excel)
NAVY, TEXT, MUTED, CARD, LINE, SUBT = "0B2545", "1F2937", "6B7280", "F3F5F8", "DDE2E8", "E3EBF5"
GREEN, AMBER, RED = "2E7D32", "C78A00", "C62828"
GREEN_BG, AMBER_BG, RED_BG = "DCF0DD", "FFF0C2", "FBDADA"
INPUT_BG, CALC_BG, LINK_BG = "FFF2CC", "EDEFF2", "DDEBF7"

THIN = Side(style="thin", color=LINE)
F = "Calibri"


def fill(hex_):
    return PatternFill("solid", fgColor=hex_, bgColor=hex_)


def font(size=10, bold=False, color=TEXT, italic=False):
    return Font(name=F, size=size, bold=bold, color=color, italic=italic)


# Fórmulas de texto independientes del idioma/región de Excel:
# FIXED usa el separador de miles del sistema; nada de TEXT con máscaras.
def N0(x):
    return f"FIXED({x},0)"


def PC(x):
    return f'ROUND(({x})*100,0)&"%"'


def MM(x):
    return f'"$ "&FIXED(({x})/1000000,0)&" M"'


def SEV(p):
    """1 rojo, 2 amarillo, 3 verde, 4 neutro (sin dato)."""
    return f"IF(NOT(ISNUMBER({p})),4,IF({p}>=$BA$1,3,IF({p}>=$BA$2,2,1)))"


# ---------------------------------------------------------------------------
# Bloques (rangos que se vinculan al PowerPoint)
# ---------------------------------------------------------------------------
class Bloque:
    def __init__(self, ws, r0, c0, widths_px, heights_pt, nombre):
        self.ws, self.r0, self.c0, self.nombre = ws, r0, c0, nombre
        self.w_px, self.h_pt = widths_px, heights_pt
        for i, px in enumerate(widths_px):
            ws.column_dimensions[get_column_letter(c0 + i)].width = px / 7
        for j, pt in enumerate(heights_pt):
            ws.row_dimensions[r0 + j].height = pt
        # todo el rango con fondo blanco explícito (vista previa y Excel iguales)
        for j in range(len(heights_pt)):
            for i in range(len(widths_px)):
                c = self.c(j, i)
                c.fill = fill("FFFFFF")
                c.font = font()

    def c(self, r, col):
        return self.ws.cell(row=self.r0 + r, column=self.c0 + col)

    def a(self, r, col):
        return f"${get_column_letter(self.c0 + col)}${self.r0 + r}"

    def rel(self, r, col):
        return f"{get_column_letter(self.c0 + col)}{self.r0 + r}"

    @property
    def area(self):
        c1 = get_column_letter(self.c0)
        c2 = get_column_letter(self.c0 + len(self.w_px) - 1)
        return f"{c1}{self.r0}:{c2}{self.r0 + len(self.h_pt) - 1}"

    @property
    def size_in(self):
        return sum(self.w_px) / 96, sum(self.h_pt) / 72

    @property
    def fin(self):
        return self.r0 + len(self.h_pt), self.c0 + len(self.w_px)

    def merge(self, r, c1, c2):
        self.ws.merge_cells(start_row=self.r0 + r, start_column=self.c0 + c1,
                            end_row=self.r0 + r, end_column=self.c0 + c2)


class Hoja:
    """Hoja PPT_n: los bloques se ubican en escalera para no compartir
    filas ni columnas (cada rango conserva sus propios altos y anchos)."""

    def __init__(self, wb, nombre):
        self.ws = wb.create_sheet(nombre)
        self.ws.sheet_view.showGridLines = False
        self.ws.sheet_properties.tabColor = NAVY
        self.r, self.c = 1, 1
        self.bloques = []
        # umbrales del semáforo (para formatos condicionales en la misma hoja)
        self.ws["BA1"] = "=Config!$B$8"
        self.ws["BA2"] = "=Config!$B$9"
        self.ws["AZ1"] = "Umbral verde"
        self.ws["AZ2"] = "Umbral amarillo"
        self.aux_col = 55  # BC en adelante: columnas auxiliares

    def bloque(self, widths, heights, nombre):
        b = Bloque(self.ws, self.r, self.c, widths, heights, nombre)
        self.r, self.c = b.fin[0] + 1, b.fin[1] + 1
        self.bloques.append(b)
        return b

    def aux(self, n=1):
        c = self.aux_col
        self.aux_col += n
        return c


def cf_expr(ws, rng, formula, color=None, bg=None, bold=None):
    dxf = DifferentialStyle(font=Font(color=color, bold=bold) if color or bold else None,
                            fill=fill(bg) if bg else None)
    ws.conditional_formatting.add(rng, Rule(type="expression", dxf=dxf, formula=[formula], stopIfTrue=True))


def cf_semaforo_celdas(ws, rng, tl, con_fondo=True):
    """Semáforo sobre el valor numérico de cada celda del rango (tl = celda sup. izq. relativa)."""
    cf_expr(ws, rng, f"AND(ISNUMBER({tl}),{tl}>=$BA$1)", GREEN, GREEN_BG if con_fondo else None, True)
    cf_expr(ws, rng, f"AND(ISNUMBER({tl}),{tl}>=$BA$2)", AMBER, AMBER_BG if con_fondo else None, True)
    cf_expr(ws, rng, f"ISNUMBER({tl})", RED, RED_BG if con_fondo else None, True)


def cf_por_severidad(ws, rng, sev_ref):
    cf_expr(ws, rng, f"{sev_ref}=1", RED)
    cf_expr(ws, rng, f"{sev_ref}=2", AMBER)
    cf_expr(ws, rng, f"{sev_ref}=3", GREEN)
    cf_expr(ws, rng, f"{sev_ref}=4", NAVY)


# ---------------------------------------------------------------------------
# Hojas de entrada
# ---------------------------------------------------------------------------
def celda(ws, ref, v, bg=None, fmt=None, bold=False, size=10, color=TEXT, align=None):
    c = ws[ref]
    c.value = v
    c.font = font(size, bold, color)
    c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    if bg:
        c.fill = fill(bg)
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = Alignment(horizontal=align, vertical="center")
    return c


def titulo(ws, t, nota):
    ws["A1"], ws["A2"] = t, nota
    ws["A1"].font = font(14, True, NAVY)
    ws["A2"].font = font(9, color=MUTED, italic=True)
    ws.sheet_view.showGridLines = False


def encabezados(ws, row, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.fill, c.font = fill(NAVY), font(10, True, "FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 32


def hoja_instrucciones(wb):
    ws = wb.active
    ws.title = "Instrucciones"
    titulo(ws, "Base del Comité de Operaciones · Ventas, Revenue y Order Intake",
           "Customer Excellence — este libro alimenta por vínculos la presentación Comite_Ventas_Vinculado.pptx")
    lineas = [
        "CADA SEMANA",
        "1. Diligencie las celdas AMARILLAS de Config, Ventas_m2, Revenue, Order_Intake, Pipeline y Comentarios.",
        "   Celdas AZULES = traen el dato de otra hoja por fórmula (puede sobrescribirlas si necesita otro valor).",
        "   Celdas GRISES = fórmulas; no las edite.",
        "2. Guarde el libro (Ctrl+G / Ctrl+S).",
        "3. Abra Comite_Ventas_Vinculado.pptx y elija 'Actualizar vínculos'.",
        "   Si PowerPoint ya estaba abierto: Archivo > Información > Editar vínculos a archivos > seleccione todos > Actualizar ahora.",
        "",
        "PRIMERA VEZ (o si cambia la carpeta)",
        f"- Deje este Excel y el PowerPoint en la MISMA carpeta. Si la carpeta es {RUTA_DEFECTO}, no hay que hacer nada más.",
        "- Si usa otra carpeta: con PowerPoint cerrado, doble clic en Actualizar_Rutas.bat (apunta todos los vínculos a esta carpeta).",
        "- No cambie el nombre de este archivo ni de las hojas PPT_1, PPT_2, PPT_3 (los vínculos usan nombres definidos).",
        "",
        "QUÉ HAY EN LAS HOJAS PPT_1, PPT_2, PPT_3",
        "- Son las diapositivas 'armadas' con fórmulas: encabezado + KPIs, tablas, alertas y los datos de los gráficos.",
        "- Cada rango vinculado tiene un nombre (Fórmulas > Administrador de nombres): S1_Encabezado, S1_Tabla, S1_Alertas, etc.",
        "- Si ajusta formato (colores, fuentes) en esos rangos, la presentación lo tomará al actualizar vínculos.",
        "",
        "SEMÁFORO: verde ≥ umbral verde, amarillo ≥ umbral amarillo, rojo por debajo (Config).",
        "ALERTAS: se escriben solas con fórmulas; los comentarios de la hoja Comentarios aparecen primero.",
    ]
    for i, t in enumerate(lineas, start=4):
        c = ws.cell(row=i, column=1, value=t)
        c.font = font(10, bold=t.isupper() and bool(t))
    ws.column_dimensions["A"].width = 130
    for i, (bg, t) in enumerate([(INPUT_BG, "Dato de entrada"), (LINK_BG, "Traído por fórmula (sobrescribible)"),
                                 (CALC_BG, "Fórmula")], start=26):
        celda(ws, f"A{i}", t, bg)


def hoja_config(wb):
    ws = wb.create_sheet("Config")
    titulo(ws, "Configuración del periodo", "Alimenta títulos, semáforos y alertas de las 3 diapositivas")
    filas = [("Mes", "Septiembre", None), ("Año", 2026, "0"), ("Semana de corte (1-4)", 4, "0"),
             ("Fecha del comité (texto o fecha)", "30/09/2026", None), ("Área que presenta", "Customer Excellence", None),
             ("Umbral verde (≥)", 0.95, "0%"), ("Umbral amarillo (≥)", 0.85, "0%"),
             ("Alcance Order Intake", "Premium (Env + Refri)", None), ("Budget Order Intake mes (m²)", 147719, "#,##0")]
    for r, (k, v, fmt) in enumerate(filas, start=3):
        celda(ws, f"A{r}", k, bold=True)
        celda(ws, f"B{r}", v, INPUT_BG, fmt)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 26


# Filas fijas de Ventas_m2 (mismo orden de la diapositiva)
V_ROWS = [  # (fila, línea, tipo, datos)
    (5, "Envolventes", "L", (103379, 83329, 91536, (25832, 16666, 19999, 20832), (28862, 19733, 17184, 17550), 80178)),
    (6, "Refrigeración", "L", (35600, 41293, 43141, (12801, 8259, 9910, 10323), (12610, 5429, 8229, 15025), 33426)),
    (7, "Premium", "S", (5, 6)),
    (8, "KingCore", "L", (29779, 25931, 27974, (3690, 5186, 10572, 6483), (4460, 1760, 9860, 9129), 25209)),
    (9, "Single Skin", "L", (23851, 4594, 20387, (721, 2800, 0, 1073), (2209, 1042, 0, 1343), 3251)),
    (10, "Ducto", "L", (520, 1044, 1385, (396, 0, 520, 0), (1044, 0, 0, 0), 1044)),
    (11, "Estándar", "S", (8, 10)),
    (12, "Línea Arquitectónica", "L", (718, 718, 3182, (0, 0, 0, 718), (0, 264, 0, 0), 264)),
    (13, "Total", "T", None),
]
V_LINEAS = [r for r, _, t, _ in V_ROWS if t == "L"]


def hoja_ventas(wb):
    ws = wb.create_sheet("Ventas_m2")
    titulo(ws, "Volumen de ventas (m²) por línea",
           "F'cast W = plan de facturación semanal · Real/Proy W = facturado en semanas cerradas y proyección en la semana en curso · "
           "Real Sales = facturado acumulado. Para agregar una línea, inserte una fila dentro del grupo y ajuste el subtotal.")
    heads = ["Línea", "F'cast Inicial", "F'cast Actual", "Budget", "F'cast W1", "F'cast W2", "F'cast W3", "F'cast W4",
             "Real/Proy W1", "Real/Proy W2", "Real/Proy W3", "Real/Proy W4", "Real Sales (facturado)",
             "Cierre proyectado", "% F'cast vs Budget", "% Real vs Budget", "Gap Real vs Budget",
             "Cumpl. W1", "Cumpl. W2", "Cumpl. W3", "Cumpl. W4"]
    encabezados(ws, 4, heads, [22] + [11] * 11 + [13, 12, 11, 11, 12, 9, 9, 9, 9])
    for r, nombre, tipo, dat in V_ROWS:
        sub = tipo != "L"
        bg = SUBT if sub else None
        celda(ws, f"A{r}", nombre, bg or INPUT_BG, bold=sub)
        for ci in range(2, 14):
            L = get_column_letter(ci)
            if tipo == "L":
                fc_ini, fc, b, fw, rw, real = dat
                vals = [fc_ini, fc, b, *fw, *rw, real]
                celda(ws, f"{L}{r}", vals[ci - 2], INPUT_BG, "#,##0")
            elif tipo == "S":
                celda(ws, f"{L}{r}", f"=SUM({L}{dat[0]}:{L}{dat[1]})", SUBT, "#,##0", True)
            else:
                celda(ws, f"{L}{r}", f"={L}7+{L}11+{L}12", SUBT, "#,##0", True)
        celda(ws, f"N{r}", f"=SUM(I{r}:L{r})", bg or CALC_BG, "#,##0", sub)
        celda(ws, f"O{r}", f'=IF(D{r}=0,"",C{r}/D{r})', bg or CALC_BG, "0%", sub)
        celda(ws, f"P{r}", f'=IF(D{r}=0,"",M{r}/D{r})', bg or CALC_BG, "0%", sub)
        celda(ws, f"Q{r}", f"=M{r}-D{r}", bg or CALC_BG, "#,##0", sub)
        for k in range(4):
            fc, re, out = "EFGH"[k], "IJKL"[k], "RSTU"[k]
            celda(ws, f"{out}{r}", f'=IF(OR({fc}{r}=0,{k + 1}>Config!$B$5),"",{re}{r}/{fc}{r})', bg or CALC_BG, "0%", sub)
    for rng in ("O5:P13", "R5:U13"):
        tl = rng.split(":")[0]
        cf_expr(ws, rng, f"AND(ISNUMBER({tl}),{tl}>=Config!$B$8)", GREEN, GREEN_BG, True)
        cf_expr(ws, rng, f"AND(ISNUMBER({tl}),{tl}>=Config!$B$9)", AMBER, AMBER_BG, True)
        cf_expr(ws, rng, f"ISNUMBER({tl})", RED, RED_BG, True)
    ws.freeze_panes = "B5"


# Revenue: filas fijas
R_ROWS = [  # fila, línea, unidad, cantidad, COP, budget cant, budget COP, tipo
    (5, "Envolventes", "m²", "=Ventas_m2!C5", 9169606489, "=Ventas_m2!D5", None, "L"),
    (6, "Refrigeración", "m²", "=Ventas_m2!C6", 4478225850, "=Ventas_m2!D6", None, "L"),
    (7, "Premium", "m²", None, None, None, None, "S5:6"),
    (8, "KingCore", "m²", "=Ventas_m2!C8", 836015895, "=Ventas_m2!D8", None, "L"),
    (9, "Single Skin", "m²", "=Ventas_m2!C9", 103998996, "=Ventas_m2!D9", None, "L"),
    (10, "Ducto", "m²", "=Ventas_m2!C10", 52372800, "=Ventas_m2!D10", None, "L"),
    (11, "Traslúcida", "m²", 2155, 573520000, None, None, "L"),
    (12, "Estándar", "m²", None, None, None, None, "S8:11"),
    (13, "Línea Arquitectónica", "m²", 264, 20500000, "=Ventas_m2!D12", None, "L"),
    (14, "Proyectos", "—", None, 328000000, None, 533000000, "L"),
    (15, "Fijaciones y Accesorios", "—", None, 1007000000, None, None, "L"),
    (16, "Consumibles", "—", None, 0, None, None, "L"),
    (17, "Puertas", "und", 17, 83050000, None, None, "L"),
    (18, "Facturación ICO", "—", None, 1251189846, None, None, "L"),
    (19, "TOTAL F'cast (COP)", "", None, None, None, None, "T"),
]
R_LINEAS = [r[0] for r in R_ROWS if r[7] == "L"]


def hoja_revenue(wb):
    ws = wb.create_sheet("Revenue")
    titulo(ws, "Revenue (COP) por línea — F'cast del mes",
           "Cantidades y Budget de las líneas de panel vienen de Ventas_m2 (azul). % Perf. = cantidad ÷ Budget cantidad "
           "o, si no hay, Revenue ÷ Budget COP. En subtotales, solo sobre las líneas que tienen Budget.")
    encabezados(ws, 4, ["Línea", "Unidad", "Cantidad F'cast", "Revenue F'cast (COP)", "Budget cantidad",
                        "Budget Revenue (COP)", "Precio promedio", "% Performance", "% Mix Revenue"],
                [26, 8, 14, 20, 14, 20, 15, 13, 12])
    for r, nombre, uni, cant, cop, bc, bcop, tipo in R_ROWS:
        sub = tipo != "L"
        bg = SUBT if sub else None
        celda(ws, f"A{r}", nombre, bg or INPUT_BG, bold=sub)
        celda(ws, f"B{r}", uni, bg or INPUT_BG, bold=sub, align="center")
        if tipo == "L":
            for L, v, fmt in (("C", cant, "#,##0"), ("D", cop, '"$" #,##0'), ("E", bc, "#,##0"), ("F", bcop, '"$" #,##0')):
                celda(ws, f"{L}{r}", v, LINK_BG if isinstance(v, str) else INPUT_BG, fmt)
            celda(ws, f"H{r}", f'=IF(N(E{r})>0,C{r}/E{r},IF(N(F{r})>0,D{r}/F{r},""))', CALC_BG, "0%")
        elif tipo.startswith("S"):
            a, b = tipo[1:].split(":")
            for L, fmt in (("C", "#,##0"), ("D", '"$" #,##0'), ("E", "#,##0"), ("F", '"$" #,##0')):
                celda(ws, f"{L}{r}", f"=SUM({L}{a}:{L}{b})", SUBT, fmt, True)
            celda(ws, f"H{r}", f'=IF(E{r}>0,SUMIF(E{a}:E{b},">0",C{a}:C{b})/E{r},IF(F{r}>0,D{r}/F{r},""))', SUBT, "0%", True)
        else:
            celda(ws, f"C{r}", "", SUBT)
            celda(ws, f"D{r}", "=D7+D12+" + "+".join(f"D{x}" for x in range(13, 19)), SUBT, '"$" #,##0', True)
            celda(ws, f"E{r}", "", SUBT)
            celda(ws, f"F{r}", "", SUBT)
            celda(ws, f"H{r}", "", SUBT)
        if tipo != "T":
            celda(ws, f"G{r}", f'=IF(N(C{r})=0,"",D{r}/C{r})', bg or CALC_BG, '"$" #,##0', sub)
        else:
            celda(ws, f"G{r}", "", SUBT)
        celda(ws, f"I{r}", f"=IF($D$19=0,\"\",D{r}/$D$19)", bg or CALC_BG, "0%", sub)
    cf_expr(ws, "H5:H18", "AND(ISNUMBER(H5),H5>=Config!$B$8)", GREEN, GREEN_BG, True)
    cf_expr(ws, "H5:H18", "AND(ISNUMBER(H5),H5>=Config!$B$9)", AMBER, AMBER_BG, True)
    cf_expr(ws, "H5:H18", "ISNUMBER(H5)", RED, RED_BG, True)
    ws.freeze_panes = "B5"


OI_ROWS = [(5, "W1", 25901, "Cerrada"), (6, "W2", 41491, "Cerrada"), (7, "W3", 45187, "Cerrada"), (8, "W4", 27588, "En curso")]


def hoja_oi(wb):
    ws = wb.create_sheet("Order_Intake")
    titulo(ws, "Order Intake semanal — Budget vs negocios ganados (m² adjudicados)",
           "Budget semanal = Budget del mes (Config) ÷ 4, sobrescribible. Estado: Cerrada / En curso / Proyectada. "
           "Real vacío = semana aún sin datos.")
    encabezados(ws, 4, ["Semana", "Budget (m²)", "Real ganado (m²)", "Estado", "Gap (m²)", "Cumplimiento",
                        "Por ganar ponderado (Pipeline)", "Budget acumulado", "Ganado acumulado", "Proyectado acumulado"],
                [10, 13, 15, 12, 12, 13, 16, 15, 15, 16])
    dv = DataValidation(type="list", formula1='"Cerrada,En curso,Proyectada"', allow_blank=True)
    ws.add_data_validation(dv)
    for r, w, real, est in OI_ROWS:
        celda(ws, f"A{r}", w, CALC_BG, bold=True, align="center")
        bud = "=Config!$B$11-SUM(B5:B7)" if r == 8 else "=ROUND(Config!$B$11/4,0)"
        celda(ws, f"B{r}", bud, LINK_BG, "#,##0")
        celda(ws, f"C{r}", real, INPUT_BG, "#,##0")
        celda(ws, f"D{r}", est, INPUT_BG, align="center")
        dv.add(f"D{r}")
        celda(ws, f"E{r}", f'=IF(C{r}="","",C{r}-B{r})', CALC_BG, "#,##0")
        celda(ws, f"F{r}", f'=IF(OR(C{r}="",N(B{r})=0),"",C{r}/B{r})', CALC_BG, "0%")
        celda(ws, f"G{r}", f'=SUMIFS(Pipeline!$H$5:$H$29,Pipeline!$F$5:$F$29,A{r})', CALC_BG, "#,##0")
        celda(ws, f"H{r}", f"=SUM($B$5:B{r})", CALC_BG, "#,##0")
        celda(ws, f"I{r}", f'=IF(C{r}="",NA(),SUM($C$5:C{r}))', CALC_BG, "#,##0")
        celda(ws, f"J{r}", f'=IF(C{r}="",NA(),I{r})+SUM($G$5:G{r})', CALC_BG, "#,##0")
    celda(ws, "A9", "TOTAL", SUBT, bold=True, align="center")
    for L, f_, fmt in (("B", "=SUM(B5:B8)", "#,##0"), ("C", "=SUM(C5:C8)", "#,##0"), ("D", "", None),
                       ("E", "=C9-B9", "#,##0"), ("F", '=IF(B9=0,"",C9/B9)', "0%"), ("G", "=SUM(G5:G8)", "#,##0")):
        celda(ws, f"{L}9", f_, SUBT, fmt, True)
    cf_expr(ws, "F5:F9", "AND(ISNUMBER(F5),F5>=Config!$B$8)", GREEN, GREEN_BG, True)
    cf_expr(ws, "F5:F9", "AND(ISNUMBER(F5),F5>=Config!$B$9)", AMBER, AMBER_BG, True)
    cf_expr(ws, "F5:F9", "ISNUMBER(F5)", RED, RED_BG, True)
    ws["A11"] = "Control: la suma semanal debe igualar el Budget del mes"
    ws["A11"].font = font(9, color=MUTED, italic=True)
    ws["D11"] = '=IF(B9=Config!B11,"OK","Revise: la suma semanal difiere del Budget del mes")'


PIPE = [("Oportunidad A (ejemplo)", "Cliente por definir", "Envolventes", 6500, 0.7, "W4", "Comercial"),
        ("Oportunidad B (ejemplo)", "Cliente por definir", "Refrigeración", 4200, 0.5, "W4", "Comercial"),
        ("Oportunidad C (ejemplo)", "Cliente por definir", "Envolventes", 12000, 0.6, "Mes siguiente", "Comercial"),
        ("Oportunidad D (ejemplo)", "Cliente por definir", "Refrigeración", 8000, 0.4, "Mes siguiente", "Comercial")]


def hoja_pipeline(wb):
    ws = wb.create_sheet("Pipeline")
    titulo(ws, "Pipeline — negocios por ganar",
           "Filas de EJEMPLO: reemplácelas por las oportunidades reales. Ponderado = m² × probabilidad. "
           "Cierre esperado: W1..W4 del mes en curso o 'Mes siguiente'.")
    encabezados(ws, 4, ["Oportunidad / Proyecto", "Cliente", "Línea", "m²", "Probabilidad", "Cierre esperado",
                        "Responsable", "m² ponderados", "Orden (aux)"], [28, 24, 16, 12, 12, 16, 16, 15, 12])
    dv = DataValidation(type="list", formula1='"W1,W2,W3,W4,Mes siguiente"', allow_blank=True)
    ws.add_data_validation(dv)
    for i in range(25):
        r = 5 + i
        d = PIPE[i] if i < len(PIPE) else (None,) * 7
        for j, v in enumerate(d):
            L = "ABCDEFG"[j]
            celda(ws, f"{L}{r}", v, INPUT_BG, "#,##0" if L == "D" else ("0%" if L == "E" else None))
        dv.add(f"F{r}")
        celda(ws, f"H{r}", f'=IF(A{r}="","",N(D{r})*N(E{r}))', CALC_BG, "#,##0")
        celda(ws, f"I{r}", f'=IF(A{r}="","",H{r}+(100-ROW())/1000000)', CALC_BG, "0.000000")
    celda(ws, "A31", "TOTAL", SUBT, bold=True)
    celda(ws, "D31", "=SUM(D5:D29)", SUBT, "#,##0", True)
    celda(ws, "H31", "=SUM(H5:H29)", SUBT, "#,##0", True)
    ws.column_dimensions["I"].hidden = True


def hoja_comentarios(wb):
    ws = wb.create_sheet("Comentarios")
    titulo(ws, "Comentarios para cada diapositiva (opcional)",
           "Aparecen primero en el recuadro de alertas de la diapositiva; deje vacío si no aplica.")
    encabezados(ws, 4, ["Diapositiva", "Comentario"], [18, 110])
    for i, s in enumerate(["1 · Volumen m²"] * 2 + ["2 · Revenue"] * 2 + ["3 · Order Intake"] * 2):
        celda(ws, f"A{5 + i}", s, CALC_BG, bold=True)
        celda(ws, f"B{5 + i}", None, INPUT_BG)


# ---------------------------------------------------------------------------
# Hojas PPT (diapositivas armadas con fórmulas)
# ---------------------------------------------------------------------------
CARD_W = [283, 19, 283, 19, 283, 19, 283]
CARD_H = [40, 22, 12, 20, 34, 22, 8]


def encabezado(h, nombre, titulo_f, subtitulo_f, tarjetas):
    """tarjetas: (etiqueta, valor_f, detalle_f, pct_f | None)."""
    b = h.bloque(CARD_W, CARD_H, nombre)
    ws = h.ws
    b.merge(0, 0, 4)
    b.merge(1, 0, 4)
    c = b.c(0, 0)
    c.value, c.font = f"={titulo_f}", font(26, True, NAVY)
    c.alignment = Alignment(vertical="center")
    c = b.c(1, 0)
    c.value, c.font = f"={subtitulo_f}", font(12, color=MUTED)
    c = b.c(0, 6)
    c.value, c.font = "Comité de Operaciones", font(11, True, NAVY)
    c.alignment = Alignment(horizontal="right", vertical="bottom")
    c = b.c(1, 6)
    c.value = ('=Config!$B$7&" · "&IF(ISNUMBER(Config!$B$6),TEXT(DAY(Config!$B$6),"00")&"/"&'
               'TEXT(MONTH(Config!$B$6),"00")&"/"&YEAR(Config!$B$6),Config!$B$6)')
    c.font, c.alignment = font(10, color=MUTED), Alignment(horizontal="right", vertical="top")
    for k, (lab, val, det, p) in enumerate(tarjetas):
        col = 2 * k
        for r in range(3, 7):
            b.c(r, col).fill = fill(CARD)
        c = b.c(3, col)
        c.value, c.font = lab.upper(), font(9, True, MUTED)
        c.alignment = Alignment(indent=1, vertical="bottom")
        c = b.c(4, col)
        c.value, c.font = f"={val}", font(22, True, NAVY)
        c.alignment = Alignment(indent=1, vertical="center")
        c = b.c(5, col)
        if p:
            sev = h.aux()
            sref = f"${get_column_letter(sev)}${b.r0 + 5}"
            ws.cell(row=b.r0 + 5, column=sev, value=f"={SEV(p)}")
            c.value = f'="● "&{PC(p)}&"  "&{det}'
            c.font = font(10, True, MUTED)
            cf_por_severidad(ws, b.rel(5, col), sref)
        else:
            c.value, c.font = f"={det}", font(10, color=MUTED)
        c.alignment = Alignment(indent=1, vertical="center")
    return b


def tabla(h, nombre, widths, header, header_h, row_h, filas, nota=None, nota_h=14, size=10):
    """filas: lista de (celdas, estilo) — celdas: fórmulas/valores; estilo: 'L' línea, 'S' subtotal, 'T' total.
    Cada celda puede ser (valor, fmt) o dict(v, fmt, pct=True, align)."""
    hs = [header_h] + [row_h] * len(filas) + ([nota_h] if nota is not None else [])
    b = h.bloque(widths, hs, nombre)
    ws = b.ws
    for i, t in enumerate(header):
        c = b.c(0, i)
        c.value, c.fill, c.font = t, fill(NAVY), font(size - 1, True, "FFFFFF")
        c.alignment = Alignment(horizontal="left" if i == 0 else "center", vertical="center", wrap_text=True,
                                indent=1 if i == 0 else 0)
    for j, (cells, est) in enumerate(filas, start=1):
        bg = {"S": SUBT, "T": NAVY}.get(est, "FFFFFF")
        color = "FFFFFF" if est == "T" else TEXT
        for i, spec in enumerate(cells):
            spec = spec if isinstance(spec, dict) else {"v": spec}
            c = b.c(j, i)
            v = spec.get("v")
            c.value = v
            c.fill = fill(bg)
            c.font = font(size, est in ("S", "T") or spec.get("bold", False), color)
            c.number_format = spec.get("fmt", "General")
            c.alignment = Alignment(horizontal=spec.get("align", "left" if i == 0 else "right"),
                                    vertical="center", indent=1 if i == 0 else 0)
            c.border = Border(bottom=Side(style="thin", color=LINE))
            if spec.get("pct"):
                c.alignment = Alignment(horizontal="center", vertical="center")
                tl = b.rel(j, i)
                cf_semaforo_celdas(ws, tl, tl)
    if nota is not None:
        c = b.c(len(filas) + 1, 0)
        c.value, c.font = nota, font(8, color=MUTED, italic=True)
    return b


def alertas(h, nombre, widths, candidatos, n, row_h, titulo_t="ALERTAS Y TENDENCIAS", size=10, wrap=False):
    """candidatos: lista de (texto_f, severidad_f). Muestra los n primeros no vacíos."""
    ws = h.ws
    hs = ([16] if titulo_t else [6]) + [row_h] * n + [4]
    b = h.bloque(widths, hs, nombre)
    for r in range(len(hs)):
        for i in range(len(widths)):
            b.c(r, i).fill = fill(CARD)
    if titulo_t:
        c = b.c(0, 0)
        c.value, c.font = titulo_t, font(9, True, MUTED)
        c.alignment = Alignment(indent=1, vertical="center")
        b.merge(0, 0, len(widths) - 1)
    ct, cs, cc = h.aux(), h.aux(), h.aux()
    r0 = b.r0 + 20  # auxiliares debajo, fuera de la vista
    for k, (tf, sf) in enumerate(candidatos):
        r = r0 + k
        ws.cell(row=r, column=ct, value=f"=IFERROR({tf},\"\")")
        ws.cell(row=r, column=cs, value=f"=IFERROR({sf},4)")
        prev = f"{get_column_letter(cc)}{r - 1}" if k else "0"
        ws.cell(row=r, column=cc, value=f'={prev}+IF({get_column_letter(ct)}{r}="",0,1)')
    L = lambda col: f"${get_column_letter(col)}${r0}:${get_column_letter(col)}${r0 + len(candidatos) - 1}"
    for k in range(n):
        dot, txt = b.c(1 + k, 0), b.c(1 + k, 1)
        m = f"MATCH({k + 1},{L(cc)},0)"
        txt.value = f'=IFERROR(INDEX({L(ct)},{m}),"")'
        txt.font = font(size)
        txt.alignment = Alignment(vertical="center" if not wrap else "top", wrap_text=wrap)
        sev_cell = ws.cell(row=b.r0 + 1 + k, column=h.aux())
        sev_cell.value = f"=IFERROR(INDEX({L(cs)},{m}),0)"
        sref = f"${get_column_letter(sev_cell.column)}${sev_cell.row}"
        dot.value = f'=IF({txt.coordinate}="","","●")'
        dot.font = font(size + 1, True, MUTED)
        dot.alignment = Alignment(horizontal="right", vertical="center" if not wrap else "top")
        cf_por_severidad(ws, dot.coordinate, sref)
    return b


def datos_grafico(h, titulo_t, categorias, series):
    """Bloque de datos de un gráfico (fuera de las vistas). Devuelve refs para el XML del gráfico."""
    ws = h.ws
    c0 = 80 + len(getattr(h, "graficos", [])) * 6  # columnas CB en adelante
    h.graficos = getattr(h, "graficos", []) + [titulo_t]
    r0 = 1
    ws.cell(row=r0, column=c0, value=titulo_t).font = font(10, True, NAVY)
    ws.column_dimensions[get_column_letter(c0)].width = 22
    hr = r0 + 1
    for j, (nombre, _) in enumerate(series):
        c = ws.cell(row=hr, column=c0 + 1 + j, value=nombre)
        c.font, c.fill = font(9, True, "FFFFFF"), fill(NAVY)
        ws.column_dimensions[get_column_letter(c0 + 1 + j)].width = 16
    for i, cat in enumerate(categorias):
        ws.cell(row=hr + 1 + i, column=c0, value=cat).font = font(9, True)
        for j, (_, vals) in enumerate(series):
            c = ws.cell(row=hr + 1 + i, column=c0 + 1 + j, value=vals[i])
            c.number_format = "#,##0.00"
            c.fill = fill(CALC_BG)
    sh = ws.title
    col = lambda k: get_column_letter(c0 + k)
    last = hr + len(categorias)
    return {
        "cat": (f"{sh}!${col(0)}${hr + 1}:${col(0)}${last}", c0, hr + 1, last),
        "series": [(f"{sh}!${col(1 + j)}${hr}", f"{sh}!${col(1 + j)}${hr + 1}:${col(1 + j)}${last}", c0 + 1 + j)
                   for j in range(len(series))],
        "sheet": sh,
    }


# ----------------------------- Diapositiva 1 --------------------------------
def ppt1(wb):
    h = Hoja(wb, "PPT_1")
    V = "Ventas_m2!"
    b1 = encabezado(
        h, "S1_Encabezado",
        '"Ventas "&Config!$B$3&" · Volumen (m²)"',
        '"F\'cast vs Budget vs Real Sales por línea y cumplimiento semanal · Corte W"&Config!$B$5',
        [("Real Sales Premium", f'{N0(V + "$M$7")}&" m²"', f'"del Budget ("&{N0(V + "$D$7")}&")"', V + "$P$7"),
         ("Cierre proyectado Premium", f'{N0(V + "$N$7")}&" m²"', f'"del Budget · F\'cast "&{N0(V + "$C$7")}',
          f"{V}$N$7/{V}$D$7"),
         ("Real Sales Estándar", f'{N0(V + "$M$11")}&" m²"', f'"del Budget ("&{N0(V + "$D$11")}&")"', V + "$P$11"),
         ("Gap total vs Budget", f'{N0(V + "$Q$13")}&" m²"',
          f'"Real "&{N0(V + "$M$13")}&" de "&{N0(V + "$D$13")}', V + "$P$13")])
    filas = []
    for r, nombre, tipo, _ in V_ROWS:
        est = {"L": "L", "S": "S", "T": "T"}[tipo]
        cells = [f"={V}A{r}"] + [{"v": f"={V}{L}{r}", "fmt": "#,##0"} for L in "DCMN"]
        cells.append({"v": f'=IF({V}P{r}="","–",{V}P{r})', "fmt": "0%", "pct": True})
        for L in "RSTU":
            cells.append({"v": f'=IF({V}{L}{r}="","–",{V}{L}{r})', "fmt": "0%", "pct": True})
        filas.append((cells, "S" if est == "T" else est))
    tb = tabla(h, "S1_Tabla", [170, 80, 82, 82, 82, 86, 44, 44, 44, 44],
               ["Línea", "Budget", "F'cast Act.", "Real Sales", "Cierre proy.", "% Real vs Bgt", "W1", "W2", "W3", "W4"],
               28, 20.5, filas, nota="W1–W4 = cumplimiento semanal: Real/Proy de la semana ÷ F'cast semanal.")
    # Alertas: Premium + 3 líneas con menor % Real vs Budget
    ws = h.ws
    ka, kn, kp = h.aux(), h.aux(), h.aux()
    for i, r in enumerate(V_LINEAS):
        rr = 40 + i
        ws.cell(row=rr, column=kn, value=f"={V}A{r}")
        ws.cell(row=rr, column=kp, value=f'=IF(ISNUMBER({V}P{r}),{V}P{r},"")')
        ws.cell(row=rr, column=ka, value=f'=IF(ISNUMBER({V}P{r}),{V}P{r}+{i + 1}/1000000,"")')
    rng = lambda col: f"${get_column_letter(col)}$40:${get_column_letter(col)}${39 + len(V_LINEAS)}"
    cand = [(f'IF(Comentarios!$B${5 + i}="","",Comentarios!$B${5 + i})', "4") for i in range(2)]
    p = f"{V}$N$7/{V}$D$7"
    cand.append((f'"Premium proyecta cerrar en "&{N0(V + "$N$7")}&" m² ("&{PC(p)}&" del Budget); brecha de "&'
                 f'{N0(V + "$N$7-" + V + "$D$7")}&" m²."', SEV(p)))
    for k in range(1, 4):
        idx = f"MATCH(SMALL({rng(ka)},{k}),{rng(ka)},0)"
        nm = f"INDEX({rng(kn)},{idx})"
        row = f"INDEX({{{','.join(str(r) for r in V_LINEAS)}}},{idx})"
        real = f"INDEX({V}$M$1:$M$20,{row})"
        bud = f"INDEX({V}$D$1:$D$20,{row})"
        pp = f"INDEX({rng(kp)},{idx})"
        cand.append((f'{nm}&": Real "&{N0(real)}&" m² = "&{PC(pp)}&" del Budget (gap "&{N0(real + "-" + bud)}&" m²)."',
                     SEV(pp)))
    al = alertas(h, "S1_Alertas", [26, 732], cand, 4, 16)
    # Gráficos
    g1 = datos_grafico(h, "Premium semanal (m²)", ["W1", "W2", "W3", "W4"],
                       [("F'cast semanal", [f"={V}{L}7" for L in "EFGH"]),
                        ("Real/Proy semanal", [f"={V}{L}7" for L in "IJKL"])])
    g2 = datos_grafico(h, "Budget vs F'cast vs Real (m²)", [f"={V}A{r}" for r in V_LINEAS],
                       [("Budget", [f"={V}D{r}" for r in V_LINEAS]), ("F'cast actual", [f"={V}C{r}" for r in V_LINEAS]),
                        ("Real Sales", [f"={V}M{r}" for r in V_LINEAS])])
    return h, {"enc": b1, "tabla": tb, "alertas": al}, (g1, g2)


# ----------------------------- Diapositiva 2 --------------------------------
def ppt2(wb):
    h = Hoja(wb, "PPT_2")
    R = "Revenue!"
    mix_p = f"{R}$D$7/({R}$D$19-{R}$D$18)"
    num_m2 = (f'(SUMIF({R}$E$5:$E$6,">0",{R}$C$5:$C$6)+SUMIF({R}$E$8:$E$11,">0",{R}$C$8:$C$11)'
              f'+IF(N({R}$E$13)>0,{R}$C$13,0))')
    bud_m2 = f"({R}$E$7+{R}$E$12+N({R}$E$13))"
    b1 = encabezado(
        h, "S2_Encabezado",
        '"Ventas "&Config!$B$3&" · m² y Revenue"',
        '"F\'cast actual del mes en m² y valorizado en COP: precio promedio, mix y performance vs Budget"',
        [("Revenue F'cast total", MM(R + "$D$19"), f'"incluye ICO "&{MM(R + "$D$18")}', None),
         ("Revenue Premium", MM(R + "$D$7"), f'{PC(mix_p)}&" del revenue de producto"', None),
         ("Volumen F'cast (m²)", f'{N0(f"{R}$C$7+{R}$C$12+N({R}$C$13)")}&" m²"',
          f'"del Budget ("&{N0(bud_m2)}&" m²)"', f"{num_m2}/{bud_m2}"),
         ("Precio prom. Premium / m²", f'"$ "&{N0(R + "$G$7")}', f'"Estándar: $ "&{N0(R + "$G$12")}&" / m²"', None)])
    filas = []
    for r, nombre, uni, *_ , tipo in R_ROWS:
        est = "T" if tipo == "T" else ("S" if tipo.startswith("S") else "L")
        cant = f'=IF(N({R}C{r})=0,"–",{N0(R + "C" + str(r))}&IF(OR({R}B{r}="m²",{R}B{r}="",{R}B{r}="—"),""," "&{R}B{r}))'
        precio = f'=IF({R}G{r}="","–","$ "&{N0(R + "G" + str(r))})'
        bud = (f'=IF(N({R}E{r})>0,{N0(R + "E" + str(r))},IF(N({R}F{r})>0,{MM(R + "F" + str(r))},"–"))')
        cells = [f"={R}A{r}", {"v": "" if est == "T" else cant}, {"v": "" if est == "T" else precio},
                 {"v": f"={R}D{r}", "fmt": '"$ "#,##0'}, {"v": f"={R}I{r}", "fmt": "0%"},
                 {"v": "" if est == "T" else bud},
                 {"v": "" if est == "T" else f'=IF({R}H{r}="","–",{R}H{r})', "fmt": "0%", "pct": est != "T"}]
        filas.append((cells, est))
    tb = tabla(h, "S2_Tabla", [178, 88, 92, 142, 56, 110, 92],
               ["Línea", "Cantidad", "Precio prom.", "Revenue (COP)", "Mix", "Budget", "% Perf."],
               22, 17, filas, size=9)
    # Alertas
    ws = h.ws
    kt = h.aux()
    for i, r in enumerate(R_LINEAS):
        ws.cell(row=40 + i, column=kt,
                value=f'=IF(AND(ISNUMBER({R}H{r}),{R}H{r}<$BA$2),{R}A{r}&" "&{PC(R + "H" + str(r))}&", ","")')
    L = get_column_letter(kt)
    lista = "&".join(f"{L}{40 + i}" for i in range(len(R_LINEAS)))
    kj = h.aux()
    ws.cell(row=40, column=kj, value=f'=IF(({lista})="","",LEFT({lista},LEN({lista})-2))')
    jref = f"${get_column_letter(kj)}$40"
    cand = [(f'IF(Comentarios!$B${7 + i}="","",Comentarios!$B${7 + i})', "4") for i in range(2)]
    cand.append((f'"Premium concentra "&{PC(mix_p)}&" del revenue de producto ("&{MM(R + "$D$7")}&")."', "4"))
    cand.append((f'IF({jref}="","Todas las líneas con Budget están sobre el umbral amarillo.",'
                 f'"Bajo umbral vs Budget: "&{jref}&".")', f'IF({jref}="",3,1)'))
    al = alertas(h, "S2_Alertas", [26, 732], cand, 2, 16, titulo_t=None, size=9)
    # Gráficos
    grupos = [("Premium", f"{R}D7"), ("Estándar", f"{R}D12"), ("Arquitectónica", f"{R}D13"), ("Proyectos", f"{R}D14"),
              ("Accesorios", f"{R}D15+{R}D16"), ("Puertas", f"{R}D17"), ("ICO", f"{R}D18")]
    g1 = datos_grafico(h, "Mix de revenue por grupo (M COP)", [g for g, _ in grupos],
                       [("Revenue (M COP)", [f"=({x})/1000000" for _, x in grupos])])
    # Revenue por línea ordenado de mayor a menor (auxiliares + LARGE)
    kn, kv, kk = h.aux(), h.aux(), h.aux()
    for i, r in enumerate(R_LINEAS):
        rr = 60 + i
        ws.cell(row=rr, column=kn, value=f"={R}A{r}")
        ws.cell(row=rr, column=kv, value=f"=N({R}D{r})/1000000")
        ws.cell(row=rr, column=kk, value=f"={get_column_letter(kv)}{rr}+({len(R_LINEAS)}-{i})/1000000000")
    rg = lambda col: f"${get_column_letter(col)}$60:${get_column_letter(col)}${59 + len(R_LINEAS)}"
    idx = lambda k: f"MATCH(LARGE({rg(kk)},{k}),{rg(kk)},0)"
    g2 = datos_grafico(h, "Revenue por línea (millones COP)", [f"=INDEX({rg(kn)},{idx(k)})" for k in range(1, len(R_LINEAS) + 1)],
                       [("Revenue (M COP)", [f"=INDEX({rg(kv)},{idx(k)})" for k in range(1, len(R_LINEAS) + 1)])])
    return h, {"enc": b1, "tabla": tb, "alertas": al}, (g1, g2)


# ----------------------------- Diapositiva 3 --------------------------------
def ppt3(wb):
    h = Hoja(wb, "PPT_3")
    O, P = "Order_Intake!", "Pipeline!"
    bm = "Config!$B$11"
    gan = f"{O}$C$9"
    pond_mes = f"{O}$G$9"
    brutos_sig = f'({P}$D$31-SUM(SUMIFS({P}$D$5:$D$29,{P}$F$5:$F$29,{{"W1","W2","W3","W4"}})))'
    b1 = encabezado(
        h, "S3_Encabezado",
        '"Order Intake "&Config!$B$3',
        'Config!$B$10&" · Negocios ganados (m² adjudicados) vs Budget y proyección por ganar"',
        [("Negocios ganados", f'{N0(gan)}&" m²"', f'"del Budget ("&{N0(bm)}&")"', f"{gan}/{bm}"),
         ("Gap vs Budget", f'{N0(gan + "-" + bm)}&" m²"', f'"acumulado a W"&COUNT({O}$C$5:$C$8)', None),
         ("Cierre proyectado mes", f'{N0(gan + "+" + pond_mes)}&" m²"', f'"con pipeline ponderado "&{N0(pond_mes)}',
          f"({gan}+{pond_mes})/{bm}"),
         ("Pipeline mes siguiente", f'{N0(P + "$H$31-" + pond_mes)}&" m²"', f'"ponderado · "&{N0(brutos_sig)}&" m² brutos"',
          None)])
    ws = h.ws
    semanas = [(5, "W1"), (6, "W2"), (7, "W3"), (8, "W4")]
    fb = ["Budget"] + [{"v": f"={O}B{r}", "fmt": "#,##0"} for r, _ in semanas] + [{"v": f"={bm}", "fmt": "#,##0"}]
    fr = ["Ganado"] + [{"v": f'=IF({O}C{r}="","–",{N0(O + "C" + str(r))}&IF({O}D{r}="En curso","*",""))'}
                       for r, _ in semanas] + [{"v": f"={gan}", "fmt": "#,##0", "bold": True}]
    fg = ["Gap"] + [{"v": f'=IF({O}C{r}="","–",{O}E{r})', "fmt": "#,##0"} for r, _ in semanas] + \
         [{"v": f"={gan}-{bm}", "fmt": "#,##0", "bold": True}]
    fc = ["% Cumpl."] + [{"v": f'=IF({O}F{r}="","–",{O}F{r})', "fmt": "0%", "pct": True} for r, _ in semanas] + \
         [{"v": f"=IF({bm}=0,\"–\",{gan}/{bm})", "fmt": "0%", "pct": True}]
    tb = tabla(h, "S3_Tabla", [110, 70, 70, 70, 70, 76], ["Order Intake", "W1", "W2", "W3", "W4", "Total"],
               24, 22, [(fb, "L"), (fr, "L"), (fg, "L"), (fc, "L")],
               nota=f'=IF(COUNTIF({O}D5:D8,"En curso")>0,"* Semana en curso: valor parcial al corte.","")')
    # formato condicional: gap verde/rojo y semana en curso resaltada
    gap_rng = f"{tb.rel(3, 1)}:{tb.rel(3, 5)}"
    tl = tb.rel(3, 1)
    cf_expr(ws, gap_rng, f"AND(ISNUMBER({tl}),{tl}>=0)", GREEN)
    cf_expr(ws, gap_rng, f"AND(ISNUMBER({tl}),{tl}<0)", RED)
    for k, (r, _) in enumerate(semanas):
        cf_expr(ws, tb.rel(2, 1 + k), f'{O}$D${r}="En curso"', None, SUBT)
    # Pipeline top 5
    filas = []
    for k in range(1, 6):
        idx = f"MATCH(LARGE({P}$I$5:$I$29,{k}),{P}$I$5:$I$29,0)"
        g = lambda col: f'IFERROR(INDEX({P}${col}$5:${col}$29,{idx}),"")'
        vacio = "" if k > 1 else f'IF(COUNT({P}$I$5:$I$29)=0,"Sin oportunidades en la hoja Pipeline","")'
        nombre = f"=IFERROR(INDEX({P}$A$5:$A$29,{idx}),{vacio or chr(34) * 2})"
        filas.append(([nombre, f"={g('C')}", {"v": f"={g('D')}", "fmt": "#,##0"},
                       {"v": f"={g('E')}", "fmt": "0%", "align": "center"},
                       {"v": f'=IF(LEFT({g("F")},3)="Mes","Mes sig.",{g("F")})', "align": "center"},
                       {"v": f"={g('H')}", "fmt": "#,##0", "bold": True}], "L"))
    tp = tabla(h, "S3_Pipeline", [168, 92, 52, 42, 58, 54], ["Negocio por ganar", "Línea", "m²", "Prob.", "Cierre", "Pond."],
               22, 19.5, filas, size=9)
    # Alertas
    kt = h.aux()
    for i, (r, w) in enumerate(semanas):
        ws.cell(row=60 + i, column=kt,
                value=f'=IF(AND(ISNUMBER({O}F{r}),{O}F{r}<$BA$2),"{w} ("&{PC(O + "F" + str(r))}&"), ","")')
    L = get_column_letter(kt)
    lista = "&".join(f"{L}{60 + i}" for i in range(4))
    kj = h.aux()
    ws.cell(row=60, column=kj, value=f'=IF(({lista})="","",LEFT({lista},LEN({lista})-2))')
    jref = f"${get_column_letter(kj)}$60"
    p = f"{gan}/{bm}"
    gap = f"({bm}-{gan})"
    cand = [(f'IF(Comentarios!$B${9 + i}="","",Comentarios!$B${9 + i})', "4") for i in range(2)]
    cand.append((f'"Acumulado "&{PC(p)}&" del Budget; gap "&{N0(gan + "-" + bm)}&" m²."', SEV(p)))
    cand.append((f'IF({jref}="","Ninguna semana bajo el umbral amarillo.","Semanas bajo umbral: "&{jref}&".")',
                 f'IF({jref}="",3,1)'))
    cand.append((f'IF({gap}<=0,"Budget del mes cumplido.","Pipeline del mes cubre "&{PC(pond_mes + "/" + gap)}&'
                 f'" del gap ("&{N0(pond_mes)}&" de "&{N0(gap)}&" m²).")',
                 f"IF({gap}<=0,3,{SEV(pond_mes + '/' + gap)})"))
    al = alertas(h, "S3_Alertas", [20, 206], cand, 3, 37, titulo_t="ALERTAS", size=9, wrap=True)
    # Gráficos
    cats = [w for _, w in semanas]
    g1 = datos_grafico(h, "Order Intake semanal (m²)", cats,
                       [("Budget", [f"={O}B{r}" for r, _ in semanas]),
                        ("Ganado", [f"=N({O}C{r})" for r, _ in semanas]),
                        ("Por ganar (ponderado)", [f"={O}G{r}" for r, _ in semanas])])
    g2 = datos_grafico(h, "Acumulado del mes vs Budget (m²)", cats,
                       [("Budget acumulado", [f"={O}H{r}" for r, _ in semanas]),
                        ("Ganado acumulado", [f"={O}I{r}" for r, _ in semanas]),
                        ("Proyectado (ganado + pipeline)", [f"={O}J{r}" for r, _ in semanas])])
    return h, {"enc": b1, "tabla": tb, "pipeline": tp, "alertas": al}, (g1, g2)


def construir_excel(path):
    wb = Workbook()
    hoja_instrucciones(wb)
    hoja_config(wb)
    hoja_ventas(wb)
    hoja_revenue(wb)
    hoja_oi(wb)
    hoja_pipeline(wb)
    hoja_comentarios(wb)
    slides = [ppt1(wb), ppt2(wb), ppt3(wb)]
    for h, bloques, _ in slides:
        for b in bloques.values():
            wb.defined_names[b.nombre] = DefinedName(b.nombre, attr_text=_abs_ref(h.ws.title, b.area))
        h.ws.sheet_view.zoomScale = 90
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return slides


def _abs_ref(sheet, area):
    import re
    a, b = area.split(":")
    fx = lambda x: re.sub(r"([A-Z]+)(\d+)", r"$\1$\2", x)
    return f"{sheet}!{fx(a)}:{fx(b)}"


# ---------------------------------------------------------------------------
# LibreOffice: recálculo y vistas previas
# ---------------------------------------------------------------------------
def soffice(args, cwd):
    sk = next(Path("/root/.claude/skills/synced").glob("*/pptx/scripts"), None)
    if sk:
        sys.path.insert(0, str(sk))
        from office.soffice import run_soffice
        return run_soffice(args, cwd=cwd, capture_output=True, timeout=240)
    return subprocess.run(["soffice"] + args, cwd=cwd, capture_output=True, timeout=240)


def recalcular_y_previas(xlsx, slides, tmp):
    tmp = Path(tmp)
    # 1) libro recalculado (valores para los cachés de los gráficos)
    shutil.copy(xlsx, tmp / "calc_in.xlsx")
    soffice(["--headless", "--convert-to", "xlsx", "--outdir", str(tmp / "out"), str(tmp / "calc_in.xlsx")], tmp)
    vals = load_workbook(tmp / "out" / "calc_in.xlsx", data_only=True)
    # 2) PDF con un área de impresión por rango vinculado
    wb = load_workbook(xlsx)
    orden = []
    for ws in wb.worksheets:
        if not ws.title.startswith("PPT_"):
            ws.sheet_state = "hidden"
            continue
        h = next(s[0] for s in slides if s[0].ws.title == ws.title)
        areas = [b.area for b in h.bloques]
        orden += [b.nombre for b in h.bloques]
        ws.print_area = areas
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.scale = 100
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=False)
        for k in ("left", "right", "top", "bottom", "header", "footer"):
            setattr(ws.page_margins, k, 0)
    wb.active = [i for i, w in enumerate(wb.worksheets) if w.sheet_state == "visible"][0]
    wb.save(tmp / "prev.xlsx")
    soffice(["--headless", "--convert-to", "pdf", "--outdir", str(tmp / "out"), str(tmp / "prev.xlsx")], tmp)
    pdf = pymupdf.open(tmp / "out" / "prev.pdf")
    if len(pdf) != len(orden):
        raise RuntimeError(f"Se esperaban {len(orden)} páginas de vista previa y hubo {len(pdf)}")
    previas = {}
    for nombre, page in zip(orden, pdf):
        rects = [d["rect"] for d in page.get_drawings()] + [pymupdf.Rect(b[:4]) for b in page.get_text("blocks")]
        bbox = functools.reduce(lambda a, b: a | b, rects)
        bbox = pymupdf.Rect(0, 0, bbox.x1, bbox.y1) & page.rect
        pix = page.get_pixmap(clip=bbox, dpi=220)
        previas[nombre] = pix.tobytes("png")
    return vals, previas


# ---------------------------------------------------------------------------
# PowerPoint
# ---------------------------------------------------------------------------
def target_xlsx(carpeta, item=None):
    ruta = carpeta.rstrip("\\") + "\\" + XLSX_NAME
    return f"file:///{ruta}" + (f"!{item}" if item else "")


def ole_vinculado(slide, x, y, w, h, png, nombre, carpeta):
    img_part, img_rid = slide.part.get_or_add_image_part(io.BytesIO(png))
    link_rid = slide.part.relate_to(target_xlsx(carpeta, nombre), RT.OLE_OBJECT, is_external=True)
    sid = max([int(s.shape_id) for s in slide.shapes] + [1]) + 1
    X, Y, W, H = (int(Inches(v)) for v in (x, y, w, h))
    xml = f"""
<p:graphicFrame {nsdecls('p', 'a', 'r')}>
  <p:nvGraphicFramePr>
    <p:cNvPr id="{sid}" name="{nombre}" descr="Rango vinculado {nombre} de {XLSX_NAME}"/>
    <p:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></p:cNvGraphicFramePr>
    <p:nvPr/>
  </p:nvGraphicFramePr>
  <p:xfrm><a:off x="{X}" y="{Y}"/><a:ext cx="{W}" cy="{H}"/></p:xfrm>
  <a:graphic>
    <a:graphicData uri="http://schemas.openxmlformats.org/presentationml/2006/ole">
      <p:oleObj name="Worksheet" r:id="{link_rid}" imgW="{W}" imgH="{H}" progId="Excel.Sheet.12">
        <p:link updateAutomatic="1"/>
        <p:pic>
          <p:nvPicPr><p:cNvPr id="0" name=""/><p:cNvPicPr/><p:nvPr/></p:nvPicPr>
          <p:blipFill><a:blip r:embed="{img_rid}"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>
          <p:spPr><a:xfrm><a:off x="{X}" y="{Y}"/><a:ext cx="{W}" cy="{H}"/></a:xfrm>
            <a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>
        </p:pic>
      </p:oleObj>
    </a:graphicData>
  </a:graphic>
</p:graphicFrame>"""
    slide.shapes._spTree.append(parse_xml(xml))


def vincular_grafico(chart, g, carpeta):
    """Cambia el libro embebido por un vínculo externo al Excel base y apunta las series a PPT_n."""
    part = chart.part
    cs = chart._chartSpace
    ext = cs.find(qn("c:externalData"))
    old = ext.get(qn("r:id"))
    part.drop_rel(old)
    rid = part.relate_to(target_xlsx(carpeta), RT.OLE_OBJECT, is_external=True)
    ext.set(qn("r:id"), rid)
    for au in ext.findall(qn("c:autoUpdate")):
        ext.remove(au)
    au = etree.SubElement(ext, qn("c:autoUpdate"))
    au.set("val", "1")
    sers = cs.findall(".//" + qn("c:ser"))
    for ser, (name_ref, val_ref, _) in zip(sers, g["series"]):
        ser.find(qn("c:tx")).find(".//" + qn("c:f")).text = name_ref
        cat = ser.find(qn("c:cat"))
        if cat is not None:
            cat.find(".//" + qn("c:f")).text = g["cat"][0]
        ser.find(qn("c:val")).find(".//" + qn("c:f")).text = val_ref


def leer(vals, g):
    ws = vals[g["sheet"]]
    _, c0, r1, r2 = g["cat"]
    cats = [str(ws.cell(row=r, column=c0).value or "") for r in range(r1, r2 + 1)]
    ser = []
    for j, (nm_ref, _, col) in enumerate(g["series"]):
        nombre = ws.cell(row=r1 - 1, column=col).value
        v = []
        for r in range(r1, r2 + 1):
            x = ws.cell(row=r, column=col).value
            v.append(x if isinstance(x, (int, float)) else None)
        ser.append((nombre, v))
    return cats, ser


def texto_fijo(slide, n, pie):
    G.texto(slide, 0.45, 7.12, 10.5, 0.25, pie, size=9, color=G.MUTED)
    G.texto(slide, 12.4, 7.12, 0.5, 0.25, str(n), size=9, color=G.MUTED, align=PP_ALIGN.RIGHT)


def poner(slide, b, previas, x, y, carpeta, ancho=None):
    w, h = b.size_in
    if ancho:
        h, w = h * ancho / w, ancho
    ole_vinculado(slide, x, y, w, h, previas[b.nombre], b.nombre, carpeta)
    return y + h


def construir_pptx(path, slides, vals, previas, carpeta):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    pie = (f"Fuente: {XLSX_NAME} (vinculado) · Para actualizar: Archivo > Información > Editar vínculos a archivos > "
           "Actualizar ahora")
    X0 = 0.45

    # ---- Diapositiva 1
    (h1, b1, (g1, g2)) = slides[0]
    s = prs.slides.add_slide(prs.slide_layouts[6])
    poner(s, b1["enc"], previas, X0, 0.25, carpeta)
    y = poner(s, b1["tabla"], previas, X0, 2.55, carpeta)
    poner(s, b1["alertas"], previas, X0, max(y + 0.1, 5.85), carpeta)
    cats, ser = leer(vals, g1)
    ch = G.barras(s, 8.6, 2.45, 4.3, 2.3, XL_CHART_TYPE.COLUMN_CLUSTERED, cats, ser, "Premium semanal (m²)",
                  [G.STEEL, G.NAVY], gap=55, overlap=-5, ocultar_valores=True)
    vincular_grafico(ch, g1, carpeta)
    cats, ser = leer(vals, g2)
    ch = G.barras(s, 8.6, 4.8, 4.3, 2.3, XL_CHART_TYPE.BAR_CLUSTERED, cats, ser, "Budget vs F'cast vs Real (m²)",
                  [G.GREY, G.STEEL, G.NAVY], gap=40, overlap=0, etiquetar=[2], ocultar_valores=True)
    ch.category_axis.reverse_order = True
    ch.category_axis.tick_labels.font.size = Pt(8)
    vincular_grafico(ch, g2, carpeta)
    texto_fijo(s, 1, pie)
    s.notes_slide.notes_text_frame.text = (
        "Diapositiva vinculada a Base_Comite_Ventas.xlsx (hoja PPT_1). Encabezado, KPIs, tabla y alertas son rangos "
        "de Excel vinculados; los gráficos toman sus datos del mismo libro.")

    # ---- Diapositiva 2
    (h2, b2, (g3, g4)) = slides[1]
    s = prs.slides.add_slide(prs.slide_layouts[6])
    poner(s, b2["enc"], previas, X0, 0.25, carpeta)
    y = poner(s, b2["tabla"], previas, X0, 2.5, carpeta)
    poner(s, b2["alertas"], previas, X0, y + 0.08, carpeta)
    cats, ser = leer(vals, g3)
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series(*ser[0])
    ch = s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(8.6), Inches(2.4), Inches(4.3), Inches(2.35), cd).chart
    G.estilo_chart(ch, "Mix de revenue por grupo")
    ch.legend.position = XL_LEGEND_POSITION.RIGHT
    plot = ch.plots[0]
    for i, g in enumerate(cats):
        pt = plot.series[0].points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = G.GROUP_COLORS.get(g, G.GREY)
    dl = plot.series[0].data_labels
    dl.show_percentage, dl.show_value = True, False
    dl.number_format, dl.number_format_is_linked = '[<0.04]"";0%', False
    dl.font.size, dl.font.color.rgb, dl.font.bold = Pt(8), G.WHITE, True
    hole = ch._chartSpace.find(".//" + qn("c:holeSize"))
    if hole is not None:
        hole.set("val", "55")
    vincular_grafico(ch, g3, carpeta)
    cats, ser = leer(vals, g4)
    ch = G.barras(s, 8.6, 4.8, 4.3, 2.3, XL_CHART_TYPE.BAR_CLUSTERED, cats, ser, "Revenue por línea (millones COP)",
                  [G.NAVY], gap=35, leyenda=False, ocultar_valores=True)
    ch.category_axis.reverse_order = True
    ch.category_axis.tick_labels.font.size = Pt(8)
    vincular_grafico(ch, g4, carpeta)
    texto_fijo(s, 2, pie)
    s.notes_slide.notes_text_frame.text = (
        "Diapositiva vinculada a Base_Comite_Ventas.xlsx (hoja PPT_2). El subtotal Estándar incluye Traslúcida; "
        "su % de performance se mide solo sobre las líneas que tienen Budget.")

    # ---- Diapositiva 3
    (h3, b3, (g5, g6)) = slides[2]
    s = prs.slides.add_slide(prs.slide_layouts[6])
    poner(s, b3["enc"], previas, X0, 0.25, carpeta)
    cats, ser = leer(vals, g5)
    ch = G.barras(s, 0.35, 2.45, 6.2, 2.45, XL_CHART_TYPE.COLUMN_CLUSTERED, cats, ser, "Order Intake semanal (m²)",
                  [G.GREY, G.NAVY, G.STEEL], gap=55, overlap=-5, etiquetar=[1, 2], ocultar_valores=True)
    vincular_grafico(ch, g5, carpeta)
    cats, ser = leer(vals, g6)
    cd = CategoryChartData()
    cd.categories = cats
    for nm, v in ser:
        cd.add_series(nm, v)
    ch = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(6.75), Inches(2.45), Inches(6.15), Inches(2.45), cd).chart
    G.estilo_chart(ch, "Acumulado del mes vs Budget (m²)")
    G.ejes(ch, ocultar_valores=True)
    for sr, c, dash in zip(ch.plots[0].series, (G.GREY, G.NAVY, G.BLUE), (False, False, True)):
        sr.smooth = False
        sr.format.line.color.rgb = c
        sr.format.line.width = Pt(2.25)
        if dash:
            sr.format.line.dash_style = 4
        sr.marker.style = XL_MARKER_STYLE.CIRCLE
        sr.marker.size = 6
        sr.marker.format.fill.solid()
        sr.marker.format.fill.fore_color.rgb = c
        sr.marker.format.line.color.rgb = c
    G.etiquetas(ch.plots[0].series[1], pos=XL_LABEL_POSITION.BELOW, color=G.NAVY, bold=True)
    G.etiquetas(ch.plots[0].series[0], pos=XL_LABEL_POSITION.ABOVE, color=G.MUTED)
    vincular_grafico(ch, g6, carpeta)
    poner(s, b3["tabla"], previas, X0, 5.1, carpeta)
    poner(s, b3["pipeline"], previas, 5.5, 5.1, carpeta)
    poner(s, b3["alertas"], previas, 10.55, 5.05, carpeta)
    texto_fijo(s, 3, pie)
    s.notes_slide.notes_text_frame.text = (
        "Diapositiva vinculada a Base_Comite_Ventas.xlsx (hoja PPT_3). Proyección = pipeline ponderado "
        "(m² × probabilidad) de la hoja Pipeline.")
    prs.save(path)


# ---------------------------------------------------------------------------
PS1 = r'''# Apunta todos los vínculos de las presentaciones de esta carpeta al archivo
# Base_Comite_Ventas.xlsx de esta misma carpeta. Cierre PowerPoint antes de ejecutar.
$ErrorActionPreference = "Stop"
$carpeta = Split-Path -Parent $MyInvocation.MyCommand.Path
$xlsx = Join-Path $carpeta "Base_Comite_Ventas.xlsx"
if (-not (Test-Path $xlsx)) { Write-Host "No se encontró $xlsx"; Read-Host "Enter para salir"; exit 1 }

function Encode-Ruta([string]$ruta) {
    $sb = New-Object System.Text.StringBuilder
    foreach ($b in [System.Text.Encoding]::UTF8.GetBytes($ruta)) {
        $ch = [char]$b
        if (($b -lt 128) -and ("ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789-_.~:\/()',;=+$!".IndexOf($ch) -ge 0)) {
            [void]$sb.Append($ch)
        } else {
            [void]$sb.Append('%' + $b.ToString('X2'))
        }
    }
    return $sb.ToString()
}

Add-Type -AssemblyName System.IO.Compression
Add-Type -AssemblyName System.IO.Compression.FileSystem
$nuevo = 'Target="file:///' + (Encode-Ruta $xlsx)
$patron = 'Target="file:///[^"!]*?Base_Comite_Ventas\.xlsx'
$utf8 = New-Object System.Text.UTF8Encoding($false)
$total = 0
foreach ($pptx in Get-ChildItem -Path $carpeta -Filter *.pptx) {
    try {
        $zip = [System.IO.Compression.ZipFile]::Open($pptx.FullName, 'Update')
    } catch {
        Write-Host ("{0}: no se pudo abrir (¿está abierto en PowerPoint? ciérrelo y repita)" -f $pptx.Name)
        continue
    }
    $n = 0
    try {
        foreach ($e in @($zip.Entries)) {
            if (-not $e.FullName.EndsWith('.rels')) { continue }
            $s = $e.Open()
            $sr = New-Object System.IO.StreamReader($s, $utf8, $true, 4096, $true)
            $txt = $sr.ReadToEnd()
            $sr.Dispose()
            $cuenta = ([regex]::Matches($txt, $patron)).Count
            if ($cuenta -gt 0) {
                $txt2 = [regex]::Replace($txt, $patron, $nuevo.Replace('$', '$$'))
                $s.SetLength(0)
                $s.Position = 0
                $sw = New-Object System.IO.StreamWriter($s, $utf8)
                $sw.Write($txt2)
                $sw.Flush()
                $sw.Dispose()
                $n += $cuenta
            } else {
                $s.Dispose()
            }
        }
    } finally { $zip.Dispose() }
    Write-Host ("{0}: {1} vínculos actualizados" -f $pptx.Name, $n)
    $total += $n
}
Write-Host ""
Write-Host "Listo. Vínculos apuntando a: $xlsx"
Write-Host "Abra la presentación y elija 'Actualizar vínculos'."
Read-Host "Enter para salir"
'''

BAT = ('@echo off\r\n'
       'powershell -NoProfile -ExecutionPolicy Bypass -File "%~dp0Actualizar_Rutas.ps1"\r\n')

LEEME = f"""COMITÉ DE OPERACIONES · VENTAS, REVENUE Y ORDER INTAKE (VERSIÓN VINCULADA)
=========================================================================

Archivos (deben estar en la MISMA carpeta):
  {XLSX_NAME}        Excel base: aquí se diligencia todo cada semana.
  {PPTX_NAME}   Presentación de 3 diapositivas vinculada al Excel.
  Actualizar_Rutas.bat / .ps1   Apunta los vínculos al Excel de esta carpeta.

PRIMERA VEZ
  a) Si la carpeta es {RUTA_DEFECTO}  -> no hay que hacer nada.
  b) Si usa otra carpeta: con PowerPoint CERRADO, doble clic en Actualizar_Rutas.bat.
     (Repítalo si mueve o renombra la carpeta.)

CADA SEMANA
  1. Abra {XLSX_NAME}. Diligencie las celdas AMARILLAS de las hojas
     Config, Ventas_m2, Revenue, Order_Intake, Pipeline y Comentarios. Guarde.
  2. Abra {PPTX_NAME}. En el aviso de seguridad elija "Actualizar vínculos".
     Si no aparece: Archivo > Información > Editar vínculos a archivos >
     seleccione todos (Ctrl+clic) > Actualizar ahora.
  3. Para guardar la versión de la semana: Archivo > Guardar como
     (p. ej. Comite_Ventas_Octubre_W1.pptx). La copia conserva los vínculos;
     si quiere congelarla, en Editar vínculos use "Romper vínculo".

QUÉ SE ACTUALIZA
  - Encabezado con mes, corte y fecha; tarjetas KPI con semáforo.
  - Tablas (volumen, revenue, order intake semanal, top 5 del pipeline).
  - Alertas (textos armados con fórmulas; primero van sus comentarios).
  - Los 6 gráficos (clic derecho > Editar datos abre el Excel base).

NOTAS
  - No cambie el nombre del Excel ni de las hojas PPT_1, PPT_2, PPT_3:
    los vínculos usan rangos con nombre (S1_Encabezado, S1_Tabla, ...).
  - PowerPoint para la web / Teams no actualiza vínculos: actualice en la
    aplicación de escritorio y luego comparta.
  - Las filas del Pipeline son de EJEMPLO: reemplácelas por las reales.
"""


def main():
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "vinculado"
    out.mkdir(parents=True, exist_ok=True)
    xlsx = out / XLSX_NAME
    slides = construir_excel(xlsx)
    with tempfile.TemporaryDirectory(dir=ROOT / "salidas" if (ROOT / "salidas").exists() else None) as tmp:
        vals, previas = recalcular_y_previas(xlsx, slides, tmp)
    construir_pptx(out / PPTX_NAME, slides, vals, previas, RUTA_DEFECTO)
    (out / "Actualizar_Rutas.ps1").write_text(PS1, encoding="utf-8-sig")
    (out / "Actualizar_Rutas.bat").write_bytes(BAT.encode("ascii"))
    (out / "LEEME.txt").write_text(LEEME.replace("\n", "\r\n"), encoding="utf-8-sig")
    print(f"Listo: {out}")


if __name__ == "__main__":
    main()
