"""
Genera automaticamente el PowerPoint semanal del Comite de Operaciones
a partir de la plantilla Excel (plantilla/Plantilla_Semanal_Comite_Operaciones.xlsx).

Flujo de uso semanal:
    1. El responsable de cada area llena la plantilla Excel (celdas amarillas).
    2. Guarda el archivo con Excel/LibreOffice (para que las formulas queden calculadas).
    3. Ejecuta:
         python scripts/generar_ppt_comite.py plantilla/Plantilla_Semanal_Comite_Operaciones.xlsx
    4. Se genera automaticamente: salidas/Comite_Operaciones_<fecha>.pptx

No requiere una plantilla .pptx fija: el diseño se construye en este script,
así que basta con actualizar el Excel cada semana para tener las diapositivas
listas, con el mismo formato y colores siempre.

Dependencias: openpyxl, python-pptx  (pip install -r requirements.txt)
"""

import sys
import datetime
from pathlib import Path

from openpyxl import load_workbook

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

# ---------------------------------------------------------------------------
# Paleta y estilo (Midnight Executive)
# ---------------------------------------------------------------------------
NAVY = RGBColor(0x1E, 0x27, 0x61)
ICE = RGBColor(0xCA, 0xDC, 0xFC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_TEXT = RGBColor(0x22, 0x22, 0x22)
GREY_TEXT = RGBColor(0x66, 0x66, 0x66)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
AMBER = RGBColor(0xE6, 0xA8, 0x17)
RED = RGBColor(0xC6, 0x28, 0x28)
LIGHT_BG = RGBColor(0xF5, 0xF7, 0xFC)

RAG_COLOR = {"Verde": GREEN, "Amarillo": AMBER, "Rojo": RED,
             "Alto": RED, "Medio": AMBER, "Bajo": GREEN,
             "Alta": RED, "Media": AMBER, "Baja": GREEN,
             "Completado": GREEN, "En curso": AMBER, "Pendiente": GREY_TEXT, "Atrasado": RED}

FONT = "Calibri"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


# ---------------------------------------------------------------------------
# Lectura de datos desde Excel
# ---------------------------------------------------------------------------
def leer_filas(ws, header_row, start_col, ncols):
    """Lee filas de una tabla hasta encontrar una fila completamente vacia."""
    filas = []
    row = header_row + 1
    while True:
        valores = [ws.cell(row=row, column=start_col + i).value for i in range(ncols)]
        if all(v is None or str(v).strip() == "" for v in valores):
            break
        filas.append(valores)
        row += 1
        if row > header_row + 500:  # limite de seguridad
            break
    return filas


def leer_lista_textos(ws, start_row, col=2, max_rows=50):
    textos = []
    for r in range(start_row, start_row + max_rows):
        v = ws.cell(row=r, column=col).value
        if v is None or str(v).strip() == "":
            break
        textos.append(str(v))
    return textos


def cargar_datos(xlsx_path):
    wb = load_workbook(xlsx_path, data_only=True)
    data = {}

    ws = wb["Portada"]
    data["comite"] = ws["D4"].value or "Comite de Operaciones"
    data["semana_inicio"] = ws["D5"].value
    data["semana_fin"] = ws["D6"].value
    data["fecha_presentacion"] = ws["D7"].value
    data["responsable"] = ws["D8"].value
    data["estado_general"] = ws["D9"].value or "Amarillo"

    ws = wb["Resumen Ejecutivo"]
    logros, atencion = [], []
    r = 5
    while True:
        v = ws.cell(row=r, column=2).value
        if v is None or str(v).strip() == "":
            break
        logros.append(str(v))
        r += 1
    r += 2  # salta la fila en blanco y el encabezado "Puntos de atencion..."
    while True:
        v = ws.cell(row=r, column=2).value
        if v is None or str(v).strip() == "":
            break
        atencion.append(str(v))
        r += 1
    data["logros"] = logros
    data["atencion"] = atencion

    ws = wb["KPIs"]
    filas = leer_filas(ws, header_row=4, start_col=2, ncols=9)
    data["kpis"] = [
        {
            "indicador": f[0], "unidad": f[1], "meta": f[2], "sentido": f[3],
            "anterior": f[4], "actual": f[5], "variacion": f[6],
            "cumplimiento": f[7], "estado": f[8],
        }
        for f in filas
    ]

    ws = wb["Proyectos"]
    filas = leer_filas(ws, header_row=4, start_col=2, ncols=7)
    data["proyectos"] = [
        {
            "nombre": f[0], "responsable": f[1], "avance": f[2], "estado": f[3],
            "hito": f[4], "fecha_hito": f[5], "comentario": f[6],
        }
        for f in filas
    ]

    ws = wb["Incidentes"]
    filas = leer_filas(ws, header_row=4, start_col=2, ncols=8)
    data["incidentes"] = [
        {
            "id": f[0], "fecha": f[1], "descripcion": f[2], "severidad": f[3],
            "estado": f[4], "responsable": f[5], "causa": f[6], "accion": f[7],
        }
        for f in filas
    ]

    ws = wb["Riesgos"]
    filas = leer_filas(ws, header_row=4, start_col=2, ncols=6)
    data["riesgos"] = [
        {
            "riesgo": f[0], "probabilidad": f[1], "impacto": f[2], "nivel": f[3],
            "mitigacion": f[4], "responsable": f[5],
        }
        for f in filas
    ]

    ws = wb["Plan de Accion"]
    filas = leer_filas(ws, header_row=4, start_col=2, ncols=5)
    data["acciones"] = [
        {
            "accion": f[0], "responsable": f[1], "fecha": f[2], "prioridad": f[3], "estado": f[4],
        }
        for f in filas
    ]

    ws = wb["Proximos Pasos"]
    data["proximos_pasos"] = leer_lista_textos(ws, start_row=4)

    return data


def fmt_fecha(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.strftime("%d/%m/%Y")
    return str(v)


# ---------------------------------------------------------------------------
# Construccion de diapositivas
# ---------------------------------------------------------------------------
def set_background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, size=14, bold=False,
                 color=DARK_TEXT, align=PP_ALIGN.LEFT, font=FONT, italic=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    return box


def add_title_bar(slide, titulo, subtitulo=None):
    add_textbox(slide, Inches(0.6), Inches(0.35), Inches(11.5), Inches(0.7),
                titulo, size=30, bold=True, color=NAVY)
    if subtitulo:
        add_textbox(slide, Inches(0.6), Inches(0.95), Inches(11.5), Inches(0.4),
                    subtitulo, size=13, color=GREY_TEXT)


def add_rag_badge(slide, left, top, width, height, texto, color_key=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RAG_COLOR.get(color_key if color_key is not None else texto, GREY_TEXT)
    shape.line.fill.background()
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = texto or ""
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = WHITE
    run.font.name = FONT
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return shape


def style_table(table, header_bg=NAVY, header_font_color=WHITE, font_size=11, header_size=12):
    for c, cell in enumerate(table.rows[0].cells):
        cell.fill.solid()
        cell.fill.fore_color.rgb = header_bg
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(header_size)
            p.font.bold = True
            p.font.color.rgb = header_font_color
            p.font.name = FONT
    for r in range(1, len(table.rows)):
        bg = LIGHT_BG if r % 2 == 0 else WHITE
        for cell in table.rows[r].cells:
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(font_size)
                p.font.name = FONT
                p.font.color.rgb = DARK_TEXT


def slide_portada(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, NAVY)
    add_textbox(slide, Inches(1), Inches(2.6), Inches(11.3), Inches(1.2),
                data["comite"], size=40, bold=True, color=WHITE)
    periodo = f'Semana del {fmt_fecha(data["semana_inicio"])} al {fmt_fecha(data["semana_fin"])}'
    add_textbox(slide, Inches(1), Inches(3.6), Inches(11.3), Inches(0.6),
                periodo, size=18, color=ICE)
    pie = f'Presentado por {data["responsable"] or ""} · {fmt_fecha(data["fecha_presentacion"])}'
    add_textbox(slide, Inches(1), Inches(6.6), Inches(11.3), Inches(0.5),
                pie, size=12, color=ICE)
    add_rag_badge(slide, Inches(1), Inches(4.3), Inches(2.6), Inches(0.55),
                  f'Estado general: {data["estado_general"]}', color_key=data["estado_general"])
    return slide


def slide_resumen(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, WHITE)
    add_title_bar(slide, "Resumen ejecutivo", "Principales avances y puntos de atencion de la semana")

    col_w = Inches(5.9)
    left1, left2 = Inches(0.6), Inches(6.8)
    top = Inches(1.6)

    add_textbox(slide, left1, top, col_w, Inches(0.4), "Logros y avances", size=18, bold=True, color=GREEN)
    y = top + Inches(0.55)
    for item in (data["logros"] or ["Sin datos registrados."])[:6]:
        box = slide.shapes.add_textbox(left1, y, col_w, Inches(0.7))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = f"•  {item}"
        run.font.size = Pt(13)
        run.font.name = FONT
        run.font.color.rgb = DARK_TEXT
        y += Inches(0.75)

    add_textbox(slide, left2, top, col_w, Inches(0.4), "Puntos de atencion", size=18, bold=True, color=RED)
    y = top + Inches(0.55)
    for item in (data["atencion"] or ["Sin datos registrados."])[:6]:
        box = slide.shapes.add_textbox(left2, y, col_w, Inches(0.7))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = f"•  {item}"
        run.font.size = Pt(13)
        run.font.name = FONT
        run.font.color.rgb = DARK_TEXT
        y += Inches(0.75)

    return slide


def slide_kpis_tabla(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, WHITE)
    add_title_bar(slide, "Indicadores operacionales (KPIs)", "Meta vs. resultado real de la semana")

    kpis = data["kpis"]
    if not kpis:
        add_textbox(slide, Inches(0.6), Inches(2), Inches(10), Inches(1), "Sin KPIs registrados esta semana.")
        return slide

    rows = len(kpis) + 1
    cols = 6
    left, top, width, height = Inches(0.6), Inches(1.6), Inches(12.1), Inches(0.5 * rows)
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table
    headers = ["Indicador", "Meta", "Sem. Anterior", "Sem. Actual", "% Cumplimiento", "Estado"]
    for c, h in enumerate(headers):
        table.cell(0, c).text = h

    for i, k in enumerate(kpis, start=1):
        table.cell(i, 0).text = f'{k["indicador"]} ({k["unidad"]})' if k["unidad"] else str(k["indicador"])
        table.cell(i, 1).text = "" if k["meta"] is None else str(k["meta"])
        table.cell(i, 2).text = "" if k["anterior"] is None else str(k["anterior"])
        table.cell(i, 3).text = "" if k["actual"] is None else str(k["actual"])
        cumpl = k["cumplimiento"]
        table.cell(i, 4).text = f"{cumpl*100:.0f}%" if isinstance(cumpl, (int, float)) else ""
        table.cell(i, 5).text = str(k["estado"] or "")

    style_table(table)
    table.columns[0].width = Inches(4.6)
    for c in range(1, 6):
        table.columns[c].width = Inches((12.1 - 4.6) / 5)

    # Colorear celda de Estado segun RAG
    for i, k in enumerate(kpis, start=1):
        cell = table.cell(i, 5)
        color = RAG_COLOR.get(k["estado"], LIGHT_BG)
        cell.fill.solid()
        cell.fill.fore_color.rgb = color
        for p in cell.text_frame.paragraphs:
            p.font.color.rgb = WHITE
            p.font.bold = True

    return slide


def slide_kpis_grafico(prs, data):
    kpis = [k for k in data["kpis"] if isinstance(k.get("meta"), (int, float)) and isinstance(k.get("actual"), (int, float))]
    if not kpis:
        return None
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, WHITE)
    add_title_bar(slide, "Tendencia de KPIs", "Comparativo semana anterior vs. semana actual vs. meta")

    chart_data = CategoryChartData()
    chart_data.categories = [k["indicador"] for k in kpis]
    chart_data.add_series("Semana anterior", tuple(k["anterior"] if isinstance(k["anterior"], (int, float)) else 0 for k in kpis))
    chart_data.add_series("Semana actual", tuple(k["actual"] for k in kpis))
    chart_data.add_series("Meta", tuple(k["meta"] for k in kpis))

    x, y, cx, cy = Inches(0.6), Inches(1.6), Inches(12.1), Inches(5.3)
    graphic_frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, cx, cy, chart_data)
    chart = graphic_frame.chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.has_title = False

    colors = [NAVY, ICE, AMBER]
    for i, series in enumerate(chart.plots[0].series):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = colors[i % len(colors)]

    cat_axis = chart.category_axis
    cat_axis.tick_labels.font.size = Pt(10)
    val_axis = chart.value_axis
    val_axis.tick_labels.font.size = Pt(10)
    val_axis.has_major_gridlines = True

    return slide


def slide_proyectos(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, WHITE)
    add_title_bar(slide, "Avance de proyectos e iniciativas", "% de avance, estado y proximos hitos")

    proyectos = data["proyectos"]
    if not proyectos:
        add_textbox(slide, Inches(0.6), Inches(2), Inches(10), Inches(1), "Sin proyectos registrados esta semana.")
        return slide

    rows = len(proyectos) + 1
    cols = 6
    left, top, width, height = Inches(0.6), Inches(1.55), Inches(12.1), Inches(0.5 * rows)
    table = slide.shapes.add_table(rows, cols, left, top, width, height).table
    headers = ["Proyecto", "Responsable", "% Avance", "Estado", "Proximo hito", "Fecha"]
    for c, h in enumerate(headers):
        table.cell(0, c).text = h

    for i, p in enumerate(proyectos, start=1):
        table.cell(i, 0).text = str(p["nombre"] or "")
        table.cell(i, 1).text = str(p["responsable"] or "")
        avance = p["avance"]
        table.cell(i, 2).text = f"{avance*100:.0f}%" if isinstance(avance, (int, float)) else str(avance or "")
        table.cell(i, 3).text = str(p["estado"] or "")
        table.cell(i, 4).text = str(p["hito"] or "")
        table.cell(i, 5).text = fmt_fecha(p["fecha_hito"])

    style_table(table)
    widths = [3.4, 2.0, 1.4, 1.4, 2.6, 1.3]
    for c, w in enumerate(widths):
        table.columns[c].width = Inches(w)

    for i, p in enumerate(proyectos, start=1):
        cell = table.cell(i, 3)
        color = RAG_COLOR.get(p["estado"], LIGHT_BG)
        cell.fill.solid()
        cell.fill.fore_color.rgb = color
        for para in cell.text_frame.paragraphs:
            para.font.color.rgb = WHITE
            para.font.bold = True

    return slide


def slide_incidentes(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, WHITE)
    add_title_bar(slide, "Incidentes y problemas", "Eventos relevantes ocurridos durante la semana")

    incidentes = data["incidentes"]
    if not incidentes:
        add_textbox(slide, Inches(0.6), Inches(2), Inches(10), Inches(1), "Sin incidentes registrados esta semana.")
        return slide

    rows = len(incidentes) + 1
    cols = 6
    left, top, width, height = Inches(0.6), Inches(1.55), Inches(12.1), Inches(0.5 * rows)
    table = slide.shapes.add_table(rows, cols, left, top, width, height).table
    headers = ["ID", "Fecha", "Descripcion", "Severidad", "Estado", "Causa raiz"]
    for c, h in enumerate(headers):
        table.cell(0, c).text = h

    for i, inc in enumerate(incidentes, start=1):
        table.cell(i, 0).text = str(inc["id"] or "")
        table.cell(i, 1).text = fmt_fecha(inc["fecha"])
        table.cell(i, 2).text = str(inc["descripcion"] or "")
        table.cell(i, 3).text = str(inc["severidad"] or "")
        table.cell(i, 4).text = str(inc["estado"] or "")
        table.cell(i, 5).text = str(inc["causa"] or "")

    style_table(table, font_size=10.5)
    widths = [1.0, 1.2, 3.6, 1.3, 1.8, 3.2]
    for c, w in enumerate(widths):
        table.columns[c].width = Inches(w)

    for i, inc in enumerate(incidentes, start=1):
        cell = table.cell(i, 3)
        color = RAG_COLOR.get(inc["severidad"], LIGHT_BG)
        cell.fill.solid()
        cell.fill.fore_color.rgb = color
        for para in cell.text_frame.paragraphs:
            para.font.color.rgb = WHITE
            para.font.bold = True

    return slide


def slide_riesgos(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, WHITE)
    add_title_bar(slide, "Matriz de riesgos operacionales", "Probabilidad, impacto y plan de mitigacion")

    riesgos = data["riesgos"]
    if not riesgos:
        add_textbox(slide, Inches(0.6), Inches(2), Inches(10), Inches(1), "Sin riesgos registrados esta semana.")
        return slide

    rows = len(riesgos) + 1
    cols = 5
    left, top, width, height = Inches(0.6), Inches(1.55), Inches(12.1), Inches(0.5 * rows)
    table = slide.shapes.add_table(rows, cols, left, top, width, height).table
    headers = ["Riesgo", "Probabilidad", "Impacto", "Nivel", "Plan de mitigacion"]
    for c, h in enumerate(headers):
        table.cell(0, c).text = h

    for i, r in enumerate(riesgos, start=1):
        table.cell(i, 0).text = str(r["riesgo"] or "")
        table.cell(i, 1).text = str(r["probabilidad"] or "")
        table.cell(i, 2).text = str(r["impacto"] or "")
        table.cell(i, 3).text = str(r["nivel"] or "")
        table.cell(i, 4).text = str(r["mitigacion"] or "")

    style_table(table)
    widths = [3.6, 1.8, 1.8, 1.4, 3.5]
    for c, w in enumerate(widths):
        table.columns[c].width = Inches(w)

    for i, r in enumerate(riesgos, start=1):
        cell = table.cell(i, 3)
        color = RAG_COLOR.get(r["nivel"], LIGHT_BG)
        cell.fill.solid()
        cell.fill.fore_color.rgb = color
        for para in cell.text_frame.paragraphs:
            para.font.color.rgb = WHITE
            para.font.bold = True

    return slide


def slide_plan_accion(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, WHITE)
    add_title_bar(slide, "Plan de accion y compromisos", "Seguimiento a los compromisos definidos")

    acciones = data["acciones"]
    if not acciones:
        add_textbox(slide, Inches(0.6), Inches(2), Inches(10), Inches(1), "Sin acciones registradas esta semana.")
        return slide

    rows = len(acciones) + 1
    cols = 5
    left, top, width, height = Inches(0.6), Inches(1.55), Inches(12.1), Inches(0.5 * rows)
    table = slide.shapes.add_table(rows, cols, left, top, width, height).table
    headers = ["Accion", "Responsable", "Fecha compromiso", "Prioridad", "Estado"]
    for c, h in enumerate(headers):
        table.cell(0, c).text = h

    for i, a in enumerate(acciones, start=1):
        table.cell(i, 0).text = str(a["accion"] or "")
        table.cell(i, 1).text = str(a["responsable"] or "")
        table.cell(i, 2).text = fmt_fecha(a["fecha"])
        table.cell(i, 3).text = str(a["prioridad"] or "")
        table.cell(i, 4).text = str(a["estado"] or "")

    style_table(table)
    widths = [4.8, 2.2, 1.8, 1.4, 1.9]
    for c, w in enumerate(widths):
        table.columns[c].width = Inches(w)

    for i, a in enumerate(acciones, start=1):
        cell = table.cell(i, 4)
        color = RAG_COLOR.get(a["estado"], LIGHT_BG)
        cell.fill.solid()
        cell.fill.fore_color.rgb = color
        for para in cell.text_frame.paragraphs:
            para.font.color.rgb = WHITE
            para.font.bold = True

    return slide


def slide_proximos_pasos(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, NAVY)
    add_textbox(slide, Inches(0.8), Inches(0.7), Inches(11), Inches(0.8),
                "Proximos pasos", size=32, bold=True, color=WHITE)

    y = Inches(1.9)
    for item in (data["proximos_pasos"] or ["Sin agenda registrada."]):
        box = slide.shapes.add_textbox(Inches(0.9), y, Inches(11), Inches(0.7))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = f"→  {item}"
        run.font.size = Pt(16)
        run.font.name = FONT
        run.font.color.rgb = WHITE
        y += Inches(0.75)

    return slide


def generar_presentacion(data):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slide_portada(prs, data)
    slide_resumen(prs, data)
    slide_kpis_tabla(prs, data)
    slide_kpis_grafico(prs, data)
    slide_proyectos(prs, data)
    slide_incidentes(prs, data)
    slide_riesgos(prs, data)
    slide_plan_accion(prs, data)
    slide_proximos_pasos(prs, data)

    return prs


def main():
    if len(sys.argv) < 2:
        print("Uso: python scripts/generar_ppt_comite.py <ruta_plantilla.xlsx> [ruta_salida.pptx]")
        sys.exit(1)

    xlsx_path = Path(sys.argv[1])
    if not xlsx_path.exists():
        print(f"No se encontro el archivo: {xlsx_path}")
        sys.exit(1)

    data = cargar_datos(xlsx_path)

    if len(sys.argv) >= 3:
        out_path = Path(sys.argv[2])
    else:
        fecha = data.get("fecha_presentacion")
        if isinstance(fecha, (datetime.date, datetime.datetime)):
            fecha_str = fecha.strftime("%Y-%m-%d")
        else:
            fecha_str = datetime.date.today().strftime("%Y-%m-%d")
        out_dir = Path("salidas")
        out_dir.mkdir(exist_ok=True)
        out_path = out_dir / f"Comite_Operaciones_{fecha_str}.pptx"

    prs = generar_presentacion(data)
    prs.save(out_path)
    print(f"Presentacion generada: {out_path}")


if __name__ == "__main__":
    main()
