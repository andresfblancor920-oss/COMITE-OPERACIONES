"""
Genera el Data Engine comercial actualizado (Kingspan Middle Americas) a partir de:
  - el Data Engine del mes (Kingspan_Commercial_Data_*.xlsx)
  - el Presupuesto de Junta (Presupuesto_Junta.xlsx)

Agrega/ajusta:
  19_OI_Regiones         OI por región (Colombia, Centro América, Caribe, Ecuador, Perú,
                         Venezuela) con selector Mes/YTD y ranking que se ordena solo.
  20_Lost_Business       Negocios perdidos en tabla filtrable + resumen por motivo/competidor.
  21_Ppto_Junta          Presupuesto de Junta mensual, % por canal, ventas reales y metas.
                         La participación por canal puede ser la del presupuesto ANTERIOR
                         (por defecto) o la de la Junta (selector).
  22_Canales_Segmentos   Motor de la nueva estructura comercial:
                           Ventas Alternativas = valor del canal actual.
                           Industrial + Distribución = total IMP − Alternativas, prorrateado.
                           Construcción / Refrigeración = total canales × % manual.
  02_OI / 04_Backlog     Columnas de m² mensuales para las etiquetas de las barras.
  RSM_Resumen_Canales    Budget de Junta por grupo (columnas T:W) y fuente del slide 13.

Uso:
  python scripts/crear_data_engine_comercial.py DATA_ENGINE.xlsx PRESUPUESTO_JUNTA.xlsx SALIDA.xlsx
"""
import sys
from copy import copy

import openpyxl
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.legend import Legend
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY = "061530"
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
SUB_FILL = PatternFill("solid", fgColor="DCE6F1")
TOTAL_FILL = PatternFill("solid", fgColor="E2EFDA")
EST_FILL = PatternFill("solid", fgColor="FCE4D6")
F_TITLE = Font(name="Arial", size=14, bold=True, color=NAVY)
F_SUB = Font(name="Arial", size=9, italic=True, color="595959")
F_HEAD = Font(name="Arial", size=9, bold=True, color="FFFFFF")
F_INPUT = Font(name="Arial", size=9, color="1F4E9A", bold=True)
F_BODY = Font(name="Arial", size=9, color="000000")
F_BOLD = Font(name="Arial", size=9, bold=True, color="000000")
F_SECTION = Font(name="Arial", size=11, bold=True, color=NAVY)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MESES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
NUM = "#,##0;-#,##0;0"
PCT = "0.0%"
PP = '+0.0" pp";-0.0" pp"'


# --------------------------------------------------------------------------- helpers
def put(ws, ref, value, font=F_BODY, fill=None, fmt=None, align=None, border=True):
    c = ws[ref]
    c.value = value
    c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = align
    if border:
        c.border = BORDER
    return c


def inp(ws, ref, value, fmt=None):
    return put(ws, ref, value, font=F_INPUT, fill=INPUT_FILL, fmt=fmt)


def header(ws, row, col0, labels, fill=HEAD_FILL):
    for i, lab in enumerate(labels):
        put(ws, f"{get_column_letter(col0 + i)}{row}", lab, font=F_HEAD, fill=fill,
            align=Alignment(horizontal="center", vertical="center", wrap_text=True))


def title(ws, text, sub):
    put(ws, "A1", text, font=F_TITLE, border=False)
    put(ws, "A2", sub, font=F_SUB, border=False)


def section(ws, ref, text):
    put(ws, ref, text, font=F_SECTION, border=False)


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def bar_chart(title_txt, horizontal=True, w=16, h=7.5):
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.title = title_txt
    ch.style = 10
    ch.width, ch.height = w, h
    ch.legend = None
    return ch


# --------------------------------------------------------------------------- main
def build(src, ppto_path, out):
    wb = openpyxl.load_workbook(src)
    pj = openpyxl.load_workbook(ppto_path, data_only=True)["Hoja1"]

    # ---------------------------------------------------------------- fixes existing
    oi = wb["02_OI"]
    oi["B4"] = 10                       # mes de corte = octubre (el dato de Oct ya está cargado)
    oi["F7"] = 1120898                  # OI YTD PY (antes 112.088: dígito faltante; ver 06 C28)
    oi["F24"] = "=B12"                  # margen OI del mes de corte
    oi["F24"].number_format = "0.0%"
    pf = wb["06_Perf_Forecast"]
    pf["B4"] = "='02_OI'!B4"
    pf["D16"] = '=IF(A16<=$B$4,"Real","Forecast")'
    ov = wb["01_Overview"]
    ov["F6"] = 1120898                  # OI YTD PY a octubre (antes 988.103 = YTD a sept)

    # m² por mes para etiquetas de barras (slide 3 OI)
    put(oi, "I14", "Etiqueta barra 2026: OI m²", font=F_HEAD, fill=HEAD_FILL,
        align=Alignment(wrap_text=True))
    put(oi, "J14", "Etiqueta barra 2025: OI m² (input)", font=F_HEAD, fill=HEAD_FILL,
        align=Alignment(wrap_text=True))
    for i in range(12):
        r = 15 + i
        put(oi, f"I{r}", f"=IF(A{r}<=$B$4,'06_Perf_Forecast'!E{7 + i},\"\")", fmt=NUM)
        inp(oi, f"J{r}", None, fmt=NUM)
    widths(oi, {"I": 16, "J": 16})

    # m² por mes para etiquetas de barras (slide 5 Backlog)
    bk = wb["04_Backlog"]
    bk["F23"] = "=B10"
    bk["F23"].number_format = "0.0%"
    put(bk, "G13", "Etiqueta barra 2026: Backlog m² (input)", font=F_HEAD, fill=HEAD_FILL,
        align=Alignment(wrap_text=True))
    put(bk, "H13", "Etiqueta barra 2025: Backlog m² (input)", font=F_HEAD, fill=HEAD_FILL,
        align=Alignment(wrap_text=True))
    backlog_m2 = {9: 498453, 10: 525529}
    backlog_m2_py = {9: 347834, 10: 360246}
    for i in range(12):
        r = 14 + i
        inp(bk, f"G{r}", backlog_m2.get(i + 1), fmt=NUM)
        inp(bk, f"H{r}", backlog_m2_py.get(i + 1), fmt=NUM)
    widths(bk, {"G": 18, "H": 18})

    build_ppto(wb, pj)
    build_motor(wb)
    build_regiones(wb)
    build_lost(wb)
    extend_rsm(wb)
    update_docs(wb)
    rebuild_charts(wb)

    order = ["LEEME", "00_Resumen_Ejecutivo", "01_Overview", "19_OI_Regiones", "02_OI",
             "03_OI_Canal", "22_Canales_Segmentos", "21_Ppto_Junta", "04_Backlog",
             "05_Backlog_Canal"]
    rest = [ws for ws in wb._sheets if ws.title not in order]
    wb._sheets = [wb[n] for n in order] + rest
    for name in ("19_OI_Regiones", "20_Lost_Business", "21_Ppto_Junta", "22_Canales_Segmentos"):
        wb[name].sheet_properties.tabColor = "00EDB6"
    wb.calculation.fullCalcOnLoad = True
    wb.save(out)
    print("OK ->", out)


# --------------------------------------------------------------------------- 21 Ppto Junta
def build_ppto(wb, pj):
    ws = wb.create_sheet("21_Ppto_Junta")
    title(ws, "21 · Presupuesto de Junta, metas y ventas reales (IMP, m²)",
          "Amarillo = input. Actualice el presupuesto mensual, los % por canal y las ventas reales; "
          "todo lo demás (YTD, asignación por canal y por segmento) es fórmula.")
    put(ws, "A4", "Mes de corte (de 02_OI B4)", font=F_BOLD)
    put(ws, "B4", "='02_OI'!B4", font=F_BOLD)
    put(ws, "A5", "Base de participación por canal", font=F_BOLD)
    inp(ws, "B5", "Anterior")
    dv = DataValidation(type="list", formula1='"Anterior,Junta"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B5")
    put(ws, "C5", "Anterior = % de cada canal en el presupuesto previo (RSM). "
                  "Junta = % mensuales del archivo de Junta (filas 17–21).", font=F_SUB, border=False)

    header(ws, 7, 1, ["Concepto"] + MESES + ["YTD (≤ corte)", "Total año"])
    put(ws, "A8", "N° mes", font=F_BOLD)
    for m in range(12):
        put(ws, f"{get_column_letter(2 + m)}8", m + 1, font=F_BOLD,
            align=Alignment(horizontal="center"))
    ytd = lambda r: f"=SUMIFS(B{r}:M{r},$B$8:$M$8,\"<=\"&$B$4)"

    put(ws, "A9", "Presupuesto Junta IMP (m²)", font=F_BOLD)
    for m in range(12):
        inp(ws, f"{get_column_letter(2 + m)}9", pj.cell(6, 6 + m).value, fmt=NUM)
    put(ws, "N9", ytd(9), font=F_BOLD, fmt=NUM, fill=TOTAL_FILL)
    put(ws, "O9", "=SUM(B9:M9)", font=F_BOLD, fmt=NUM, fill=TOTAL_FILL)

    put(ws, "A10", "Ventas reales IMP (m²) – mensual", font=F_BOLD)
    ventas = {9: 125933, 10: 141231}
    for m in range(12):
        inp(ws, f"{get_column_letter(2 + m)}10", ventas.get(m + 1), fmt=NUM)
    put(ws, "N10", ytd(10), fmt=NUM)
    put(ws, "O10", "=SUM(B10:M10)", fmt=NUM)
    put(ws, "A11", "% cumplimiento mes", font=F_BODY)
    for m in range(12):
        c = get_column_letter(2 + m)
        put(ws, f"{c}11", f'=IF(OR({c}10="",{c}9=0),"",{c}10/{c}9)', fmt=PCT)
    put(ws, "A12", "Ventas YTD reales – cifra oficial (01_Overview E5)", font=F_BOLD)
    put(ws, "N12", "='01_Overview'!E5", font=F_BOLD, fmt=NUM, fill=TOTAL_FILL)
    put(ws, "A13", "% cumplimiento YTD vs Presupuesto Junta", font=F_BOLD)
    put(ws, "N13", "=IFERROR(N12/N9,0)", font=F_BOLD, fmt=PCT, fill=TOTAL_FILL)
    put(ws, "B12", "Si diligencia todos los meses de la fila 10, N10 debe igualar N12.",
        font=F_SUB, border=False)

    header(ws, 15, 1, ["% participación – Presupuesto Junta"] + MESES + ["YTD ponderado", "Año"])
    junta_rows = [("Distribución", 7), ("Industrial", 8), ("Otras Geografías", 9),
                  ("SITES (Site Cali)", 10), ("Ventas Alternativas", 11)]
    for i, (lab, src_row) in enumerate(junta_rows):
        r = 17 + i
        put(ws, f"A{r}", lab, font=F_BOLD)
        for m in range(12):
            inp(ws, f"{get_column_letter(2 + m)}{r}", pj.cell(src_row, 6 + m).value, fmt=PCT)
        put(ws, f"N{r}",
            f"=IFERROR(SUMPRODUCT(B{r}:M{r},$B$9:$M$9,--($B$8:$M$8<=$B$4))/$N$9,0)", fmt=PCT)
        put(ws, f"O{r}", f"=IFERROR(SUMPRODUCT(B{r}:M{r},$B$9:$M$9)/$O$9,0)", fmt=PCT)
    put(ws, "A16", "(los meses deben sumar 100%; ver control fila 22)", font=F_SUB, border=False)
    put(ws, "A22", "Control Σ %", font=F_BOLD)
    for m in range(14):
        c = get_column_letter(2 + m)
        put(ws, f"{c}22", f"=SUM({c}17:{c}21)", fmt=PCT)

    header(ws, 24, 1, ["Canal (estructura anterior)", "Budget anterior YTD (RSM)",
                       "% part. anterior", "% part. Junta YTD", "% part. aplicada"])
    legacy = ["Industrial", "Distribution", "Other Geographies", "Cold Storage", "Site Cali",
              "Alternative Sales"]
    junta_map = {"Industrial": "N18", "Distribution": "N17", "Other Geographies": "N19",
                 "Cold Storage": None, "Site Cali": "N20", "Alternative Sales": "N21"}
    junta_month_row = {"Industrial": 18, "Distribution": 17, "Other Geographies": 19,
                       "Cold Storage": None, "Site Cali": 20, "Alternative Sales": 21}
    for i, lab in enumerate(legacy):
        r = 25 + i
        put(ws, f"A{r}", lab, font=F_BOLD)
        put(ws, f"B{r}", f"=RSM_Resumen_Canales!E{5 + i}", fmt=NUM)
        put(ws, f"C{r}", f"=IFERROR(B{r}/$B$31,0)", fmt=PCT)
        put(ws, f"D{r}", f"={junta_map[lab]}" if junta_map[lab] else 0, fmt=PCT)
        put(ws, f"E{r}", f'=IF($B$5="Junta",D{r},C{r})', fmt=PCT, font=F_BOLD)
    put(ws, "A31", "TOTAL", font=F_BOLD, fill=TOTAL_FILL)
    for c in "BCDE":
        put(ws, f"{c}31", f"=SUM({c}25:{c}30)", font=F_BOLD, fill=TOTAL_FILL,
            fmt=NUM if c == "B" else PCT)

    header(ws, 33, 1, ["Presupuesto Junta por canal anterior (m²)"] + MESES + ["YTD (≤ corte)", "Total año"])
    for i, lab in enumerate(legacy):
        r = 35 + i
        put(ws, f"A{r}", lab, font=F_BOLD)
        jr = junta_month_row[lab]
        for m in range(12):
            c = get_column_letter(2 + m)
            junta_expr = f"{c}$9*{c}{jr}" if jr else "0"
            put(ws, f"{c}{r}", f'=IF($B$5="Junta",{junta_expr},{c}$9*$C{25 + i})', fmt=NUM)
        put(ws, f"N{r}", ytd(r), fmt=NUM, font=F_BOLD)
        put(ws, f"O{r}", f"=SUM(B{r}:M{r})", fmt=NUM)
    put(ws, "A41", "TOTAL", font=F_BOLD, fill=TOTAL_FILL)
    for m in range(14):
        c = get_column_letter(2 + m)
        put(ws, f"{c}41", f"=SUM({c}35:{c}40)", font=F_BOLD, fill=TOTAL_FILL, fmt=NUM)
    put(ws, "A34", "(base = selector B5)", font=F_SUB, border=False)

    header(ws, 43, 1, ["Presupuesto por nueva estructura (m²)"] + MESES + ["YTD (≤ corte)", "Total año"])
    rows = [
        ("Ventas Alternativas", lambda c: f"={c}40"),
        ("Industrial + Distribución (total − Alternativas)", lambda c: f"={c}41-{c}44"),
        ("% Industrial (prorrateo, hoja 22)", lambda c: "='22_Canales_Segmentos'!$F$17"),
        ("Industrial", lambda c: f"={c}45*{c}46"),
        ("Distribución", lambda c: f"={c}45*(1-{c}46)"),
        ("Construcción", lambda c: f"={c}41*'22_Canales_Segmentos'!$F$25"),
        ("Refrigeración", lambda c: f"={c}41*'22_Canales_Segmentos'!$F$26"),
    ]
    for i, (lab, f) in enumerate(rows):
        r = 44 + i
        put(ws, f"A{r}", lab, font=F_BOLD)
        for m in range(12):
            c = get_column_letter(2 + m)
            put(ws, f"{c}{r}", f(c), fmt=PCT if r == 46 else NUM)
        if r == 46:
            put(ws, f"N{r}", "='22_Canales_Segmentos'!$F$17", fmt=PCT)
        else:
            put(ws, f"N{r}", ytd(r), fmt=NUM, font=F_BOLD)
            put(ws, f"O{r}", f"=SUM(B{r}:M{r})", fmt=NUM)
    widths(ws, {"A": 44, **{get_column_letter(c): 10.5 for c in range(2, 14)}, "N": 13, "O": 12})
    ws.freeze_panes = "B9"

    ch = LineChart()
    ch.title = "Presupuesto Junta vs ventas reales (m²)"
    ch.style = 12
    ch.width, ch.height = 22, 7
    data = Reference(ws, min_col=1, max_col=13, min_row=9, max_row=10)
    ch.add_data(data, from_rows=True, titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, max_col=13, min_row=7))
    ws.add_chart(ch, "Q7")


# --------------------------------------------------------------------------- 22 Motor
def build_motor(wb):
    ws = wb.create_sheet("22_Canales_Segmentos")
    title(ws, "22 · Motor de canales (Industrial · Distribución · Ventas Alternativas) y segmentos "
              "(Construcción · Refrigeración)",
          "Regla: Ventas Alternativas = valor del canal actual · Industrial + Distribución = total IMP "
          "− Alternativas, prorrateado con la participación Industrial/(Industrial+Distribución) del canal "
          "actual (o el % manual) · Segmentos = total de los 3 canales × % asignado manualmente.")
    section(ws, "A4", "PARÁMETROS")
    put(ws, "A5", "% Construcción por defecto (Refrigeración = 1 − Construcción)", font=F_BOLD)
    inp(ws, "B5", 0.70, fmt=PCT)
    put(ws, "C5", "Asigne el % real del mes. Se puede ajustar por métrica en la fila 24.",
        font=F_SUB, border=False)

    cols = ["OI mes CM", "OI mes PY", "OI YTD", "Ventas YTD", "Presupuesto YTD (Junta)",
            "Backlog CM", "Backlog PY", "OI YTD (RSM)", "Ventas YTD (RSM)"]
    header(ws, 8, 1, ["Concepto (m²)"] + cols)
    total = ["='02_OI'!B7", "='02_OI'!C7", "='02_OI'!E7", "='01_Overview'!E5", "='21_Ppto_Junta'!N41",
             "='04_Backlog'!B5", "='04_Backlog'!C5", "=RSM_Resumen_Canales!C11",
             "=RSM_Resumen_Canales!D11"]
    alt = ["='03_OI_Canal'!B10", "='03_OI_Canal'!C10", "='03_OI_Canal'!H10", "=RSM_Resumen_Canales!D10",
           "='21_Ppto_Junta'!N40", "='05_Backlog_Canal'!B10", "='05_Backlog_Canal'!C10",
           "=RSM_Resumen_Canales!C10", "=RSM_Resumen_Canales!D10"]
    ind = ["='03_OI_Canal'!B5", "='03_OI_Canal'!C5", "='03_OI_Canal'!H5", "=RSM_Resumen_Canales!D5",
           "='21_Ppto_Junta'!N35", "='05_Backlog_Canal'!B5", "='05_Backlog_Canal'!C5",
           "=RSM_Resumen_Canales!C5", "=RSM_Resumen_Canales!D5"]
    dist = ["='03_OI_Canal'!B6", "='03_OI_Canal'!C6", "='03_OI_Canal'!H6", "=RSM_Resumen_Canales!D6",
            "='21_Ppto_Junta'!N36", "='05_Backlog_Canal'!B6", "='05_Backlog_Canal'!C6",
            "=RSM_Resumen_Canales!C6", "=RSM_Resumen_Canales!D6"]
    src = ["02_OI B7", "02_OI C7", "02_OI E7", "01_Overview E5", "21_Ppto_Junta N41", "04_Backlog B5",
           "04_Backlog C5", "RSM C11", "RSM D11"]
    labels = {
        9: "Fuente del total", 10: "Total IMP (m²)", 11: "Ventas Alternativas – canal actual",
        12: "Industrial + Distribución = Total − Alternativas",
        13: "Base prorrateo: Industrial (canal actual)", 14: "Base prorrateo: Distribución (canal actual)",
        15: "% Industrial – prorrateo automático", 16: "% Industrial – ajuste manual (opcional)",
        17: "% Industrial aplicado", 18: "% Distribución aplicado", 19: "INDUSTRIAL", 20: "DISTRIBUCIÓN",
        21: "VENTAS ALTERNATIVAS", 22: "TOTAL CANALES", 23: "Control vs total IMP (debe ser 0)",
        24: "% Construcción – ajuste manual (opcional)", 25: "% Construcción aplicado",
        26: "% Refrigeración aplicado", 27: "CONSTRUCCIÓN", 28: "REFRIGERACIÓN",
        29: "TOTAL SEGMENTOS",
    }
    for r, lab in labels.items():
        bold = r in (10, 17, 19, 20, 21, 22, 25, 27, 28, 29)
        put(ws, f"A{r}", lab, font=F_BOLD if bold else F_BODY,
            fill=TOTAL_FILL if r in (22, 29) else (SUB_FILL if r in (19, 20, 21, 27, 28) else None))
    for j in range(len(cols)):
        c = get_column_letter(2 + j)
        put(ws, f"{c}9", src[j], font=F_SUB)
        put(ws, f"{c}10", total[j], fmt=NUM, font=F_BOLD)
        put(ws, f"{c}11", alt[j], fmt=NUM)
        put(ws, f"{c}12", f"={c}10-{c}11", fmt=NUM)
        put(ws, f"{c}13", ind[j], fmt=NUM)
        put(ws, f"{c}14", dist[j], fmt=NUM)
        put(ws, f"{c}15", f"=IFERROR({c}13/({c}13+{c}14),0.5)", fmt=PCT)
        inp(ws, f"{c}16", None, fmt=PCT)
        put(ws, f"{c}17", f'=IF({c}16="",{c}15,{c}16)', fmt=PCT, font=F_BOLD)
        put(ws, f"{c}18", f"=1-{c}17", fmt=PCT)
        put(ws, f"{c}19", f"={c}12*{c}17", fmt=NUM, font=F_BOLD, fill=SUB_FILL)
        put(ws, f"{c}20", f"={c}12*{c}18", fmt=NUM, font=F_BOLD, fill=SUB_FILL)
        put(ws, f"{c}21", f"={c}11", fmt=NUM, font=F_BOLD, fill=SUB_FILL)
        put(ws, f"{c}22", f"=SUM({c}19:{c}21)", fmt=NUM, font=F_BOLD, fill=TOTAL_FILL)
        put(ws, f"{c}23", f"=ROUND({c}22-{c}10,0)", fmt=NUM)
        inp(ws, f"{c}24", None, fmt=PCT)
        put(ws, f"{c}25", f'=IF({c}24="",$B$5,{c}24)', fmt=PCT, font=F_BOLD)
        put(ws, f"{c}26", f"=1-{c}25", fmt=PCT)
        put(ws, f"{c}27", f"={c}22*{c}25", fmt=NUM, font=F_BOLD, fill=SUB_FILL)
        put(ws, f"{c}28", f"={c}22*{c}26", fmt=NUM, font=F_BOLD, fill=SUB_FILL)
        put(ws, f"{c}29", f"={c}27+{c}28", fmt=NUM, font=F_BOLD, fill=TOTAL_FILL)

    # Margins (volume-weighted blend: Industrial/Distribución absorben su parte prorrateada de
    # Otras Geografías, Cold Storage y Site Cali)
    section(ws, "A31", "MARGEN BRUTO (GM %) POR CANAL NUEVO – ponderado por volumen")
    header(ws, 32, 1, ["Canal", "OI GM CM", "OI GM PY", "Backlog GM CM", "Backlog GM PY"])
    specs = [  # (sheet, vol col, gm col, p cell)
        ("'03_OI_Canal'", "B", "E", "$B$17"), ("'03_OI_Canal'", "C", "F", "$C$17"),
        ("'05_Backlog_Canal'", "B", "E", "$G$17"), ("'05_Backlog_Canal'", "C", "F", "$H$17"),
    ]
    for j, (sh, v, g, p) in enumerate(specs):
        c = get_column_letter(2 + j)
        others = f"SUMPRODUCT({sh}!{v}7:{v}9,{sh}!{g}7:{g}9)"
        others_v = f"SUM({sh}!{v}7:{v}9)"
        put(ws, f"{c}33", f'=IFERROR(({sh}!{v}5*{sh}!{g}5+{p}*{others})/({sh}!{v}5+{p}*{others_v}),"")', fmt=PCT)
        put(ws, f"{c}34", f'=IFERROR(({sh}!{v}6*{sh}!{g}6+(1-{p})*{others})/({sh}!{v}6+(1-{p})*{others_v}),"")', fmt=PCT)
        put(ws, f"{c}35", f'=IFERROR({sh}!{g}10,"")', fmt=PCT)
        put(ws, f"{c}36", f'=IFERROR(SUMPRODUCT({sh}!{v}5:{v}10,{sh}!{g}5:{g}10)/SUM({sh}!{v}5:{v}10),"")',
            fmt=PCT, font=F_BOLD, fill=TOTAL_FILL)
        inp(ws, f"{c}37", None, fmt=PCT)
        inp(ws, f"{c}38", None, fmt=PCT)
    for r, lab in zip(range(33, 39), ["Industrial", "Distribución", "Ventas Alternativas",
                                       "TOTAL (ponderado)", "Construcción (input opcional)",
                                       "Refrigeración (input opcional)"]):
        put(ws, f"A{r}", lab, font=F_BOLD)
    put(ws, "A39", "Si el GM por segmento queda en blanco, el slide muestra '—'.", font=F_SUB, border=False)

    # ---- Slide tables
    def yoy(a, b):
        return f'=IF(OR({a}="",{b}=""),"",IF({b}=0,"n.a.",{a}/{b}-1))'

    def pp(a, b):
        return f'=IF(OR({a}="",{b}=""),"",({a}-{b})*100)'

    section(ws, "A41", "TABLA SLIDE · ORDER INTAKE BY CHANNEL – IMP")
    header(ws, 42, 1, ["Canal", "OI CM", "OI PY", "YoY", "GM CM", "GM PY", "Var GM (pp)", "OI YTD",
                       "Mix YTD", "Ventas YTD", "Budget YTD", "% Budget"])
    chan = [("Industrial", 19, 33), ("Distribución", 20, 34), ("Ventas Alternativas", 21, 35),
            ("TOTAL", 22, 36)]
    for i, (lab, mr, gr) in enumerate(chan):
        r = 43 + i
        fill = TOTAL_FILL if lab == "TOTAL" else None
        put(ws, f"A{r}", lab, font=F_BOLD, fill=fill)
        vals = [f"=B{mr}", f"=C{mr}", yoy(f"B{r}", f"C{r}"), f"=B{gr}", f"=C{gr}", pp(f"E{r}", f"F{r}"),
                f"=D{mr}", f"=IFERROR(H{r}/$H$46,0)", f"=E{mr}", f"=F{mr}", f"=IFERROR(J{r}/K{r},0)"]
        fmts = [NUM, NUM, PCT, PCT, PCT, "0.0", NUM, "0%", NUM, NUM, "0%"]
        for j, (v, fm) in enumerate(zip(vals, fmts)):
            put(ws, f"{get_column_letter(2 + j)}{r}", v, fmt=fm, fill=fill,
                font=F_BOLD if fill else F_BODY)
    header(ws, 48, 1, ["Segmento", "OI CM", "OI PY", "YoY", "GM CM", "GM PY", "Var GM (pp)", "OI YTD",
                       "Mix YTD", "Ventas YTD", "Budget YTD", "% Budget"])
    seg = [("Construcción", 27, 37), ("Refrigeración", 28, 38), ("TOTAL", 29, 36)]
    for i, (lab, mr, gr) in enumerate(seg):
        r = 49 + i
        fill = TOTAL_FILL if lab == "TOTAL" else None
        put(ws, f"A{r}", lab, font=F_BOLD, fill=fill)
        vals = [f"=B{mr}", f"=C{mr}", yoy(f"B{r}", f"C{r}"), f'=IF(B{gr}="","",B{gr})',
                f'=IF(C{gr}="","",C{gr})', pp(f"E{r}", f"F{r}"), f"=D{mr}", f"=IFERROR(H{r}/$H$51,0)",
                f"=E{mr}", f"=F{mr}", f"=IFERROR(J{r}/K{r},0)"]
        fmts = [NUM, NUM, PCT, PCT, PCT, "0.0", NUM, "0%", NUM, NUM, "0%"]
        for j, (v, fm) in enumerate(zip(vals, fmts)):
            put(ws, f"{get_column_letter(2 + j)}{r}", v, fmt=fm, fill=fill,
                font=F_BOLD if fill else F_BODY)

    section(ws, "A53", "TABLA SLIDE · BACKLOG BY CHANNEL – IMP")
    header(ws, 54, 1, ["Canal", "Backlog CM", "Backlog PY", "YoY", "GM CM", "GM PY", "Δ GM (pp)", "Mix %"])
    for i, (lab, mr, gr) in enumerate(chan):
        r = 55 + i
        fill = TOTAL_FILL if lab == "TOTAL" else None
        put(ws, f"A{r}", lab, font=F_BOLD, fill=fill)
        vals = [f"=G{mr}", f"=H{mr}", yoy(f"B{r}", f"C{r}"), f"=D{gr}", f"=E{gr}", pp(f"E{r}", f"F{r}"),
                f"=IFERROR(B{r}/$B$58,0)"]
        for j, (v, fm) in enumerate(zip(vals, [NUM, NUM, PCT, PCT, PCT, "0.0", "0%"])):
            put(ws, f"{get_column_letter(2 + j)}{r}", v, fmt=fm, fill=fill,
                font=F_BOLD if fill else F_BODY)
    header(ws, 60, 1, ["Segmento", "Backlog CM", "Backlog PY", "YoY", "GM CM", "GM PY", "Δ GM (pp)", "Mix %"])
    for i, (lab, mr, gr) in enumerate(seg):
        r = 61 + i
        fill = TOTAL_FILL if lab == "TOTAL" else None
        put(ws, f"A{r}", lab, font=F_BOLD, fill=fill)
        gcm = f'=IF(D{gr}="","",D{gr})' if lab != "TOTAL" else f"=D{gr}"
        gpy = f'=IF(E{gr}="","",E{gr})' if lab != "TOTAL" else f"=E{gr}"
        vals = [f"=G{mr}", f"=H{mr}", yoy(f"B{r}", f"C{r}"), gcm, gpy, pp(f"E{r}", f"F{r}"),
                f"=IFERROR(B{r}/$B$63,0)"]
        for j, (v, fm) in enumerate(zip(vals, [NUM, NUM, PCT, PCT, PCT, "0.0", "0%"])):
            put(ws, f"{get_column_letter(2 + j)}{r}", v, fmt=fm, fill=fill,
                font=F_BOLD if fill else F_BODY)

    section(ws, "A65", "TABLA SLIDE · RSM PERFORMANCE – CANALES Y SEGMENTOS (YTD, resumen)")
    header(ws, 66, 1, ["Canal", "OI YTD", "Ventas YTD", "Budget YTD (Junta)", "Brecha", "% Budget",
                       "Backlog", "Margen backlog", "Mix backlog"])
    for i, (lab, mr, gr) in enumerate(chan):
        r = 67 + i
        fill = TOTAL_FILL if lab == "TOTAL" else None
        put(ws, f"A{r}", lab, font=F_BOLD, fill=fill)
        vals = [f"=I{mr}", f"=J{mr}", f"=F{mr}", f"=C{r}-D{r}", f"=IFERROR(C{r}/D{r},0)", f"=G{mr}",
                f"=D{gr}", f"=IFERROR(G{r}/$G$70,0)"]
        for j, (v, fm) in enumerate(zip(vals, [NUM, NUM, NUM, NUM, "0%", NUM, PCT, "0%"])):
            put(ws, f"{get_column_letter(2 + j)}{r}", v, fmt=fm, fill=fill,
                font=F_BOLD if fill else F_BODY)
    header(ws, 72, 1, ["Segmento", "OI YTD", "Ventas YTD", "Budget YTD (Junta)", "Brecha", "% Budget",
                       "Backlog", "Margen backlog", "Mix backlog"])
    for i, (lab, mr, gr) in enumerate(seg):
        r = 73 + i
        fill = TOTAL_FILL if lab == "TOTAL" else None
        put(ws, f"A{r}", lab, font=F_BOLD, fill=fill)
        gm = f'=IF(D{gr}="","",D{gr})' if lab != "TOTAL" else f"=D{gr}"
        vals = [f"=I{mr}", f"=J{mr}", f"=F{mr}", f"=C{r}-D{r}", f"=IFERROR(C{r}/D{r},0)", f"=G{mr}",
                gm, f"=IFERROR(G{r}/$G$75,0)"]
        for j, (v, fm) in enumerate(zip(vals, [NUM, NUM, NUM, NUM, "0%", NUM, PCT, "0%"])):
            put(ws, f"{get_column_letter(2 + j)}{r}", v, fmt=fm, fill=fill,
                font=F_BOLD if fill else F_BODY)

    widths(ws, {"A": 46, **{get_column_letter(c): 14 for c in range(2, 13)}})
    ws.freeze_panes = "B9"

    ch = bar_chart("OI YTD por canal (m²)", horizontal=False, w=14, h=7)
    ch.add_data(Reference(ws, min_col=8, min_row=42, max_row=45), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=43, max_row=45))
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showVal = True
    ws.add_chart(ch, "N41")
    ch2 = bar_chart("Ventas vs Budget YTD por canal (m²)", horizontal=False, w=14, h=7)
    ch2.legend = Legend()
    ch2.add_data(Reference(ws, min_col=10, max_col=11, min_row=42, max_row=45), titles_from_data=True)
    ch2.set_categories(Reference(ws, min_col=1, min_row=43, max_row=45))
    ws.add_chart(ch2, "N56")


# --------------------------------------------------------------------------- 19 Regiones
def build_regiones(wb):
    ws = wb.create_sheet("19_OI_Regiones")
    title(ws, "19 · Order Intake por región (slide Commercial Overview)",
          "Diligencie el OI del mes y el YTD por región (amarillo). Seleccione la vista en B4: "
          "el ranking y el gráfico se reordenan solos de mayor a menor volumen.")
    put(ws, "A4", "Vista del gráfico", font=F_BOLD)
    inp(ws, "B4", "Mes")
    dv = DataValidation(type="list", formula1='"Mes,YTD"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B4")

    # Estimación inicial: Colombia = OI total − OI Otras Geografías; Otras Geografías repartidas con
    # el mix de VENTAS por país (01_Overview A19:B22) y "Others" entre Centro América y Perú según el
    # pipeline. Debe reemplazarse con el OI real por país del ERP.
    oi_cm, oi_og = 160547, 55560
    sales = {"Ecuador": 8828, "Caribe": 15144, "Venezuela": 9798, "Others": 6421}
    tot = sum(sales.values())
    est = {k: round(oi_og * v / tot) for k, v in sales.items()}
    ca_share = 204000 / (204000 + 64000)          # pipeline m²: Guatemala+El Salvador+Costa Rica vs Perú
    est["Centro América"] = round(est["Others"] * ca_share)
    est["Perú"] = est.pop("Others") - est["Centro América"]
    est["Colombia"] = oi_cm - oi_og
    regions = ["Colombia", "Centro América", "Caribe", "Ecuador", "Perú", "Venezuela"]
    notes = {
        "Colombia": "Estimado: OI IMP mes − OI Otras Geografías (03_OI_Canal B7)",
        "Centro América": "Estimado: Otras Geo × mix ventas 'Others' × peso pipeline CA",
        "Caribe": "Estimado: Otras Geo × mix ventas Rep. Dominicana",
        "Ecuador": "Estimado: Otras Geo × mix ventas Ecuador",
        "Perú": "Estimado: Otras Geo × mix ventas 'Others' × peso pipeline Perú",
        "Venezuela": "Estimado: Otras Geo × mix ventas Venezuela",
    }
    header(ws, 6, 1, ["Región", "OI mes (m²)", "OI YTD (m²)", "Fuente / estado del dato", "Valor vista",
                      "Clave orden (aux)"])
    for i, reg in enumerate(regions):
        r = 7 + i
        put(ws, f"A{r}", reg, font=F_BOLD)
        inp(ws, f"B{r}", est[reg], fmt=NUM)
        inp(ws, f"C{r}", None, fmt=NUM)
        put(ws, f"D{r}", notes[reg], font=F_SUB, fill=EST_FILL)
        put(ws, f"E{r}", f'=IF($B$4="YTD",C{r},B{r})', fmt=NUM)
        put(ws, f"F{r}", f"=N(E{r})+(100-ROW())/1000000000")
    put(ws, "A13", "TOTAL", font=F_BOLD, fill=TOTAL_FILL)
    for c in "BCE":
        put(ws, f"{c}13", f"=SUM({c}7:{c}12)", font=F_BOLD, fill=TOTAL_FILL, fmt=NUM)
    put(ws, "A14", "Control: Σ regiones mes − OI IMP mes (02_OI B7)", font=F_BODY)
    put(ws, "B14", "=B13-'02_OI'!B7", fmt=NUM)
    put(ws, "A15", "Control: Σ regiones YTD − OI IMP YTD (02_OI E7)", font=F_BODY)
    put(ws, "C15", "=IF(C13=0,\"Pendiente\",C13-'02_OI'!E7)", fmt=NUM)

    section(ws, "A17", "RANKING – fuente del gráfico (se ordena solo)")
    header(ws, 18, 1, ["#", "Región", "m²", "% del total"])
    for k in range(6):
        r = 19 + k
        put(ws, f"A{r}", k + 1, font=F_BOLD)
        m = f"MATCH(LARGE($F$7:$F$12,A{r}),$F$7:$F$12,0)"
        put(ws, f"B{r}", f"=INDEX($A$7:$A$12,{m})", font=F_BOLD)
        put(ws, f"C{r}", f"=INDEX($E$7:$E$12,{m})", fmt=NUM)
        put(ws, f"D{r}", f"=IFERROR(C{r}/$E$13,0)", fmt=PCT)
    put(ws, "A26", "Las celdas en salmón son estimaciones iniciales: reemplácelas con el OI real por país.",
        font=F_SUB, border=False)
    widths(ws, {"A": 46, "B": 16, "C": 14, "D": 58, "E": 13, "F": 12})

    ch = bar_chart("OI por región (m²) – vista seleccionada", horizontal=True, w=16, h=8)
    ch.add_data(Reference(ws, min_col=3, min_row=18, max_row=24), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=2, min_row=19, max_row=24))
    ch.y_axis.scaling.orientation = "minMax"
    ch.x_axis.scaling.orientation = "maxMin"   # mayor volumen arriba
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showVal = True
    ws.add_chart(ch, "H4")


# --------------------------------------------------------------------------- 20 Lost Business
LOST = [
    ("EC BUILDING", "Bodega Mz 100, Zona Franca Zofia", "Kingroof / Kingwall", 5000, 158110.27,
     "Termopainel", "Precio", "Price was the reported reason for losing the opportunity.",
     "Strengthen the commercial relationship with the company's CEO to support future opportunities.",
     "Price", "Build CEO-level relationship"),
    ("Industrias Refridcol", "Proyecto Frigorífico Yumbo", "Kingfrigo 125 mm", 2127, 131198.77, "Metecno",
     "Capacidad de producción",
     "Although the purchase order had been received, production could not proceed because the required "
     "14.28 m panel length exceeded current manufacturing capabilities.",
     "Follow up on plant upgrades and assess future feasibility for special-length panel requirements.",
     "Panel length > plant capability", "Assess special-length capability"),
    ("Interfrigo", "—", "Kingfrigo 80", 2000, 75892.93, "Termopainel", "Precio",
     "Price was the reported reason for losing the opportunity.",
     "Strengthen customer follow-up and encourage timely, clear feedback to identify objections earlier.",
     "Price", "Earlier follow-up and feedback"),
    ("Centurión", "Mezzanine y Oficinas C5", "Kingfrigo 60", 658.75, 20456.07, "Not specified",
     "Presupuesto del cliente", "The opportunity was not secured due to budget constraints.",
     "Not specified in the original report; review budget alignment and explore alternatives with the customer.",
     "Customer budget", "Align budget, offer alternatives"),
    ("Andexco Colombia", "Cubierta Finca Lorica", "Kingroof PIR30", 372, 17439.48, "UPVC sheets",
     "Tecnología alternativa",
     "The customer selected an alternative UPVC roofing technology instead of the proposed solution.",
     "Provide product and value-proposition training to the distributor to strengthen technical and "
     "commercial selling.", "Alternative tech (UPVC)", "Distributor value training"),
    ("Bold", "Cubierta Color Madera James", "Kingroof PIR30", 264, 10149.10, "UPVC sheets",
     "Tecnología alternativa",
     "The customer selected an alternative UPVC roofing technology instead of the proposed solution.",
     "Provide product and value-proposition training to the distributor to strengthen technical and "
     "commercial selling.", "Alternative tech (UPVC)", "Distributor value training"),
    ("Andexco Colombia", "Finca Ganadera Piñón Magdalena", "Kingfrigo PIR50", 154.7, 6479.70, "Metecno",
     "Plazo de entrega",
     "The customer required immediate material availability and could not wait for the proposed delivery timeline.",
     "Confirm the customer's required delivery date early and assess availability before committing to a lead time.",
     "Lead time", "Confirm date & stock upfront"),
]
LOST_REASONS = ["Precio", "Capacidad de producción", "Presupuesto del cliente", "Tecnología alternativa",
                "Plazo de entrega"]
LOST_COMPETITORS = ["Termopainel", "Metecno", "UPVC sheets", "Not specified"]


def build_lost(wb):
    ws = wb.create_sheet("20_Lost_Business")
    title(ws, "20 · Lost Business – negocios perdidos del mes",
          "Una fila por oportunidad perdida (tabla con filtros: use las flechas del encabezado). "
          "El resumen por motivo y competidor y los KPIs del slide se calculan solos.")
    cols = ["Cliente", "Proyecto", "Producto", "m² perdidos", "Valor perdido (USD)", "Competidor",
            "Motivo (categoría)", "Motivo (detalle)", "Acción correctiva / aprendizaje",
            "Motivo – resumen slide", "Acción – resumen slide"]
    header(ws, 4, 1, cols)
    n_rows = 30
    for i in range(n_rows):
        r = 5 + i
        row = LOST[i] if i < len(LOST) else (None,) * 11
        for j, v in enumerate(row):
            c = f"{get_column_letter(1 + j)}{r}"
            inp(ws, c, v, fmt=NUM if j == 3 else ("$#,##0" if j == 4 else None))
            ws[c].alignment = Alignment(wrap_text=j in (7, 8), vertical="top")
    last = 4 + n_rows
    tbl = Table(displayName="tblLostBusiness", ref=f"A4:K{last}")
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(tbl)
    dv = DataValidation(type="list", formula1='"' + ",".join(LOST_REASONS) + ',Otro"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G5:G{last}")

    rng = lambda col: f"${col}$5:${col}${last}"
    section(ws, "M4", "KPIs DEL SLIDE")
    kpis = [
        ("Total m² perdidos", f"=SUM({rng('D')})", NUM),
        ("Valor perdido (USD)", f"=SUM({rng('E')})", "$#,##0"),
        ("N° oportunidades", f"=COUNTA({rng('A')})", "0"),
        ("Ticket promedio (USD)", "=IFERROR(N6/N7,0)", "$#,##0"),
        ("Precio medio perdido (USD/m²)", "=IFERROR(N6/N5,0)", "$#,##0.0"),
        ("Motivo principal (por valor)", "=INDEX($M$15:$M$19,MATCH(MAX($O$15:$O$19),$O$15:$O$19,0))", None),
        ("% del valor – motivo principal", "=IFERROR(MAX($O$15:$O$19)/N6,0)", "0%"),
        ("Competidor principal (por valor)", "=INDEX($M$24:$M$27,MATCH(MAX($O$24:$O$27),$O$24:$O$27,0))", None),
    ]
    for i, (lab, f, fm) in enumerate(kpis):
        put(ws, f"M{5 + i}", lab, font=F_BOLD)
        put(ws, f"N{5 + i}", f, fmt=fm, font=F_BOLD, fill=TOTAL_FILL)

    header(ws, 14, 13, ["Motivo", "m²", "Valor USD", "N°", "% valor"])
    for i, reason in enumerate(LOST_REASONS):
        r = 15 + i
        put(ws, f"M{r}", reason, font=F_BOLD)
        put(ws, f"N{r}", f"=SUMIF({rng('G')},M{r},{rng('D')})", fmt=NUM)
        put(ws, f"O{r}", f"=SUMIF({rng('G')},M{r},{rng('E')})", fmt="$#,##0")
        put(ws, f"P{r}", f"=COUNTIF({rng('G')},M{r})", fmt="0")
        put(ws, f"Q{r}", f"=IFERROR(O{r}/$N$6,0)", fmt="0%")
    header(ws, 23, 13, ["Competidor", "m²", "Valor USD", "N°", "% valor"])
    for i, comp in enumerate(LOST_COMPETITORS):
        r = 24 + i
        put(ws, f"M{r}", comp, font=F_BOLD)
        put(ws, f"N{r}", f"=SUMIF({rng('F')},M{r},{rng('D')})", fmt=NUM)
        put(ws, f"O{r}", f"=SUMIF({rng('F')},M{r},{rng('E')})", fmt="$#,##0")
        put(ws, f"P{r}", f"=COUNTIF({rng('F')},M{r})", fmt="0")
        put(ws, f"Q{r}", f"=IFERROR(O{r}/$N$6,0)", fmt="0%")

    section(ws, "M30", "CONSULTA INTERACTIVA")
    put(ws, "M31", "Competidor (seleccione)", font=F_BOLD)
    inp(ws, "N31", "Todos")
    dv2 = DataValidation(type="list", formula1='"Todos,' + ",".join(LOST_COMPETITORS) + '"')
    ws.add_data_validation(dv2)
    dv2.add("N31")
    put(ws, "M32", "Motivo (seleccione)", font=F_BOLD)
    inp(ws, "N32", "Todos")
    dv3 = DataValidation(type="list", formula1='"Todos,' + ",".join(LOST_REASONS) + '"')
    ws.add_data_validation(dv3)
    dv3.add("N32")
    crit = (f'{rng("F")},IF($N$31="Todos","*",$N$31),{rng("G")},IF($N$32="Todos","*",$N$32)')
    put(ws, "M33", "m² perdidos (filtro)", font=F_BOLD)
    put(ws, "N33", f"=SUMIFS({rng('D')},{crit})", fmt=NUM, fill=TOTAL_FILL)
    put(ws, "M34", "Valor USD (filtro)", font=F_BOLD)
    put(ws, "N34", f"=SUMIFS({rng('E')},{crit})", fmt="$#,##0", fill=TOTAL_FILL)
    put(ws, "M35", "N° oportunidades (filtro)", font=F_BOLD)
    put(ws, "N35", f"=COUNTIFS({crit})", fmt="0", fill=TOTAL_FILL)

    widths(ws, {"A": 20, "B": 28, "C": 20, "D": 11, "E": 14, "F": 14, "G": 22, "H": 46, "I": 46,
                "J": 26, "K": 28, "L": 3, "M": 32, "N": 18, "O": 14, "P": 6, "Q": 9})
    ws.freeze_panes = "B5"
    ch = bar_chart("Valor perdido por motivo (USD)", horizontal=True, w=14, h=6.5)
    ch.add_data(Reference(ws, min_col=15, min_row=14, max_row=19), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=13, min_row=15, max_row=19))
    ch.x_axis.scaling.orientation = "maxMin"
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showVal = True
    ws.add_chart(ch, "S4")


# --------------------------------------------------------------------------- RSM summary
def extend_rsm(wb):
    ws = wb["RSM_Resumen_Canales"]
    ws["E4"] = "Budget anterior YTD (m²)"
    header(ws, 4, 20, ["Budget Junta YTD (m²)", "Brecha vs Junta (m²)", "% Cumpl. Junta", "Semáforo Junta"])
    for i in range(6):
        r = 5 + i
        put(ws, f"T{r}", f"='21_Ppto_Junta'!N{35 + i}", fmt=NUM)
        put(ws, f"U{r}", f"=D{r}-T{r}", fmt=NUM)
        put(ws, f"V{r}", f"=IF(T{r}=0,0,D{r}/T{r})", fmt="0%")
        put(ws, f"W{r}", f'=IF(ROUND(V{r}*100,0)<=60,"Rojo",IF(ROUND(V{r}*100,0)<=84,"Naranja",'
                         f'IF(ROUND(V{r}*100,0)<=100,"Verde claro","Verde oscuro")))')
    put(ws, "T11", "=SUM(T5:T10)", fmt=NUM, font=F_BOLD, fill=TOTAL_FILL)
    put(ws, "U11", "=SUM(U5:U10)", fmt=NUM, font=F_BOLD, fill=TOTAL_FILL)
    put(ws, "V11", "=IF(T11=0,0,D11/T11)", fmt="0%", font=F_BOLD, fill=TOTAL_FILL)
    put(ws, "W11", '=IF(ROUND(V11*100,0)<=60,"Rojo",IF(ROUND(V11*100,0)<=84,"Naranja",'
                   'IF(ROUND(V11*100,0)<=100,"Verde claro","Verde oscuro")))', font=F_BOLD, fill=TOTAL_FILL)
    put(ws, "T12", "Budget Junta = total Junta YTD × % participación del presupuesto anterior "
                   "(o de la Junta, según 21_Ppto_Junta B5).", font=F_SUB, border=False)
    # slide "Sales vs Budget & margin": budget de Junta
    ws["C29"] = "Budget Junta YTD (k m²)"
    ws["G29"] = "% Cumplimiento (Junta)"
    for i in range(7):
        r = 30 + i
        ws[f"C{r}"] = f"=T{5 + i}/1000"
        ws[f"G{r}"] = f"=V{5 + i}"
    ws["C16"] = "=T11"
    ws["C17"] = "=U11"
    ws["C18"] = "=V11"
    ws["C19"] = '=COUNTIF(V5:V10,">=1")&" de 6"'
    ws["A16"] = "Budget Junta YTD (m²)"
    ws["A17"] = "Brecha vs Junta (m²)"
    ws["A18"] = "% Cumplimiento vs Junta"
    ws["A19"] = "Canales ≥100% budget Junta"
    ws["C23"] = "=IFERROR((U5+U7)/U11,0)"
    ws.column_dimensions["T"].width = 16
    ws.column_dimensions["U"].width = 16
    ws.column_dimensions["V"].width = 12
    ws.column_dimensions["W"].width = 14
    res = wb["00_Resumen_Ejecutivo"]
    res["C6"] = "=RSM_Resumen_Canales!T11"


# --------------------------------------------------------------------------- docs
def update_docs(wb):
    ws = wb["LEEME"]
    lines = [
        "Actualización canales y segmentos (Oct-2026)",
        "• Nueva estructura: 3 canales (Industrial, Distribución, Ventas Alternativas) y 2 segmentos "
        "(Construcción, Refrigeración). Motor: hoja 22_Canales_Segmentos.",
        "• Ventas Alternativas = valor del canal actual. Industrial + Distribución = total IMP − Alternativas, "
        "prorrateado con Industrial/(Industrial+Distribución) del canal actual; % manual opcional en fila 16.",
        "• Segmentos = total de los 3 canales × % asignado (22_Canales_Segmentos B5, o por métrica en fila 24).",
        "• Presupuesto: 21_Ppto_Junta (Presupuesto de Junta mensual, % por canal, ventas reales y metas). "
        "B5 elige la participación por canal: Anterior (por defecto) o Junta.",
        "• OI por región (Overview): 19_OI_Regiones, con selector Mes/YTD y ranking automático.",
        "• Lost Business: 20_Lost_Business (tabla con filtros, resumen por motivo/competidor y consulta interactiva).",
        "• Etiquetas m² de las barras (OI y Backlog): 02_OI I:J y 04_Backlog G:H.",
        "• Para regenerar el deck: python scripts/actualizar_deck_comercial.py <este Excel> <deck base> <salida>.",
    ]
    for i, t in enumerate(lines):
        c = ws[f"A{26 + i}"]
        c.value = t
        c.font = copy(ws["A16"].font) if i == 0 else copy(ws["A17"].font)

    v = wb["Validaciones"]
    v["E3"] = '=COUNTIF(E5:E40,"Abierto")'
    new = [
        ("Commercial Overview / OI", "OI YTD PY en 02_OI F7 = 112.088 (falta un dígito) y 01_Overview F6 = 988.103 "
         "(YTD a septiembre).", "Se usa 1.120.898 (06_Perf_Forecast C28): YoY YTD +24,9%.", "Cerrado"),
        ("Commercial Overview", "Ventas YTD PY (01_Overview F5 = 1.120.888) da +7% YoY; el deck muestra +18%.",
         "Confirmar PY de ventas YTD con el ERP.", "Abierto"),
        ("Lost Business", "Total m² perdidos en el slide = 10.437,45; la suma de las filas es 10.576,45.",
         "El slide usa la suma por fórmula (20_Lost_Business).", "Cerrado"),
        ("OI por región", "No hay OI por país en la fuente: Otras Geografías se repartió con el mix de ventas por "
         "país y 'Others' con el pipeline (CA vs Perú).", "Reemplazar 19_OI_Regiones B7:C12 con OI real por país.",
         "Abierto"),
        ("Canales / segmentos", "% Construcción / Refrigeración es un supuesto inicial (70% / 30%).",
         "Asignar el % real en 22_Canales_Segmentos B5 (o por métrica, fila 24).", "Abierto"),
        ("RSM by groups", "Detalle RSM (ventas 1.058.016 m², OI 1.269.124 m²) no cuadra con el IMP a octubre "
         "(ventas 1.199.504, OI 1.400.088): hojas 09–14 parecen cortadas a septiembre.",
         "Actualizar hojas 09–14 a octubre; el budget de Junta usa el mismo mes de corte.", "Abierto"),
    ]
    start = 21
    for i, (sl, hall, trat, est) in enumerate(new):
        r = start + i
        for col, val in zip("ABCDE", (r - 4, sl, hall, trat, est)):
            c = v[f"{col}{r}"]
            c.value = val
            c.font = copy(v["C20"].font)
            c.alignment = copy(v["C20"].alignment)
            c.border = copy(v["C20"].border)


# --------------------------------------------------------------------------- charts
def rebuild_charts(wb):
    """openpyxl no conserva los gráficos existentes: se recrean (uno por hoja)."""
    def combo(ws, title_txt, cats, bars, line, anchor):
        b = BarChart()
        b.type = "col"
        b.title = title_txt
        b.style = 10
        b.width, b.height = 18, 8
        for ref in bars:
            b.add_data(Reference(ws, range_string=ref[0]), titles_from_data=True)
        b.set_categories(Reference(ws, range_string=cats))
        if line:
            ln = LineChart()
            ln.add_data(Reference(ws, range_string=line), titles_from_data=True)
            ln.y_axis.axId = 200
            ln.y_axis.number_format = "0%"
            ln.y_axis.crosses = "max"
            b += ln
        ws.add_chart(b, anchor)

    combo(wb["02_OI"], "OI – Ingresos CY vs PY (COP MM)", "'02_OI'!$B$15:$B$26",
          [("'02_OI'!$G$14:$G$26",), ("'02_OI'!$H$14:$H$26",)], "'02_OI'!$F$14:$F$26", "L4")
    combo(wb["04_Backlog"], "Backlog CY vs PY (COP MM)", "'04_Backlog'!$B$14:$B$25",
          [("'04_Backlog'!$C$13:$C$25",), ("'04_Backlog'!$D$13:$D$25",)], None, "J3")
    combo(wb["06_Perf_Forecast"], "OI 2026 – Real vs Forecast (m²)", "'06_Perf_Forecast'!$B$7:$B$18",
          [("'06_Perf_Forecast'!$I$6:$I$18",), ("'06_Perf_Forecast'!$J$6:$J$18",)], None, "L3")
    combo(wb["RSM_Resumen_Canales"], "Sales vs Budget Junta (k m²) y margen de backlog",
          "RSM_Resumen_Canales!$A$30:$A$35",
          [("RSM_Resumen_Canales!$B$29:$B$35",), ("RSM_Resumen_Canales!$C$29:$C$35",)],
          "RSM_Resumen_Canales!$E$29:$E$35", "I27")

    def hbar(ws, title_txt, cats, vals, anchor, pct=False):
        ch = bar_chart(title_txt, horizontal=True, w=13, h=7.5)
        ch.add_data(Reference(ws, range_string=vals), titles_from_data=True)
        ch.set_categories(Reference(ws, range_string=cats))
        ch.x_axis.scaling.orientation = "maxMin"
        if pct:
            ch.y_axis.number_format = "0%"
        ws.add_chart(ch, anchor)

    hbar(wb["03_OI_Canal"], "OI YTD por canal anterior (m²)", "'03_OI_Canal'!$A$5:$A$10",
         "'03_OI_Canal'!$H$4:$H$10", "A16")
    hbar(wb["05_Backlog_Canal"], "Backlog por canal anterior (m²)", "'05_Backlog_Canal'!$A$5:$A$10",
         "'05_Backlog_Canal'!$B$4:$B$10", "A15")
    hbar(wb["RSM_Resumen_Canales"], "% Cumplimiento budget por canal", "RSM_Resumen_Canales!$Q$5:$Q$10",
         "RSM_Resumen_Canales!$R$4:$R$10", "Y4", pct=True)
    for name, last in [("09_RSM_Industrial", 13), ("10_RSM_Distribucion", 14), ("11_RSM_Otras_Geo", 12),
                       ("12_RSM_Cold_Storage", 7), ("13_RSM_Site_Cali", 6), ("14_RSM_Vtas_Alt", 6)]:
        hbar(wb[name], "% Cumplimiento budget por RSM (detalle interno)", f"'{name}'!$M$5:$M${last}",
             f"'{name}'!$N$4:$N${last}", "L22", pct=True)


if __name__ == "__main__":
    if len(sys.argv) != 4:
        print(__doc__)
        sys.exit(1)
    build(*sys.argv[1:])
