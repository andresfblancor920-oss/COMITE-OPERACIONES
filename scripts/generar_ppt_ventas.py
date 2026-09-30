"""
Genera las 3 diapositivas de Ventas / Revenue / Order Intake del Comité de
Operaciones (vista Customer Excellence) a partir de la plantilla Excel.

    python scripts/generar_ppt_ventas.py plantilla/Plantilla_Ventas_OrderIntake.xlsx [salida.pptx]

Diapositivas:
    1. Volumen de ventas (m²): F'cast vs Budget vs Real, cumplimiento semanal.
    2. Volumen (m²) y Revenue (COP): precio promedio, mix y performance.
    3. Order Intake: negocios ganados vs Budget y proyección de negocios a ganar.

Todos los gráficos son nativos de PowerPoint. Su Excel embebido (clic derecho >
Editar datos) trae celdas amarillas editables y fórmulas, de modo que el
gráfico también puede ajustarse directamente en la presentación sin macros.
Los semáforos, tablas y alertas se recalculan al volver a ejecutar el script.

El script lee solo las celdas de entrada y hace todos los cálculos en Python,
así que no importa si el Excel se guardó sin recalcular fórmulas.
"""

import io
import sys
import datetime
from pathlib import Path

from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# ---------------------------------------------------------------------------
# Estilo
# ---------------------------------------------------------------------------
NAVY = RGBColor(0x0B, 0x25, 0x45)
BLUE = RGBColor(0x1F, 0x6F, 0xB2)
STEEL = RGBColor(0x8F, 0xB3, 0xD9)
GREY = RGBColor(0xC3, 0xC9, 0xD1)
TEXT = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
CARD = RGBColor(0xF3, 0xF5, 0xF8)
LINE = RGBColor(0xDD, 0xE2, 0xE8)
SUBTOTAL = RGBColor(0xE3, 0xEB, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
AMBER = RGBColor(0xC7, 0x8A, 0x00)
RED = RGBColor(0xC6, 0x28, 0x28)
GREEN_BG = RGBColor(0xDC, 0xF0, 0xDD)
AMBER_BG = RGBColor(0xFF, 0xF0, 0xC2)
RED_BG = RGBColor(0xFB, 0xDA, 0xDA)
FONT = "Calibri"

# Colores de gráficos por grupo (mix revenue)
GROUP_COLORS = {
    "Premium": NAVY, "Estándar": BLUE, "Arquitectónica": RGBColor(0x5F, 0x9E, 0xA0),
    "Proyectos": STEEL, "Accesorios": RGBColor(0x9A, 0xA5, 0xB1), "Puertas": RGBColor(0xD9, 0xB3, 0x6C),
    "ICO": RGBColor(0xE0, 0x7A, 0x5F), "Otros": GREY,
}


# ---------------------------------------------------------------------------
# Formatos (convención colombiana: punto de miles)
# ---------------------------------------------------------------------------
def n(v):
    if v is None or v == "":
        return "–"
    s = f"{abs(v):,.0f}".replace(",", ".")
    return f"-{s}" if v < 0 else s


def pct(v):
    return "–" if v is None else f"{v * 100:.0f}%"


def cop(v):
    return "–" if v is None else "$ " + n(v)


def mm(v):
    """Millones de pesos."""
    return "–" if v is None else "$ " + n(v / 1e6) + " M"


def div(a, b):
    return a / b if a is not None and b else None


# ---------------------------------------------------------------------------
# Lectura del Excel
# ---------------------------------------------------------------------------
def num(v):
    if v is None or isinstance(v, str):
        try:
            return float(str(v).replace(".", "").replace(",", ".")) if v not in (None, "") else None
        except ValueError:
            return None
    return float(v)


def cargar(xlsx):
    wb = load_workbook(xlsx, data_only=False)
    d = {}

    ws = wb["Config"]
    cfg = {}
    for r in range(3, 30):
        k = ws.cell(row=r, column=1).value
        if k:
            cfg[str(k).strip()] = ws.cell(row=r, column=2).value
    d["mes"] = cfg.get("Mes", "")
    d["anio"] = cfg.get("Año", "")
    d["corte"] = int(num(cfg.get("Semana de corte (1-5)")) or 4)
    fecha = cfg.get("Fecha del comité")
    d["fecha"] = fecha.strftime("%d/%m/%Y") if isinstance(fecha, (datetime.date, datetime.datetime)) else (fecha or "")
    d["area"] = cfg.get("Área que presenta", "Customer Excellence")
    d["verde"] = num(cfg.get("Umbral verde (≥)")) or 0.95
    d["amarillo"] = num(cfg.get("Umbral amarillo (≥)")) or 0.85
    d["oi_alcance"] = cfg.get("Alcance Order Intake", "Premium")
    d["oi_budget_mes"] = num(cfg.get("Budget Order Intake mes (m²)"))

    ws = wb["Ventas_m2"]
    ventas = []
    for r in range(5, 19):
        linea = ws.cell(row=r, column=1).value
        if not linea:
            continue
        v = [num(ws.cell(row=r, column=c).value) for c in range(3, 15)]
        ventas.append({
            "linea": str(linea), "grupo": ws.cell(row=r, column=2).value or "Otros",
            "fc_ini": v[0], "fc": v[1] or 0, "budget": v[2] or 0,
            "fc_w": [x or 0 for x in v[3:7]], "real_w": [x or 0 for x in v[7:11]], "real": v[11] or 0,
        })
    for x in ventas:
        x["cierre"] = sum(x["real_w"])
    d["ventas"] = ventas

    ws = wb["Revenue"]
    rev = []
    for r in range(5, 21):
        linea = ws.cell(row=r, column=1).value
        if not linea:
            continue
        rev.append({
            "linea": str(linea), "grupo": ws.cell(row=r, column=2).value or "Otros",
            "unidad": ws.cell(row=r, column=3).value or "",
            "cant": num(ws.cell(row=r, column=4).value), "cop": num(ws.cell(row=r, column=5).value) or 0,
            "b_cant": num(ws.cell(row=r, column=6).value), "b_cop": num(ws.cell(row=r, column=7).value),
        })
    d["revenue"] = rev

    ws = wb["Order_Intake"]
    oi = []
    for r in range(5, 10):
        sem = ws.cell(row=r, column=1).value
        b, real = num(ws.cell(row=r, column=2).value), num(ws.cell(row=r, column=3).value)
        if sem and (b or real is not None):
            oi.append({"sem": str(sem), "budget": b or 0, "real": real, "estado": ws.cell(row=r, column=4).value or ""})
    d["oi"] = oi

    ws = wb["Pipeline"]
    pipe = []
    for r in range(5, 30):
        nombre = ws.cell(row=r, column=1).value
        if not nombre:
            continue
        m2, prob = num(ws.cell(row=r, column=4).value) or 0, num(ws.cell(row=r, column=5).value) or 0
        if prob > 1:
            prob /= 100
        pipe.append({"nombre": str(nombre), "cliente": ws.cell(row=r, column=2).value or "",
                     "linea": ws.cell(row=r, column=3).value or "", "m2": m2, "prob": prob,
                     "cierre": str(ws.cell(row=r, column=6).value or "Mes siguiente").strip(),
                     "pond": m2 * prob})
    d["pipeline"] = pipe

    ws = wb["Comentarios"]
    com = {}
    for r in range(5, 30):
        s, t = ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value
        if s and t:
            com.setdefault(str(s).strip(), []).append(str(t))
    d["comentarios"] = com
    return d


def subtotal(ventas, grupos=None, nombre="Total"):
    rows = [v for v in ventas if grupos is None or v["grupo"] in grupos]
    s = {"linea": nombre, "grupo": nombre, "sub": True}
    for k in ("fc_ini", "fc", "budget", "real", "cierre"):
        s[k] = sum((r[k] or 0) for r in rows)
    s["fc_w"] = [sum(r["fc_w"][i] for r in rows) for i in range(4)]
    s["real_w"] = [sum(r["real_w"][i] for r in rows) for i in range(4)]
    return s


# ---------------------------------------------------------------------------
# Helpers de dibujo
# ---------------------------------------------------------------------------
def rag(p, d):
    if p is None:
        return MUTED
    return GREEN if p >= d["verde"] else AMBER if p >= d["amarillo"] else RED


def rag_bg(p, d):
    if p is None:
        return None
    return GREEN_BG if p >= d["verde"] else AMBER_BG if p >= d["amarillo"] else RED_BG


def texto(slide, x, y, w, h, runs, size=11, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP, margin=0):
    """runs: str o lista de párrafos; cada párrafo str o lista de (texto, {opciones})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = o.get("color", color)
    return tb


def caja(slide, x, y, w, h, fill=CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def encabezado(slide, d, titulo, subtitulo, num_slide):
    texto(slide, 0.45, 0.28, 8.8, 0.55, titulo, size=26, bold=True, color=NAVY)
    texto(slide, 0.45, 0.8, 8.8, 0.3, subtitulo, size=12, color=MUTED)
    texto(slide, 9.2, 0.32, 3.7, 0.5,
          [[("Comité de Operaciones", {"bold": True, "color": NAVY, "size": 11})],
           [(f"{d['area']} · {d['fecha']}", {"color": MUTED, "size": 10})]],
          align=PP_ALIGN.RIGHT)
    texto(slide, 0.45, 7.12, 9, 0.25,
          f"Fuente: Plantilla_Ventas_OrderIntake.xlsx · Corte W{d['corte']} · Semáforo: verde ≥ {pct(d['verde'])}, "
          f"amarillo ≥ {pct(d['amarillo'])}, rojo < {pct(d['amarillo'])}", size=9, color=MUTED)
    texto(slide, 12.4, 7.12, 0.5, 0.25, str(num_slide), size=9, color=MUTED, align=PP_ALIGN.RIGHT)


def kpi(slide, x, y, w, h, etiqueta, valor, detalle, p, d):
    caja(slide, x, y, w, h)
    texto(slide, x + 0.18, y + 0.1, w - 0.36, 0.25, etiqueta.upper(), size=9, bold=True, color=MUTED)
    texto(slide, x + 0.18, y + 0.33, w - 0.36, 0.45, valor, size=22, bold=True, color=NAVY)
    if p is not None:
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.18), Inches(y + h - 0.27),
                                     Inches(0.13), Inches(0.13))
        dot.fill.solid()
        dot.fill.fore_color.rgb = rag(p, d)
        dot.line.fill.background()
        texto(slide, x + 0.38, y + h - 0.34, w - 0.5, 0.26,
              [[(pct(p) + " ", {"bold": True, "color": rag(p, d)}), (detalle, {"color": MUTED})]], size=10)
    else:
        texto(slide, x + 0.18, y + h - 0.34, w - 0.36, 0.26, detalle, size=10, color=MUTED)


def alertas(slide, x, y, w, h, items, titulo="Alertas y tendencias", size=10):
    caja(slide, x, y, w, h)
    top = 0.1
    if titulo:
        texto(slide, x + 0.18, y + 0.1, w - 0.36, 0.28, titulo.upper(), size=9, bold=True, color=MUTED)
        top = 0.38
    paras = []
    for color, t in items:
        paras.append([("● ", {"color": color, "bold": True}), (t, {})])
    tb = texto(slide, x + 0.18, y + top, w - 0.3, h - top - 0.07, paras, size=size)
    for p in tb.text_frame.paragraphs:
        p.space_after = Pt(4)


def _borde(cell, color="DDE2E8", w=6350, lados=("B",)):
    tcPr = cell._tc.get_or_add_tcPr()
    for lado in ("L", "R", "T", "B"):
        tag = qn(f"a:ln{lado}")
        for old in tcPr.findall(tag):
            tcPr.remove(old)
    for lado in ("L", "R", "T", "B"):
        ln = etree.SubElement(tcPr, qn(f"a:ln{lado}"), w=str(w if lado in lados else 0))
        if lado in lados:
            sf = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr"), val=color)
        else:
            etree.SubElement(ln, qn("a:noFill"))


def tabla(slide, x, y, w, row_h, filas, anchos, size=10, header_fill=NAVY):
    """filas: lista de listas de celdas; celda = str o dict(t, bold, color, fill, align)."""
    nr, nc = len(filas), len(filas[0])
    gs = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    tbl = gs.table
    # sin estilo de tabla del tema: todo el formato es explícito
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    total = sum(anchos)
    for i, a in enumerate(anchos):
        tbl.columns[i].width = Inches(w * a / total)
    for r in range(nr):
        tbl.rows[r].height = Inches(row_h)
        for c in range(nc):
            spec = filas[r][c]
            spec = spec if isinstance(spec, dict) else {"t": spec}
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            fill = header_fill if r == 0 else spec.get("fill")
            if fill is None:
                cell.fill.background()
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = fill
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = spec.get("align", PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT if r else PP_ALIGN.CENTER)
            run = p.add_run()
            run.text = str(spec.get("t", ""))
            f = run.font
            f.name = FONT
            f.size = Pt(spec.get("size", size - 1 if r == 0 else size))
            f.bold = True if r == 0 else spec.get("bold", False)
            f.color.rgb = WHITE if r == 0 else spec.get("color", TEXT)
            _borde(cell, color="FFFFFF" if r == 0 else "DDE2E8")
    return gs


def celda_pct(p, d, bold=False):
    if p is None:
        return {"t": "–", "color": MUTED, "align": PP_ALIGN.CENTER}
    return {"t": pct(p), "fill": rag_bg(p, d), "color": rag(p, d), "bold": True, "align": PP_ALIGN.CENTER}


# ---------------------------------------------------------------------------
# Gráficos nativos + Excel embebido editable
# ---------------------------------------------------------------------------
def excel_embebido(chart, categorias, series, extras=None, nota=None, fmt="#,##0"):
    """Reemplaza el libro embebido del gráfico por uno con formato de entrada.

    series: lista de (nombre, valores | lista de fórmulas '=...{r}').
    extras: columnas auxiliares a la derecha (nombre, valores/fórmulas, fmt).
    La hoja se llama 'Sheet1' y conserva el layout que usa python-pptx
    (categorías en A2:A.., series en B.., nombres en fila 1), así las
    referencias del gráfico siguen apuntando a las mismas celdas.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    amarillo = PatternFill("solid", fgColor="FFF2CC")
    gris = PatternFill("solid", fgColor="EDEFF2")
    head = PatternFill("solid", fgColor="0B2545")
    cols = list(series) + list(extras or [])
    ws.cell(row=1, column=1, value="").fill = head
    for i, cat in enumerate(categorias):
        ws.cell(row=i + 2, column=1, value=cat).font = Font(bold=True)
    for j, col in enumerate(cols):
        nombre, valores = col[0], col[1]
        f = col[2] if len(col) > 2 else fmt
        h = ws.cell(row=1, column=j + 2, value=nombre)
        h.fill, h.font = head, Font(bold=True, color="FFFFFF")
        for i, v in enumerate(valores):
            r = i + 2
            es_formula = isinstance(v, str) and v.startswith("=")
            c = ws.cell(row=r, column=j + 2, value=v.format(r=r) if es_formula else v)
            c.fill = gris if es_formula else amarillo
            c.number_format = f
        ws.column_dimensions[chr(ord("B") + j)].width = max(12, len(str(nombre)) + 2)
    ws.column_dimensions["A"].width = max(12, max(len(str(c)) for c in categorias) + 2)
    r = len(categorias) + 3
    ws.cell(row=r, column=1, value="Amarillo = dato editable · Gris = fórmula. "
                                   "El gráfico se actualiza al cerrar esta ventana.").font = Font(italic=True, color="666666")
    if nota:
        ws.cell(row=r + 1, column=1, value=nota).font = Font(italic=True, color="666666")
    buf = io.BytesIO()
    wb.save(buf)
    chart.part.chart_workbook.update_from_xlsx_blob(buf.getvalue())


def estilo_chart(chart, titulo, leyenda=True, size=9):
    chart.font.name = FONT
    chart.font.size = Pt(size)
    chart.font.color.rgb = TEXT
    chart.has_title = True
    tf = chart.chart_title.text_frame
    tf.text = titulo
    r = tf.paragraphs[0].runs[0]
    r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(12), True, NAVY, FONT
    chart.has_legend = leyenda
    if leyenda:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(9)


def ejes(chart, fmt="#,##0", ocultar_valores=False):
    va = chart.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = LINE
    va.major_gridlines.format.line.width = Pt(0.5)
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(8)
    va.tick_labels.font.color.rgb = MUTED
    va.tick_labels.number_format = fmt
    va.tick_labels.number_format_is_linked = False
    if ocultar_valores:
        va.visible = False
    ca = chart.category_axis
    ca.format.line.color.rgb = GREY
    ca.tick_labels.font.size = Pt(9)
    ca.tick_labels.font.color.rgb = TEXT
    ca.has_major_gridlines = False


def colorear(plot, colores):
    for s, c in zip(plot.series, colores):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = c


def etiquetas(serie, fmt="#,##0;-#,##0;;", pos=XL_LABEL_POSITION.OUTSIDE_END, size=8, color=TEXT, bold=False):
    dl = serie.data_labels
    dl.show_value = True
    dl.number_format = fmt
    dl.number_format_is_linked = False
    dl.position = pos
    dl.font.size = Pt(size)
    dl.font.color.rgb = color
    dl.font.bold = bold


def barras(slide, x, y, w, h, tipo, categorias, series, titulo, colores, gap=60, overlap=-10,
           etiquetar=None, fmt="#,##0", leyenda=True, ocultar_valores=False):
    cd = CategoryChartData()
    cd.categories = categorias
    for nombre, vals in series:
        cd.add_series(nombre, vals)
    ch = slide.shapes.add_chart(tipo, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    estilo_chart(ch, titulo, leyenda)
    ejes(ch, fmt, ocultar_valores)
    plot = ch.plots[0]
    plot.gap_width = gap
    plot.overlap = overlap
    plot.vary_by_categories = False
    colorear(plot, colores)
    for i in (etiquetar if etiquetar is not None else range(len(series))):
        etiquetas(plot.series[i], f"{fmt};-{fmt};;")  # sin etiqueta para ceros
    return ch


# ---------------------------------------------------------------------------
# Slide 1 — Volumen de ventas (m²)
# ---------------------------------------------------------------------------
def slide_ventas(prs, d):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    v = d["ventas"]
    prem = subtotal(v, {"Premium"}, "Premium")
    est = subtotal(v, {"Estándar"}, "Estándar")
    tot = subtotal(v, None, "Total")
    corte = d["corte"]
    encabezado(s, d, f"Ventas {d['mes']} · Volumen (m²)",
               f"F'cast vs Budget vs Real Sales por línea y cumplimiento semanal · Corte W{corte}", 1)

    # KPIs
    y, h, w, gap = 1.25, 1.05, 2.95, 0.2
    kpi(s, 0.45, y, w, h, "Real Sales Premium", f"{n(prem['real'])} m²",
        f"del Budget ({n(prem['budget'])})", div(prem["real"], prem["budget"]), d)
    kpi(s, 0.45 + (w + gap), y, w, h, "Cierre proyectado Premium", f"{n(prem['cierre'])} m²",
        f"del Budget · F'cast {n(prem['fc'])}", div(prem["cierre"], prem["budget"]), d)
    kpi(s, 0.45 + 2 * (w + gap), y, w, h, "Real Sales Estándar", f"{n(est['real'])} m²",
        f"del Budget ({n(est['budget'])})", div(est["real"], est["budget"]), d)
    kpi(s, 0.45 + 3 * (w + gap), y, w, h, "Gap total vs Budget", f"{n(tot['real'] - tot['budget'])} m²",
        f"Real {n(tot['real'])} de {n(tot['budget'])}", div(tot["real"], tot["budget"]), d)

    # Tabla de cumplimiento (heatmap)
    head = ["Línea", "Budget", "F'cast Act.", "Real Sales", "Cierre proy.", "% Real vs Bgt",
            "W1", "W2", "W3", "W4"]
    filas = [head]

    def fila(x, bold=False, fill=None):
        base = {"bold": bold, "fill": fill}
        cells = [dict(base, t=x["linea"]), dict(base, t=n(x["budget"])), dict(base, t=n(x["fc"])),
                 dict(base, t=n(x["real"])), dict(base, t=n(x["cierre"]))]
        cells.append(celda_pct(div(x["real"], x["budget"]), d))
        for i in range(4):
            c = celda_pct(div(x["real_w"][i], x["fc_w"][i]), d)
            if i + 1 > corte:
                c = {"t": "–", "color": MUTED, "align": PP_ALIGN.CENTER}
            if fill and c.get("fill") is None:
                c["fill"] = fill
            cells.append(c)
        return cells

    for g, sub in (("Premium", prem), ("Estándar", est)):
        for x in [r for r in v if r["grupo"] == g]:
            filas.append(fila(x))
        filas.append(fila(sub, True, SUBTOTAL))
    for x in [r for r in v if r["grupo"] not in ("Premium", "Estándar")]:
        filas.append(fila(x))
    filas.append(fila(tot, True, SUBTOTAL))
    rh = min(0.29, 2.95 / len(filas))
    tabla(s, 0.45, 2.5, 7.9, rh, filas, [1.75, 0.9, 0.95, 0.95, 0.95, 1.0, 0.72, 0.72, 0.72, 0.72], size=9.5)
    texto(s, 0.45, 2.5 + rh * len(filas) + 0.03, 7.9, 0.22,
          "W1–W4 = cumplimiento semanal: Real/Proy de la semana ÷ F'cast semanal.", size=8, color=MUTED)

    # Alertas
    al = [(NAVY, t) for t in d["comentarios"].get("Ventas m2", [])]
    ver, ama = d["verde"], d["amarillo"]
    p_prem = div(prem["cierre"], prem["budget"])
    al.append((rag(p_prem, d), f"Premium proyecta cerrar en {n(prem['cierre'])} m² ({pct(p_prem)} del Budget); "
                               f"brecha de {n(prem['cierre'] - prem['budget'])} m²."))
    rojas = sorted([x for x in v if x["budget"] and div(x["real"], x["budget"]) < ama],
                   key=lambda x: x["real"] - x["budget"])
    for x in rojas[:3]:
        al.append((RED, f"{x['linea']}: Real {n(x['real'])} m² = {pct(div(x['real'], x['budget']))} del Budget "
                        f"(gap {n(x['real'] - x['budget'])} m²)."))
    caidas = [x for x in v if x["fc_ini"] and x["fc"] < 0.9 * x["fc_ini"] and x not in rojas[:3]]
    for x in sorted(caidas, key=lambda x: x["fc"] - x["fc_ini"])[:1]:
        al.append((AMBER, f"{x['linea']}: F'cast ajustado {pct(x['fc'] / x['fc_ini'] - 1)} vs inicial "
                          f"({n(x['fc_ini'])} → {n(x['fc'])} m²)."))
    sem = [(i, div(prem["real_w"][i], prem["fc_w"][i])) for i in range(min(corte, 4))]
    malas = [f"W{i + 1} ({pct(p)})" for i, p in sem if p is not None and p < ver]
    if malas:
        al.append((AMBER, f"Premium bajo el F'cast semanal en {', '.join(malas)}."))
    mejores = [(div(x["real_w"][i], x["fc_w"][i]), x["linea"], i) for x in v for i in range(min(corte, 4))
               if x["fc_w"][i] >= 1000]
    if mejores:
        p, l, i = max(mejores)
        if p >= ver:
            al.append((GREEN, f"Mejor semana: {l} W{i + 1} con {pct(p)} de su F'cast."))
    alertas(s, 0.45, 5.83, 7.9, 1.22, al[:4], size=9.5)

    # Gráficos
    cats = [f"W{i + 1}" for i in range(4)]
    ch = barras(s, 8.6, 2.45, 4.3, 2.3, XL_CHART_TYPE.COLUMN_CLUSTERED, cats,
                [("F'cast semanal", prem["fc_w"]), ("Real/Proy semanal", prem["real_w"])],
                "Premium semanal (m²)", [STEEL, NAVY], gap=55, overlap=-5, ocultar_valores=True)
    excel_embebido(ch, cats, [("F'cast semanal", prem["fc_w"]), ("Real/Proy semanal", prem["real_w"])],
                   extras=[("Cumplimiento", ["=IF(B{r}=0,\"\",C{r}/B{r})"] * 4, "0%")],
                   nota="Premium = Envolventes + Refrigeración.")

    lineas = [x for x in v]
    nombres = [x["linea"] for x in lineas]
    ser = [("Budget", [x["budget"] for x in lineas]), ("F'cast actual", [x["fc"] for x in lineas]),
           ("Real Sales", [x["real"] for x in lineas])]
    ch = barras(s, 8.6, 4.8, 4.3, 2.3, XL_CHART_TYPE.BAR_CLUSTERED, nombres, ser,
                "Budget vs F'cast vs Real (m²)", [GREY, STEEL, NAVY], gap=40, overlap=0,
                etiquetar=[2], ocultar_valores=True)
    ch.category_axis.reverse_order = True
    ch.category_axis.tick_labels.font.size = Pt(8)
    excel_embebido(ch, nombres, ser, extras=[("% Real vs Budget", ["=IF(B{r}=0,\"\",D{r}/B{r})"] * len(nombres), "0%")])
    s.notes_slide.notes_text_frame.text = (
        "Volumen de ventas en m². Real Sales = facturado acumulado; Cierre proyectado = Real de semanas cerradas "
        "+ proyección de la semana en curso. Los gráficos se editan con clic derecho > Editar datos.")


# ---------------------------------------------------------------------------
# Slide 2 — Volumen (m²) y Revenue (COP)
# ---------------------------------------------------------------------------
GRUPOS_ORDEN = ["Premium", "Estándar", "Arquitectónica", "Proyectos", "Accesorios", "Puertas", "Otros", "ICO"]


def slide_revenue(prs, d):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rev = d["revenue"]
    encabezado(s, d, f"Ventas {d['mes']} · m² y Revenue",
               "F'cast actual del mes en m² y valorizado en COP: precio promedio, mix y performance vs Budget", 2)

    grupos = [g for g in GRUPOS_ORDEN if any(r["grupo"] == g for r in rev)] + \
             sorted({r["grupo"] for r in rev} - set(GRUPOS_ORDEN))

    def sub(g):
        rows = [r for r in rev if r["grupo"] == g]
        cant = sum(r["cant"] or 0 for r in rows)
        cant_b = sum(r["cant"] or 0 for r in rows if r["b_cant"])
        b_cant = sum(r["b_cant"] or 0 for r in rows)
        b_cop = sum(r["b_cop"] or 0 for r in rows)
        cop_ = sum(r["cop"] for r in rows)
        p = div(cant_b, b_cant) if b_cant else div(cop_, b_cop)
        return {"linea": f"Subtotal {g}", "cant": cant, "cop": cop_, "b_cant": b_cant or None,
                "b_cop": b_cop or None, "perf": p, "n": len(rows)}

    subs = {g: sub(g) for g in grupos}
    total_prod = sum(r["cop"] for r in rev if r["grupo"] != "ICO")
    total = sum(r["cop"] for r in rev)
    ico = sum(r["cop"] for r in rev if r["grupo"] == "ICO")

    def precio(r):
        return div(r["cop"], r["cant"]) if r.get("cant") else None

    prem, estd = subs.get("Premium"), subs.get("Estándar")
    m2_rows = [r for r in rev if r["unidad"] == "m²"]
    m2_tot = sum(r["cant"] or 0 for r in m2_rows)
    b_m2 = sum(r["b_cant"] or 0 for r in m2_rows)
    m2_con_b = sum(r["cant"] or 0 for r in m2_rows if r["b_cant"])

    y, h, w, gap = 1.25, 1.05, 2.95, 0.2
    kpi(s, 0.45, y, w, h, "Revenue F'cast total", mm(total), f"incluye ICO {mm(ico)}", None, d)
    if prem:
        kpi(s, 0.45 + (w + gap), y, w, h, "Revenue Premium", mm(prem["cop"]),
            f"{pct(div(prem['cop'], total_prod))} del revenue de producto", None, d)
    kpi(s, 0.45 + 2 * (w + gap), y, w, h, "Volumen F'cast (m²)", f"{n(m2_tot)} m²",
        f"del Budget ({n(b_m2)} m²)", div(m2_con_b, b_m2), d)
    if prem and estd:
        kpi(s, 0.45 + 3 * (w + gap), y, w, h, "Precio prom. Premium / m²", cop(precio(prem)),
            f"Estándar: {cop(precio(estd))} / m²", None, d)

    # Tabla
    filas = [["Línea", "Cantidad", "Precio prom.", "Revenue (COP)", "Mix", "Budget", "% Perf."]]
    for g in grupos:
        rows = [r for r in rev if r["grupo"] == g and (r["cop"] or r["cant"])]
        mostrar_sub = len(rows) > 1
        for r in rows:
            unidad = f" {r['unidad']}" if r["unidad"] not in ("m²", "", "—") else ""
            perf = div(r["cant"], r["b_cant"]) if r["b_cant"] else div(r["cop"], r["b_cop"])
            filas.append([r["linea"], (n(r["cant"]) + unidad) if r["cant"] else "–", cop(precio(r)) if precio(r) else "–",
                          cop(r["cop"]), pct(div(r["cop"], total)),
                          n(r["b_cant"]) if r["b_cant"] else (mm(r["b_cop"]) if r["b_cop"] else "–"),
                          celda_pct(perf, d)])
        if mostrar_sub:
            sb = subs[g]
            b = {"bold": True, "fill": SUBTOTAL}
            c = celda_pct(sb["perf"], d)
            c.setdefault("fill", SUBTOTAL)
            if c.get("fill") is None:
                c["fill"] = SUBTOTAL
            filas.append([dict(b, t=g), dict(b, t=n(sb["cant"]) if sb["cant"] else "–"), dict(b, t=cop(precio(sb)) if precio(sb) else "–"),
                          dict(b, t=cop(sb["cop"])), dict(b, t=pct(div(sb["cop"], total))),
                          dict(b, t=n(sb["b_cant"]) if sb["b_cant"] else (mm(sb["b_cop"]) if sb["b_cop"] else "–")), c])
    b = {"bold": True, "fill": NAVY, "color": WHITE}
    filas.append([dict(b, t="TOTAL F'cast (COP)"), dict(b, t=""), dict(b, t=""), dict(b, t=cop(total)),
                  dict(b, t="100%"), dict(b, t=""), dict(b, t="")])
    rh = min(0.26, 3.75 / len(filas))
    tabla(s, 0.45, 2.45, 7.9, rh, filas, [2.2, 1.05, 1.1, 1.55, 0.6, 1.0, 0.75], size=9)
    y_fin = 2.45 + rh * len(filas)

    # Alertas
    al = [(NAVY, t) for t in d["comentarios"].get("Revenue", [])]
    if prem:
        al.append((NAVY, f"Premium concentra {pct(div(prem['cop'], total_prod))} del revenue de producto "
                         f"({mm(prem['cop'])})."))
    bajas = sorted([r for r in rev if (div(r["cant"], r["b_cant"]) if r["b_cant"] else div(r["cop"], r["b_cop"])) is not None],
                   key=lambda r: div(r["cant"], r["b_cant"]) if r["b_cant"] else div(r["cop"], r["b_cop"]))
    rojas = []
    for r in bajas:
        p = div(r["cant"], r["b_cant"]) if r["b_cant"] else div(r["cop"], r["b_cop"])
        if p < d["amarillo"]:
            rojas.append(f"{r['linea']} {pct(p)}")
    if rojas:
        al.append((RED, "Bajo umbral vs Budget: " + ", ".join(rojas) + "."))
    ya = max(y_fin + 0.12, 6.3)
    alertas(s, 0.45, ya, 7.9, 7.05 - ya, al[:max(1, int((7.05 - ya - 0.15) / 0.22))], titulo=None, size=9)

    # Doughnut mix
    gs = [g for g in grupos if subs[g]["cop"] > 0]
    vals = [subs[g]["cop"] / 1e6 for g in gs]
    cd = CategoryChartData()
    cd.categories = gs
    cd.add_series("Revenue (M COP)", vals)
    ch = s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(8.6), Inches(2.4), Inches(4.3), Inches(2.35), cd).chart
    estilo_chart(ch, "Mix de revenue por grupo")
    ch.legend.position = XL_LEGEND_POSITION.RIGHT
    plot = ch.plots[0]
    for i, g in enumerate(gs):
        pt = plot.series[0].points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = GROUP_COLORS.get(g, GREY)
    dl = plot.series[0].data_labels
    dl.show_percentage, dl.show_value = True, False
    dl.number_format, dl.number_format_is_linked = "0%", False
    dl.font.size, dl.font.color.rgb, dl.font.bold = Pt(8), WHITE, True
    # sin etiqueta en porciones pequeñas (< 4%) para que no se encimen
    tot_v = sum(vals) or 1
    ser_dl = plot.series[0]._element.get_or_add_dLbls()
    for i, val in enumerate(vals):
        if val / tot_v < 0.04:
            dl_ = etree.SubElement(ser_dl, qn("c:dLbl"))
            etree.SubElement(dl_, qn("c:idx"), val=str(i))
            etree.SubElement(dl_, qn("c:delete"), val="1")
            ser_dl.remove(dl_)
            ser_dl.insert(0, dl_)
    hole = ch._chartSpace.find(".//" + qn("c:holeSize"))
    if hole is not None:
        hole.set("val", "55")
    excel_embebido(ch, gs, [("Revenue (M COP)", [f"=C{{r}}/1000000"] * len(gs), "#,##0")],
                   extras=[("Revenue (COP)", [subs[g]["cop"] for g in gs], '"$" #,##0')],
                   nota="Edite la columna C (COP); la B la convierte a millones.")

    # Barras revenue por línea
    lr = sorted([r for r in rev if r["cop"] > 0], key=lambda r: -r["cop"])
    nombres = [r["linea"] for r in lr]
    vals = [r["cop"] / 1e6 for r in lr]
    ch = barras(s, 8.6, 4.8, 4.3, 2.3, XL_CHART_TYPE.BAR_CLUSTERED, nombres, [("Revenue (M COP)", vals)],
                "Revenue por línea (millones COP)", [NAVY], gap=35, leyenda=False, ocultar_valores=True)
    ch.category_axis.reverse_order = True
    ch.category_axis.tick_labels.font.size = Pt(8)
    for i, r in enumerate(lr):
        pt = ch.plots[0].series[0].points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = GROUP_COLORS.get(r["grupo"], GREY)
    excel_embebido(ch, nombres, [("Revenue (M COP)", ["=C{r}/1000000"] * len(lr))],
                   extras=[("Revenue (COP)", [r["cop"] for r in lr], '"$" #,##0')],
                   nota="Edite la columna C (COP); la B la convierte a millones.")
    s.notes_slide.notes_text_frame.text = (
        "Revenue = F'cast actual del mes valorizado en COP. El subtotal Estándar incluye Traslúcida en m² y COP; "
        "su % de performance se mide solo sobre las líneas que tienen Budget.")


# ---------------------------------------------------------------------------
# Slide 3 — Order Intake
# ---------------------------------------------------------------------------
def slide_oi(prs, d):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    oi, pipe = d["oi"], d["pipeline"]
    encabezado(s, d, f"Order Intake {d['mes']}",
               f"{d['oi_alcance']} · Negocios ganados (m² adjudicados) vs Budget y proyección por ganar", 3)
    sems = [o["sem"] for o in oi]
    budget = [o["budget"] for o in oi]
    real = [o["real"] or 0 for o in oi]
    pond_sem = [sum(p["pond"] for p in pipe if p["cierre"].upper() == w.upper()) for w in sems]
    pond_mes = sum(pond_sem)
    pond_sig = sum(p["pond"] for p in pipe if p["cierre"].upper() not in {w.upper() for w in sems})
    b_mes = d["oi_budget_mes"] or sum(budget)
    ganado = sum(real)
    proy = ganado + pond_mes

    y, h, w, gap = 1.25, 1.05, 2.95, 0.2
    kpi(s, 0.45, y, w, h, "Negocios ganados", f"{n(ganado)} m²", f"del Budget ({n(b_mes)})", div(ganado, b_mes), d)
    kpi(s, 0.45 + (w + gap), y, w, h, "Gap vs Budget", f"{n(ganado - b_mes)} m²",
        f"acumulado a W{len([r for r in oi if r['real'] is not None])}", None, d)
    kpi(s, 0.45 + 2 * (w + gap), y, w, h, "Cierre proyectado mes", f"{n(proy)} m²",
        f"con pipeline ponderado {n(pond_mes)}", div(proy, b_mes), d)
    kpi(s, 0.45 + 3 * (w + gap), y, w, h, "Pipeline mes siguiente", f"{n(pond_sig)} m²",
        f"ponderado · {n(sum(p['m2'] for p in pipe if p['cierre'].upper() not in {x.upper() for x in sems}))} m² brutos",
        None, d)

    # Gráfico semanal
    ser = [("Budget", budget), ("Ganado", real), ("Por ganar (ponderado)", pond_sem)]
    ch = barras(s, 0.35, 2.45, 6.2, 2.45, XL_CHART_TYPE.COLUMN_CLUSTERED, sems, ser,
                "Order Intake semanal (m²)", [GREY, NAVY, STEEL], gap=55, overlap=-5, etiquetar=[1, 2],
                ocultar_valores=True)
    excel_embebido(ch, sems, ser, extras=[("Cumplimiento", ["=IF(B{r}=0,\"\",C{r}/B{r})"] * len(sems), "0%")],
                   nota="Por ganar = pipeline ponderado (m² × probabilidad) con cierre esperado en esa semana.")

    # Gráfico acumulado (fórmulas sobre columnas auxiliares)
    cats = sems
    acum_b = [sum(budget[:i + 1]) for i in range(len(sems))]
    acum_r = [sum(real[:i + 1]) for i in range(len(sems))]
    acum_p = [acum_r[i] + sum(pond_sem[:i + 1]) for i in range(len(sems))]
    cd = CategoryChartData()
    cd.categories = cats
    for nm, vv in (("Budget acumulado", acum_b), ("Ganado acumulado", acum_r), ("Proyectado (ganado + pipeline)", acum_p)):
        cd.add_series(nm, vv)
    ch = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(6.75), Inches(2.45), Inches(6.15), Inches(2.45), cd).chart
    estilo_chart(ch, "Acumulado del mes vs Budget (m²)")
    ejes(ch, ocultar_valores=True)
    for sr, c, dash in zip(ch.plots[0].series, (GREY, NAVY, BLUE), (False, False, True)):
        sr.smooth = False
        sr.format.line.color.rgb = c
        sr.format.line.width = Pt(2.25)
        if dash:
            sr.format.line.dash_style = 4  # MSO_LINE.DASH
        sr.marker.style = XL_MARKER_STYLE.CIRCLE
        sr.marker.size = 6
        sr.marker.format.fill.solid()
        sr.marker.format.fill.fore_color.rgb = c
        sr.marker.format.line.color.rgb = c
    etiquetas(ch.plots[0].series[1], pos=XL_LABEL_POSITION.BELOW, color=NAVY, bold=True)
    etiquetas(ch.plots[0].series[0], pos=XL_LABEL_POSITION.ABOVE, color=MUTED)
    k = len(cats)
    excel_embebido(ch, cats,
                   [("Budget acumulado", ["=SUM($E$2:E{r})"] * k),
                    ("Ganado acumulado", ["=SUM($F$2:F{r})"] * k),
                    ("Proyectado (ganado + pipeline)", ["=C{r}+SUM($G$2:G{r})"] * k)],
                   extras=[("Budget semanal", budget), ("Ganado semanal", real), ("Por ganar semanal", pond_sem)],
                   nota="Edite las columnas E-G (semanales); las acumuladas B-D se calculan solas.")

    # Tabla semanal
    head = ["Order Intake"] + sems + ["Total"]
    fb = ["Budget"] + [n(x) for x in budget] + [n(b_mes)]
    fr = ["Ganado"]
    for o in oi:
        c = {"t": n(o["real"]) if o["real"] is not None else "–"}
        if str(o["estado"]).lower().startswith("en curso"):
            c.update(fill=SUBTOTAL, t=c["t"] + "*")
        fr.append(c)
    fr.append({"t": n(ganado), "bold": True})
    fg = ["Gap"] + [{"t": n(r - b), "color": GREEN if r >= b else RED} for r, b in zip(real, budget)] + \
         [{"t": n(ganado - b_mes), "color": GREEN if ganado >= b_mes else RED, "bold": True}]
    fc = ["% Cumpl."] + [celda_pct(div(r, b), d) for r, b in zip(real, budget)] + [celda_pct(div(ganado, b_mes), d)]
    tabla(s, 0.45, 5.15, 4.85, 0.3, [head, fb, fr, fg, fc], [1.35] + [1] * len(sems) + [1.1], size=10)
    en_curso = [o["sem"] for o in oi if str(o["estado"]).lower().startswith("en curso")]
    if en_curso:
        texto(s, 0.45, 6.85, 4.85, 0.22, f"* {', '.join(en_curso)} en curso: valor parcial al corte.", size=8, color=MUTED)

    # Pipeline
    top = sorted(pipe, key=lambda p: -p["pond"])[:5]
    filas = [["Negocio por ganar", "Línea", "m²", "Prob.", "Cierre", "Pond."]]
    for p in top:
        filas.append([p["nombre"], p["linea"], n(p["m2"]), {"t": pct(p["prob"]), "align": PP_ALIGN.CENTER},
                      {"t": "Mes sig." if p["cierre"].lower().startswith("mes") else p["cierre"], "align": PP_ALIGN.CENTER},
                      {"t": n(p["pond"]), "bold": True}])
    if not top:
        filas.append(["Sin oportunidades registradas en la hoja Pipeline", "", "", "", "", ""])
    tabla(s, 5.5, 5.15, 4.85, 0.29, filas, [1.95, 1.05, 0.65, 0.55, 0.7, 0.65], size=8.5)

    # Alertas
    al = [(NAVY, t) for t in d["comentarios"].get("Order Intake", [])]
    p_acum = div(ganado, b_mes)
    al.append((rag(p_acum, d), f"Acumulado {pct(p_acum)} del Budget; gap {n(ganado - b_mes)} m²."))
    bajas = [f"{o['sem']} ({pct(div(o['real'], o['budget']))})" for o in oi
             if o["real"] is not None and o["budget"] and div(o["real"], o["budget"]) < d["amarillo"]]
    if bajas:
        al.append((RED, f"Semanas bajo umbral: {', '.join(bajas)}."))
    gap_ = b_mes - ganado
    if gap_ > 0:
        cob = div(pond_mes, gap_)
        al.append((rag(cob, d), f"Pipeline del mes cubre {pct(cob)} del gap ({n(pond_mes)} de {n(gap_)} m²)."))
    alertas(s, 10.55, 5.1, 2.35, 1.95, al[:4], titulo="Alertas", size=9)
    s.notes_slide.notes_text_frame.text = (
        "Order Intake = m² adjudicados (negocios ganados). Proyección = pipeline ponderado por probabilidad "
        "(hoja Pipeline). Los gráficos se editan con clic derecho > Editar datos.")


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    xlsx = Path(sys.argv[1])
    d = cargar(xlsx)
    root = Path(__file__).resolve().parent.parent
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else \
        root / "salidas" / f"Comite_Ventas_{d['mes']}_W{d['corte']}.pptx"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide_ventas(prs, d)
    slide_revenue(prs, d)
    slide_oi(prs, d)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(f"Presentación generada: {out}")


if __name__ == "__main__":
    main()
