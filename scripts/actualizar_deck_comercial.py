"""
Actualiza el deck "Commercial Review" (Kingspan Middle Americas) desde el Data Engine.

  python scripts/actualizar_deck_comercial.py DATA_ENGINE.xlsx DECK_BASE.pptx SALIDA.pptx

Lo que hace (todo se lee del Excel; ver hoja LEEME):
  · Commercial Overview: ventas YTD por canal (Industrial / Distribución / Ventas Alternativas) y
    por segmento (Construcción / Refrigeración) + gráfico de OI por región ordenado por volumen.
  · Order Intake y Backlog del mes: misma estructura de slide, gráfico actualizado y etiqueta de
    m² centrada en cada barra.
  · OI by channel y Backlog by channel: tablas por 3 canales y 2 segmentos.
  · Lost Business: tarjetas KPI, tabla ejecutiva y gráfico de valor perdido por motivo.
  · RSM Performance by groups: budget de Junta por grupo + slide resumen por canales y segmentos.

El script es idempotente: puede correrse sobre el deck original o sobre el deck generado el mes
anterior (las formas que crea llevan el prefijo AUTO_ y se regeneran en cada corrida).

Si el Excel fue generado por script y nunca se abrió en Excel (sin valores calculados), se
recalcula automáticamente con LibreOffice (soffice) si está instalado.
"""
import copy
import os
import re
import shutil
import subprocess
import sys
import tempfile

import openpyxl
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# ----------------------------------------------------------------------------- palette
HEAD = "004990"
ROW = "0B2246"
TOTAL = "043E4C"
TEAL = "00EDB6"
SKY = "19B9F6"
CREAM = "F8F2E7"
MUTED = "A9BCD0"
NEG = "FF570F"
AMBER = "FFB81C"
CARD = "0E2A52"
CARD_LINE = "1F4E8C"
SEMAFORO = [(60, "EE2424"), (84, "FF8A1F"), (100, "7FD36B"), (10 ** 9, "1E9E4A")]

MONTHS_EN = ["January", "February", "March", "April", "May", "June", "July", "August",
             "September", "October", "November", "December"]
REASON_EN = {"Precio": "Price", "Capacidad de producción": "Plant capability",
             "Presupuesto del cliente": "Customer budget", "Tecnología alternativa": "Alternative tech",
             "Plazo de entrega": "Lead time", "Otro": "Other"}
YEAR = 2026


# ----------------------------------------------------------------------------- formatting
def es_int(v):
    if v is None or v == "":
        return "—"
    s = f"{round(v):,}".replace(",", ".")
    return s.replace("-", "−")


def es_dec(v, d=1):
    s = f"{v:,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return s.replace("-", "−")


def es_k(v, d=0):
    return es_dec(v / 1000, d) + "k"


def pct(v, d=1, sign=False):
    if v is None or v == "" or isinstance(v, str):
        return "n.a." if v == "n.a." else "—"
    s = es_dec(v * 100, d) + "%"
    return ("+" + s) if sign and v > 0 else s


def pp(v, d=1):
    if v is None or v == "" or isinstance(v, str):
        return "n.a." if v == "n.a." else "—"
    s = es_dec(v, d) + " pp"
    return ("+" + s) if v > 0 else s


def delta_color(v, base=None):
    if v is None or v == "" or isinstance(v, str):
        return base or "FFFFFF"
    return TEAL if v >= 0 else NEG


def sem_color(ratio):
    p = round(ratio * 100)
    for limit, col in SEMAFORO:
        if p <= limit:
            return col
    return SEMAFORO[-1][1]


# ----------------------------------------------------------------------------- excel
class Engine:
    def __init__(self, path):
        self.wb = openpyxl.load_workbook(path, data_only=True)
        if self.wb["22_Canales_Segmentos"]["B22"].value is None:
            self.wb = self._recalc(path)

    @staticmethod
    def _recalc(path):
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice:
            sys.exit("El Excel no tiene valores calculados: ábralo y guárdelo en Excel, o instale LibreOffice.")
        out = tempfile.mkdtemp()
        subprocess.run([soffice, "--headless", "--convert-to", "xlsx", "--outdir", out, path],
                       check=True, capture_output=True, timeout=300)
        return openpyxl.load_workbook(os.path.join(out, os.path.basename(path)), data_only=True)

    def v(self, sheet, ref):
        return self.wb[sheet][ref].value

    def row(self, sheet, row, c0, c1):
        ws = self.wb[sheet]
        return [ws.cell(row, c).value for c in range(c0, c1 + 1)]


# ----------------------------------------------------------------------------- pptx helpers
def find_slide(prs, pattern, exclude=None):
    for s in prs.slides:
        t = s.shapes.title.text_frame.text if s.shapes.title is not None else ""
        t = t.upper()
        if re.search(pattern, t) and not (exclude and re.search(exclude, t)):
            return s
    return None


def shapes_named(slide, name):
    return [sh for sh in slide.shapes if sh.name == name]


def remove(shape):
    shape._element.getparent().remove(shape._element)


def clear_auto(slide, legacy=()):
    for sh in list(slide.shapes):
        if sh.name.startswith("AUTO_") or sh.name in legacy:
            remove(sh)


def set_runs(shape, runs, keep_first_format=True):
    """Replace the shape text with runs [(text, {color, bold, size})], reusing the first run format."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]._p
    r0 = p0.find(qn("a:r"))
    tmpl = copy.deepcopy(r0.find(qn("a:rPr"))) if (r0 is not None and keep_first_format) else None
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for r in p0.findall(qn("a:r")) + p0.findall(qn("a:br")) + p0.findall(qn("a:fld")):
        p0.remove(r)
    end = p0.find(qn("a:endParaRPr"))
    for text, fmt in runs:
        r = etree.SubElement(p0, qn("a:r"))
        rpr = copy.deepcopy(tmpl) if tmpl is not None else etree.SubElement(r, qn("a:rPr"))
        if tmpl is not None:
            r.append(rpr)
        rpr.attrib.pop("err", None)
        if fmt.get("size"):
            rpr.set("sz", str(int(fmt["size"] * 100)))
        if "bold" in fmt:
            rpr.set("b", "1" if fmt["bold"] else "0")
        if fmt.get("color"):
            for f in rpr.findall(qn("a:solidFill")):
                rpr.remove(f)
            sf = etree.Element(qn("a:solidFill"))
            etree.SubElement(sf, qn("a:srgbClr")).set("val", fmt["color"])
            rpr.insert(0, sf)
        t = etree.SubElement(r, qn("a:t"))
        t.text = text
        if end is not None:
            p0.remove(end)
            p0.append(end)


def set_text(shape, text, **fmt):
    set_runs(shape, [(text, fmt)])


def replace_in_runs(prs, subs):
    for s in prs.slides:
        for sh in s.shapes:
            frames = []
            if sh.has_text_frame:
                frames.append(sh.text_frame)
            if sh.has_table:
                frames += [c.text_frame for r in sh.table.rows for c in r.cells]
            for tf in frames:
                for p in tf.paragraphs:
                    runs = p.runs
                    for r in runs:                      # run level first: keeps per-run colors
                        new = r.text
                        for pat, rep in subs:
                            new = re.sub(pat, rep, new)
                        if new != r.text:
                            r.text = new
                    full = "".join(r.text for r in runs)
                    new = full
                    for pat, rep in subs:
                        new = re.sub(pat, rep, new)
                    if new != full:                     # pattern split across runs
                        runs[0].text = new
                        for r in runs[1:]:
                            r.text = ""


def set_notes(slide, lines, source):
    body = "KEY MESSAGES\n" + "\n".join(f"- {ln}" for ln in lines) + f"\n\nUPDATE\n- Data Engine › {source}"
    tf = slide.notes_slide.notes_text_frame
    if tf is not None:                          # notes master without body placeholder: skip
        tf.text = body


def add_textbox(slide, name, x, y, w, h, runs, size=9, color="FFFFFF", bold=False, align=PP_ALIGN.LEFT,
                anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(runs, str):
        runs = [[(runs, {})]]
    elif runs and isinstance(runs[0], tuple):
        runs = [runs]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for text, fmt in para:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(fmt.get("size", size))
            r.font.bold = fmt.get("bold", bold)
            r.font.color.rgb = RGBColor.from_string(fmt.get("color", color))
        if i < len(runs) - 1:
            p.space_after = Pt(3)
    return tb


def add_box(slide, name, x, y, w, h, fill=CARD, line=CARD_LINE):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.name = name
    sh.adjustments[0] = 0.08
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        sh.line.color.rgb = RGBColor.from_string(line)
        sh.line.width = Pt(0.75)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def add_card(slide, key, x, y, w, h, label, value, sub, sub_color=MUTED):
    add_box(slide, f"AUTO_{key}_card", x, y, w, h)
    add_textbox(slide, f"AUTO_{key}_lbl", x + 0.14, y + 0.08, w - 0.28, 0.16, label, size=7.5, color=CREAM,
                bold=True)
    add_textbox(slide, f"AUTO_{key}_val", x + 0.14, y + 0.27, w - 0.28, 0.32, value, size=16, bold=True)
    add_textbox(slide, f"AUTO_{key}_sub", x + 0.14, y + h - 0.27, w - 0.28, 0.18, sub, size=8, color=sub_color,
                bold=True)


def add_insight(slide, key, x, y, w, h, lines, label="KEY INSIGHTS", size=9.5):
    add_box(slide, f"AUTO_{key}_box", x, y, w, h, fill=TOTAL, line=None)
    add_textbox(slide, f"AUTO_{key}_lbl", x + 0.17, y + 0.1, 2.0, 0.2, label, size=8, color=TEAL, bold=True)
    paras = [[(ln, {"bold": i == 0})] if isinstance(ln, str) else ln for i, ln in enumerate(lines)]
    add_textbox(slide, f"AUTO_{key}_txt", x + 0.17, y + 0.34, w - 0.34, h - 0.42, paras, size=size)


def cell_fmt(cell, text, color="FFFFFF", fill=ROW, bold=False, align=PP_ALIGN.RIGHT, size=9):
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor.from_string(fill)
    cell.margin_left = cell.margin_right = Emu(45720)
    cell.margin_top = cell.margin_bottom = Emu(9144)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)


def add_table(slide, name, x, y, widths, row_h, header, rows, header_fills=None, size=9, header_size=None):
    """rows: list of (cells, kind) where cells = [text | (text, color)] and kind in {'row','total'}."""
    gf = slide.shapes.add_table(len(rows) + 1, len(widths), Inches(x), Inches(y),
                                Inches(sum(widths)), Inches(row_h * (len(rows) + 1)))
    gf.name = name
    tbl = gf.table
    tbl.first_row = True
    tbl.horz_banding = False
    for i, w in enumerate(widths):
        tbl.columns[i].width = Inches(w)
    for r in tbl.rows:
        r.height = Inches(row_h)
    for j, h in enumerate(header):
        fill, color = (header_fills[j] if header_fills else (HEAD, "FFFFFF"))
        cell_fmt(tbl.cell(0, j), h, color=color, fill=fill, bold=True,
                 align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER, size=header_size or size)
    for i, (cells, kind) in enumerate(rows, start=1):
        fill = TOTAL if kind == "total" else ROW
        for j, c in enumerate(cells):
            text, color = (c if isinstance(c, tuple) else (c, "FFFFFF"))
            cell_fmt(tbl.cell(i, j), text, color=color, fill=fill, bold=(j == 0 or kind == "total"),
                     align=PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT, size=size)
    return gf


def old_widths(slide, name):
    sh = shapes_named(slide, name)
    if sh and sh[0].has_table:
        return [c.width / 914400 for c in sh[0].table.columns]
    return None


def style_chart_text(chart, color, size=8):
    tx = chart._chartSpace.find(qn("c:txPr"))
    if tx is None:
        tx = etree.SubElement(chart._chartSpace, qn("c:txPr"))
        etree.SubElement(tx, qn("a:bodyPr"))
        etree.SubElement(tx, qn("a:lstStyle"))
        p = etree.SubElement(tx, qn("a:p"))
        ppr = etree.SubElement(p, qn("a:pPr"))
        d = etree.SubElement(ppr, qn("a:defRPr"))
        d.set("sz", str(int(size * 100)))
        sf = etree.SubElement(d, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", color)
        etree.SubElement(p, qn("a:endParaRPr")).set("lang", "es-CO")
        # c:txPr must precede c:externalData / c:userShapes
        ext = chart._chartSpace.find(qn("c:externalData"))
        if ext is not None:
            ext.addprevious(tx)


def nowrap(data_label):
    data_label.text_frame._txBody.find(qn("a:bodyPr")).set("wrap", "none")


def set_chart_title(chart, text, color, size=9):
    chart.has_title = True
    tf = chart.chart_title.text_frame
    tf.text = text
    r = tf.paragraphs[0].runs[0]
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(color)


def bar_labels(series, labels, color, size=7):
    """Custom centered data label per point (e.g. m² volume on a revenue bar)."""
    for i, text in enumerate(labels):
        dl = series.points[i].data_label
        if not text:
            continue
        dl.has_text_frame = True
        dl.text_frame.text = text
        nowrap(dl)
        run = dl.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(size)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(color)
        dl.position = XL_LABEL_POSITION.CENTER


# ----------------------------------------------------------------------------- slide builders
class DeckUpdater:
    def __init__(self, eng, prs):
        self.e = eng
        self.prs = prs
        self.cut = int(eng.v("02_OI", "B4"))
        self.month = MONTHS_EN[self.cut - 1]
        self.mon = self.month[:3]
        self.cm = f"{self.mon}-{str(YEAR)[2:]}"
        self.py = f"{self.mon}-{str(YEAR - 1)[2:]}"
        m = "22_Canales_Segmentos"
        self.p_oi = eng.v(m, "B17")
        self.p_bk = eng.v(m, "G17")
        self.p_sales = eng.v(m, "E17")
        self.seg_c = eng.v(m, "B25")

    # -- common ------------------------------------------------------------------------
    def method_note(self, extra="", p=None, what="OI CM"):
        p = self.p_oi if p is None else p
        return (f"Channels: Alternative Sales = current channel; Industrial + Distribution = IMP total excl. "
                f"Alternative Sales, prorated by the current Industrial/Distribution mix "
                f"({pct(p, 0)} / {pct(1 - p, 0)} {what}). Segments: Construction "
                f"{pct(self.seg_c, 0)} / Refrigeration {pct(1 - self.seg_c, 0)} (assigned). {extra}").strip()

    def fix_labels(self):
        subs = [(r"OCTUBER", "OCTOBER"), (r"Octuber", "October"), (r"Ovtuber", "October"),
                (r"\boctubre\b", "October"), (r"\bOctubre\b", "October"), (r"Octub-", "Oct-"),
                (r"Octu-", "Oct-"), (r"Jan–Octub\b", "Jan–Oct"), (r"MOM VS A OCT", "MOM VS SEP"),
                (r"MOM VS OCT", "MOM VS SEP"), (r"MoM vs Aug-26", "MoM vs Sep-26"),
                (r"REVENUE SEPT IMP", "REVENUE OCT IMP"), (r"GROSS MARGIN SEPT", "GROSS MARGIN OCT")]
        replace_in_runs(self.prs, subs)

    def renumber(self):
        for i, s in enumerate(self.prs.slides, start=1):
            for sh in shapes_named(s, "Slide number"):
                set_text(sh, str(i))

    # -- 2 · Commercial overview -------------------------------------------------------
    def overview(self):
        s = find_slide(self.prs, r"COMMERCIAL OVERVIEW")
        if s is None:
            return
        e, m = self.e, "22_Canales_Segmentos"
        clear_auto(s)
        charts = [sh for sh in s.shapes if sh.has_chart and sh.name != "AUTO_oi_regiones"]
        charts.sort(key=lambda sh: sh.left)
        ventas = [e.v(m, "E19"), e.v(m, "E20"), e.v(m, "E21")]
        tot = sum(ventas)
        self._donut(charts[0].chart, ["Industrial", "Distribution", "Alt. Sales"],
                    [v / tot for v in ventas], [TEAL, SKY, AMBER], "Sales YTD by channel (m²)")
        seg = [e.v(m, "E27"), e.v(m, "E28")]
        self._donut(charts[1].chart, ["Construction", "Refrigeration"], [v / sum(seg) for v in seg],
                    [TEAL, SKY], "Sales YTD by segment (m²)")
        # OI by region (sorted by volume; source 19_OI_Regiones ranking)
        reg = [(e.v("19_OI_Regiones", f"B{r}"), e.v("19_OI_Regiones", f"C{r}"), e.v("19_OI_Regiones", f"D{r}"))
               for r in range(19, 25)]
        view = e.v("19_OI_Regiones", "B4")
        cd = CategoryChartData()
        cd.categories = [r[0] for r in reg]
        cd.add_series("OI m²", [r[1] or 0 for r in reg])
        x, y, w, h = 7.19, 3.72, 4.87, 1.84
        gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        gf.name = "AUTO_oi_regiones"
        ch = gf.chart
        style_chart_text(ch, CREAM, 8)
        set_chart_title(ch, f"OI by region (m²) · {self.cm if view != 'YTD' else 'YTD ' + str(YEAR)}", CREAM)
        ch.has_legend = False
        ca, va = ch.category_axis, ch.value_axis
        ca.reverse_order = True
        ca.format.line.fill.background()
        ca.tick_labels.font.size = Pt(8)
        ca.tick_labels.font.color.rgb = RGBColor.from_string(CREAM)
        va.visible = False
        va.has_major_gridlines = False
        va.maximum_scale = max(r[1] or 0 for r in reg) * 1.6
        plot = ch.plots[0]
        plot.gap_width = 45
        ser = plot.series[0]
        for i, r in enumerate(reg):
            pt = ser.points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = RGBColor.from_string(TEAL if i < 2 else SKY)
            dl = pt.data_label
            dl.has_text_frame = True
            dl.text_frame.text = f"{es_k(r[1] or 0, 1)} · {pct(r[2], 0)}"
            nowrap(dl)
            run = dl.text_frame.paragraphs[0].runs[0]
            run.font.size = Pt(7.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string(CREAM)
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.startswith("* Unit as reported"):
                set_text(sh, "* Unit as reported in source deck – to be validated. OI by region: Data Engine › "
                             "19_OI_Regiones (estimated split of Other Geographies until ERP data is loaded).")

    def _donut(self, chart, cats, vals, colors, title):
        cd = CategoryChartData()
        cd.categories = cats
        cd.add_series("Mix", vals)
        chart.replace_data(cd)
        ser = chart.plots[0].series[0]
        sx = ser._element
        dlbls = sx.find(qn("c:dLbls"))
        if dlbls is not None:
            for d in dlbls.findall(qn("c:dLbl")):
                dlbls.remove(d)
            nf = dlbls.find(qn("c:numFmt"))
            if nf is not None:
                nf.set("formatCode", "0%")
        for i, col in enumerate(colors):
            pt = ser.points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = RGBColor.from_string(col)
        for i in range(len(cats)):
            dl = ser.points[i].data_label
            dl.has_text_frame = True
            dl.text_frame.text = pct(vals[i], 0)
            nowrap(dl)
            r = dl.text_frame.paragraphs[0].runs[0]
            r.font.size = Pt(8)
            r.font.bold = True
            r.font.color.rgb = RGBColor.from_string(ROW if colors[i] in (TEAL, AMBER) else "FFFFFF")
        set_chart_title(chart, title, CREAM)
        if chart.has_legend:
            chart.legend.font.size = Pt(7.5)
            chart.legend.font.color.rgb = RGBColor.from_string(CREAM)

    # -- 3 / 5 · OI and backlog month --------------------------------------------------
    def _bar_chart(self, chart, sheet, r0, cy, py, mg, lbl_cy, lbl_py, names):
        e = self.e
        cats, s_cy, s_py, s_mg, l_cy, l_py = [], [], [], [], [], []
        for i in range(12):
            r = r0 + i
            cats.append(e.v(sheet, f"B{r}"))
            vcy = e.v(sheet, f"{cy}{r}")
            s_cy.append(vcy if (vcy not in ("", None) and i < self.cut) else None)
            s_py.append(e.v(sheet, f"{py}{r}"))
            vm = e.v(sheet, f"{mg}{r}")
            s_mg.append(vm if (vm not in ("", None) and i < self.cut) else None)
            a, b = e.v(sheet, f"{lbl_cy}{r}"), e.v(sheet, f"{lbl_py}{r}")
            l_cy.append(es_k(a) if isinstance(a, (int, float)) and i < self.cut else "")
            l_py.append(es_k(b) if isinstance(b, (int, float)) else "")
        cd = CategoryChartData()
        cd.categories = cats
        cd.add_series(names[0], s_cy)
        cd.add_series(names[1], s_py)
        cd.add_series(names[2], s_mg)
        chart.replace_data(cd)
        bars = chart.plots[0].series
        bar_labels(bars[0], l_cy, ROW)
        bar_labels(bars[1], l_py, "FFFFFF")

    def oi_month(self):
        s = find_slide(self.prs, r"^ORDER INTAKE\s*[–-]", exclude=r"BY CHANNEL|FORECAST")
        if s is None:
            return
        ch = [sh for sh in s.shapes if sh.has_chart][0]
        self._bar_chart(ch.chart, "02_OI", 15, "C", "D", "F", "I", "J",
                        [f"{YEAR} (COP MM)", f"{YEAR - 1} (COP MM)", f"Margin {YEAR} %"])
        set_text(s.shapes.title, f"ORDER INTAKE – {self.month.upper()} {YEAR}")
        sub = shapes_named(s, "Subtitle")
        if sub:
            set_text(sub[0], "IMP · Current month (CM) vs prior year (PY) and YTD · Revenue in COP MM · "
                             "bar labels = OI volume (k m²)")

    def backlog_month(self):
        s = find_slide(self.prs, r"^BACKLOG\s*[–-]", exclude=r"BY CHANNEL")
        if s is None:
            return
        # same frame as the Order Intake slide (positions in inches)
        def place(sh, x, y, w=None, h=None):
            sh.left, sh.top = Inches(x), Inches(y)
            if w is not None:
                sh.width = Inches(w)
            if h is not None:
                sh.height = Inches(h)

        oi = find_slide(self.prs, r"^ORDER INTAKE\s*[–-]", exclude=r"BY CHANNEL|FORECAST")
        if oi is not None and s.slide_layout.part is not oi.slide_layout.part:
            for rel in s.part.rels.values():
                if rel.reltype.endswith("/slideLayout"):
                    rel._target = oi.slide_layout.part
        place(s.shapes.title, 1.21, 0.6)
        for sh in shapes_named(s, "Subtitle"):
            place(sh, 1.21, 1.03, 10.5, 0.26)
        cards = sorted([sh for sh in s.shapes if sh.name.startswith("KPI ")], key=lambda sh: sh.left)
        texts = [sh for sh in s.shapes if sh.name.startswith("TextBox") and sh.top < Inches(2.75)
                 and sh.has_text_frame]
        xs = [0.85, 3.79, 6.72, 9.66]
        if cards and cards[0].width < Inches(2.75):            # not yet moved
            for card, nx in zip(cards, xs):
                cx0, cx1 = card.left, card.left + card.width
                group = sorted([t for t in texts if cx0 - Inches(0.05) <= t.left < cx1], key=lambda t: t.top)
                place(card, nx, 1.58, 2.79, 0.96)
                for t, ny in zip(group, [1.66, 1.9, 2.25]):
                    place(t, nx + 0.15, ny, 2.49)
        chart = [sh for sh in s.shapes if sh.has_chart][0]
        place(chart, 0.91, 2.68, 11.31, 2.7)
        for sh in s.shapes:
            if sh.has_table:
                place(sh, 1.06, 5.47)
                t = sh.table
                for col, w in zip(t.columns, [1.75, 0.8, 0.8, 0.86, 0.8, 0.96]):
                    col.width = Inches(w)
                for row in t.rows:
                    row.height = Inches(0.245)
                sh.width = Inches(5.97)
            if sh.name == "Insight":
                place(sh, 0.11, 6.64, 11.6, 0.61)
        ins_txt = [sh for sh in s.shapes if sh.has_text_frame and sh.name.startswith("TextBox")
                   and sh.top > Inches(5.5)]
        for sh in ins_txt:
            if sh.text_frame.text.strip().upper() == "KEY INSIGHT":
                place(sh, 0.32, 6.67, 1.3, 0.61)
            else:
                place(sh, 1.31, 6.9, 10.0, 0.18)
        self._bar_chart(chart.chart, "04_Backlog", 14, "C", "D", "F", "G", "H",
                        [f"{YEAR} (COP MM)", f"{YEAR - 1} (COP MM)", f"Margin {YEAR} %"])
        set_text(s.shapes.title, f"BACKLOG – {self.month.upper()} {YEAR}")
        for sh in shapes_named(s, "Subtitle"):
            set_text(sh, "IMP · Current month vs prior year and vs 1 January · Value in COP MM · "
                         "bar labels = backlog volume (k m²)")

    # -- 4 · OI by channel -------------------------------------------------------------
    def _rows(self, sheet, r0, n, spec):
        out = []
        for i in range(n):
            r = r0 + i
            cells = [self.e.v(sheet, f"A{r}")]
            for col, kind in spec:
                v = self.e.v(sheet, f"{col}{r}")
                cells.append(self._cell(v, kind))
            out.append(cells)
        return out

    @staticmethod
    def _cell(v, kind):
        if kind == "int":
            return es_int(v)
        if kind == "k":
            return es_int(v / 1000) if isinstance(v, (int, float)) else "—"
        if kind == "pct":
            return pct(v)
        if kind == "pct0":
            return pct(v, 0)
        if kind == "yoy":
            return (pct(v, 1, sign=True), delta_color(v))
        if kind == "pp":
            return (pp(v), delta_color(v))
        if kind == "budget":
            if not isinstance(v, (int, float)):
                return "—"
            return (pct(v, 0), TEAL if v >= 1 else (AMBER if v >= 0.85 else NEG))
        if kind == "sem":
            return (pct(v, 0), sem_color(v)) if isinstance(v, (int, float)) else "—"
        if kind == "gapk":
            if not isinstance(v, (int, float)):
                return "—"
            return (es_int(v / 1000), "FFFFFF" if round(v / 1000) == 0 else delta_color(v))
        return str(v)

    @staticmethod
    def _label_rows(rows, names):
        for r, n in zip(rows, names):
            r[0] = n
        return rows

    def channel_slide(self, pattern, kind):
        s = find_slide(self.prs, pattern)
        if s is None:
            return
        e, m = self.e, "22_Canales_Segmentos"
        dark = "Dark" in s.slide_layout.name
        note_color = MUTED if dark else "595959"
        label_color = TEAL if dark else HEAD
        legacy = ("Table 5", "Insight", "TextBox 9", "TextBox 10")
        widths = old_widths(s, "Table 5")
        clear_auto(s, legacy)
        chan_names = ["Industrial", "Distribution", "Alternative Sales", "Total"]
        seg_names = ["Construction", "Refrigeration", "Total"]
        if kind == "oi":
            widths = widths if widths and len(widths) == 12 else [1.75, .85, .85, .75, .75, .75, .75, 1.1, .85,
                                                                  1.1, 1.1, 1.0]
            spec = [("B", "int"), ("C", "int"), ("D", "yoy"), ("E", "pct"), ("F", "pct"), ("G", "pp"),
                    ("H", "int"), ("I", "pct0"), ("J", "int"), ("K", "int"), ("L", "budget")]
            hdr = ["OI CM", "OI PY", "YoY", "GM CM", "GM PY", "Var GM", "OI YTD", "Mix YTD", "Sales YTD",
                   "Budget YTD", "% Budget"]
            fills = [(HEAD, "FFFFFF")] * 9 + [(TEAL, ROW)] * 3
            r_ch, r_seg = 43, 49
            subtitle = (f"{self.cm} CM vs PY, YoY and gross margin by channel and segment · YTD volume, sales and "
                        f"Board budget compliance (m²)")
        else:
            widths = widths if widths and len(widths) == 8 else [2.2, 1.35, 1.35, 1.35, 1.35, 1.35, 1.35, 1.3]
            spec = [("B", "int"), ("C", "int"), ("D", "yoy"), ("E", "pct"), ("F", "pct"), ("G", "pp"),
                    ("H", "pct0")]
            hdr = ["Backlog CM", "Backlog PY", "YoY", "GM CM", "GM PY", "Δ GM", "Mix %"]
            fills = None
            r_ch, r_seg = 55, 61
            subtitle = f"{self.cm} backlog vs PY, YoY and gross margin by channel and segment (m²)"
        rows_c = self._label_rows(self._rows(m, r_ch, 4, spec), chan_names)
        rows_s = self._label_rows(self._rows(m, r_seg, 3, spec), seg_names)
        x0, rh, y1 = 0.83, 0.28, 1.55
        add_textbox(s, "AUTO_lbl_ch", x0, y1, 4, 0.2, "BY CHANNEL", size=9, color=label_color, bold=True)
        add_table(s, "AUTO_tbl_channels", x0, y1 + 0.22, widths, rh, ["Channel"] + hdr,
                  [(r, "total" if i == 3 else "row") for i, r in enumerate(rows_c)],
                  header_fills=[(HEAD, "FFFFFF")] + fills if fills else None)
        y2 = y1 + 0.22 + rh * 5 + 0.16
        add_textbox(s, "AUTO_lbl_seg", x0, y2, 4, 0.2, "BY SEGMENT", size=9, color=label_color, bold=True)
        add_table(s, "AUTO_tbl_segments", x0, y2 + 0.22, widths, rh, ["Segment"] + hdr,
                  [(r, "total" if i == 2 else "row") for i, r in enumerate(rows_s)],
                  header_fills=[(HEAD, "FFFFFF")] + fills if fills else None)
        y3 = y2 + 0.22 + rh * 4 + 0.18
        lines = self._oi_insights() if kind == "oi" else self._backlog_insights()
        add_insight(s, "ins", x0, y3, sum(widths), 1.2, lines, size=9)
        set_notes(s, lines + [self.method_note()], "22_Canales_Segmentos (tablas filas 42–63) · 21_Ppto_Junta")
        extra = "Budget = Board budget (Jan–" + self.mon + ")." if kind == "oi" else ""
        add_textbox(s, "AUTO_note", x0, y3 + 1.27, sum(widths) if dark else 10.6, 0.3, self.method_note(extra, *((self.p_oi, "OI CM") if kind == "oi" else (self.p_bk, "backlog CM"))),
                    size=7.5, color=note_color)
        for sh in shapes_named(s, "Subtitle"):
            set_text(sh, subtitle)

    def _oi_insights(self):
        e, m = self.e, "22_Canales_Segmentos"
        tot, py = e.v(m, "B46"), e.v(m, "C46")
        ytd, ytd_py = e.v("02_OI", "E7"), e.v("02_OI", "F7")
        yoy_ytd = ytd / ytd_py - 1 if ytd_py else None
        return [
            f"OI {self.month} {es_int(tot)} m² ({pct(tot / py - 1, 1, True)} vs {self.py}); YTD {es_int(ytd)} m² "
            f"({pct(yoy_ytd, 1, True)} YoY).",
            f"YTD mix: Industrial {pct(e.v(m, 'I43'), 0)} · Distribution {pct(e.v(m, 'I44'), 0)} · "
            f"Alternative Sales {pct(e.v(m, 'I45'), 0)}; Construction {pct(e.v(m, 'I49'), 0)} · Refrigeration "
            f"{pct(e.v(m, 'I50'), 0)}.",
            f"Weighted OI margin {pct(e.v(m, 'E46'))} vs {pct(e.v(m, 'F46'))} PY ({pp(e.v(m, 'G46'))}); "
            f"Alternative Sales leads at {pct(e.v(m, 'E45'))}.",
            f"Sales YTD at {pct(e.v(m, 'L46'), 0)} of the Board budget ({es_int(e.v(m, 'J46'))} vs "
            f"{es_int(e.v(m, 'K46'))} m²): Industrial {pct(e.v(m, 'L43'), 0)} · Distribution "
            f"{pct(e.v(m, 'L44'), 0)} · Alternative Sales {pct(e.v(m, 'L45'), 0)}.",
        ]

    def _backlog_insights(self):
        e, m = self.e, "22_Canales_Segmentos"
        return [
            f"Backlog IMP {es_int(e.v(m, 'B58'))} m² ({pct(e.v(m, 'D58'), 1, True)} vs {es_int(e.v(m, 'C58'))} m² "
            f"PY); Industrial holds {pct(e.v(m, 'H55'), 0)} ({pct(e.v(m, 'D55'), 1, True)}).",
            f"Distribution {es_int(e.v(m, 'B56'))} m² ({pct(e.v(m, 'D56'), 1, True)}) · Alternative Sales "
            f"{es_int(e.v(m, 'B57'))} m² ({pct(e.v(m, 'D57'), 1, True)}).",
            f"Weighted backlog margin {pct(e.v(m, 'E58'))} vs {pct(e.v(m, 'F58'))} PY ({pp(e.v(m, 'G58'))}).",
            f"Segments: Construction {es_int(e.v(m, 'B61'))} m² · Refrigeration {es_int(e.v(m, 'B62'))} m² "
            f"(assigned split {pct(self.e.v(m, 'G25'), 0)} / {pct(self.e.v(m, 'G26'), 0)}).",
        ]

    # -- 9 · Lost business ------------------------------------------------------------
    def lost_business(self):
        s = find_slide(self.prs, r"LOST BUSINESS")
        if s is None:
            return
        e, L = self.e, "20_Lost_Business"
        clear_auto(s, ("Tabla 3", "CuadroTexto 20"))
        set_text(s.shapes.title, f"LOST BUSINESS – {self.month.upper()} {YEAR}")
        rows = []
        for r in range(5, 35):
            if not e.v(L, f"A{r}"):
                continue
            rows.append({k: e.v(L, f"{c}{r}") for k, c in zip(
                ("cli", "proj", "prod", "m2", "usd", "comp", "cat", "det", "acc", "mot_s", "acc_s"), "ABCDEFGHIJK")})
        rows.sort(key=lambda d: -(d["usd"] or 0))
        m2, usd, n = e.v(L, "N5"), e.v(L, "N6"), e.v(L, "N7")
        reason, reason_pct = e.v(L, "N10"), e.v(L, "N11")
        comp = e.v(L, "N12")
        comp_val = max(e.v(L, f"O{r}") or 0 for r in range(24, 28))
        top = rows[0] if rows else None
        add_textbox(s, "AUTO_sub", 1.21, 1.28, 10.5, 0.22,
                    "Opportunities lost in the month · volume (m²), value (USD), competitor, reason and corrective "
                    "action · source: Data Engine › 20_Lost_Business", size=9, color=TEAL)
        x0, gap, w = 0.94, 0.15, (11.45 - 4 * 0.15) / 5
        cards = [
            ("M² LOST", f"{es_int(m2)} m²", f"{n} opportunities"),
            ("VALUE LOST (USD)", f"US$ {es_k(usd, 1)}", f"Avg ticket US$ {es_k(usd / n if n else 0, 1)}"),
            ("MAIN REASON", REASON_EN.get(reason, reason), f"{pct(reason_pct, 0)} of lost value"),
            ("MAIN COMPETITOR", comp, f"{pct(comp_val / usd if usd else 0, 0)} of lost value"),
            ("LARGEST LOSS", top["cli"] if top else "—",
             f"US$ {es_k(top['usd'], 1)} · {REASON_EN.get(top['cat'], top['cat'])}" if top else ""),
        ]
        for i, (lab, val, sub) in enumerate(cards):
            add_card(s, f"lb{i}", x0 + i * (w + gap), 1.62, w, 0.86, lab, val, sub,
                     sub_color=TEAL if i in (0, 1) else MUTED)
        widths = [2.35, 1.2, 0.62, 0.85, 1.0, 1.15, 1.98]
        body = []
        for d in rows:
            proj = d["proj"] if d["proj"] and d["proj"] != "—" else ""
            body.append(([f"{d['cli']}" + (f" · {proj}" if proj else ""), (d["prod"] or "", "FFFFFF"),
                          es_dec(d["m2"], 0 if float(d["m2"]).is_integer() else 1), es_int(d["usd"]),
                          (d["comp"] or "", MUTED), (REASON_EN.get(d["cat"], d["cat"] or ""), AMBER),
                          (d["acc_s"] or d["acc"] or "", "FFFFFF")], "row"))
        body.append((["Total", "", es_dec(m2, 0), es_int(usd), "", "", ""], "total"))
        gf = add_table(s, "AUTO_lb_table", x0, 2.66, widths, 0.27,
                       ["Customer · Project", "Product", "m²", "USD", "Competitor", "Reason", "Corrective action"],
                       body, size=8)
        for i in range(len(body) + 1):           # text columns left aligned
            for j in (1, 4, 5, 6):
                gf.table.cell(i, j).text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
        # value by reason chart
        reasons = [(REASON_EN.get(e.v(L, f"M{r}"), e.v(L, f"M{r}")), e.v(L, f"O{r}") or 0) for r in range(15, 20)]
        reasons = sorted([r for r in reasons if r[1] > 0], key=lambda r: -r[1])
        cd = CategoryChartData()
        cd.categories = [r[0] for r in reasons]
        cd.add_series("USD", [r[1] / 1000 for r in reasons])
        cx = x0 + sum(widths) + 0.2
        gfc = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(cx), Inches(2.6),
                                 Inches(0.94 + 11.45 - cx), Inches(2.75), cd)
        gfc.name = "AUTO_lb_chart"
        ch = gfc.chart
        style_chart_text(ch, CREAM, 8)
        set_chart_title(ch, "Lost value by reason · US$ k", CREAM, size=8.5)
        ch.has_legend = False
        ch.category_axis.reverse_order = True
        ch.category_axis.format.line.fill.background()
        ch.category_axis.tick_labels.font.size = Pt(8)
        ch.category_axis.tick_labels.font.color.rgb = RGBColor.from_string(CREAM)
        ch.value_axis.visible = False
        ch.value_axis.has_major_gridlines = False
        ch.value_axis.maximum_scale = (reasons[0][1] / 1000) * 1.6 if reasons else 1
        ch.plots[0].gap_width = 50
        ser = ch.plots[0].series[0]
        for i, r in enumerate(reasons):
            pt = ser.points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = RGBColor.from_string(NEG if i == 0 else SKY)
            dl = pt.data_label
            dl.has_text_frame = True
            dl.text_frame.text = f"{es_dec(r[1] / 1000, 1)} · {pct(r[1] / usd, 0)}"
            nowrap(dl)
            run = dl.text_frame.paragraphs[0].runs[0]
            run.font.size = Pt(7.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string(CREAM)
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
        second = reasons[1] if len(reasons) > 1 else None
        txt = (f"{REASON_EN.get(reason, reason)} explains {pct(reason_pct, 0)} of the US$ {es_k(usd, 1)} lost "
               f"({comp} the main competitor)")
        if second:
            txt += f"; {second[0].lower()} adds US$ {es_dec(second[1] / 1000, 1)}k ({pct(second[1] / usd, 0)})"
        txt += f". {es_int(m2)} m² lost across {n} opportunities."
        add_insight(s, "lb_ins", x0, 5.62, 11.45, 0.62, [[(txt, {"bold": True})]], label="KEY INSIGHT")
        set_notes(s, [txt], "20_Lost_Business (tabla con filtros y consulta interactiva por competidor/motivo)")
        lbl = [sh for sh in s.shapes if sh.name == "AUTO_lb_ins_txt"][0]
        lbl.left, lbl.top, lbl.width, lbl.height = Inches(x0 + 1.3), Inches(5.62 + 0.12), Inches(10.0), Inches(0.4)
        lb = [sh for sh in s.shapes if sh.name == "AUTO_lb_ins_lbl"][0]
        lb.top, lb.height = Inches(5.62 + 0.2), Inches(0.2)

    # -- 11 · RSM by groups ------------------------------------------------------------
    def rsm_groups(self):
        s = find_slide(self.prs, r"RSM PERFORMANCE BY (CHANNEL|GROUPS)")
        if s is None:
            return None
        e, R = self.e, "RSM_Resumen_Canales"
        set_text(s.shapes.title, "RSM PERFORMANCE BY GROUPS (YTD)")
        for sh in shapes_named(s, "Subtitle"):
            set_text(sh, f"RSM groups · order intake, sales, Board budget compliance and backlog · thousands of m² "
                         f"(k m²), Jan–{self.mon} {YEAR}")
        widths = old_widths(s, "Table 25") or [1.5, .7, .8, .85, .85, .75, .75, .85, .85]
        clear_auto(s, ("Table 25",))
        names = ["Industrial", "Distribution", "Other Geographies", "Cold Storage", "Site Cali",
                 "Alternative Sales", "Total"]
        body = []
        for i in range(7):
            r = 5 + i
            body.append(([names[i], es_int(e.v(R, f"C{r}") / 1000), es_int(e.v(R, f"D{r}") / 1000),
                          es_int(e.v(R, f"T{r}") / 1000), self._cell(e.v(R, f"U{r}"), "gapk"),
                          self._cell(e.v(R, f"V{r}"), "sem"), es_int(e.v(R, f"I{r}") / 1000),
                          pct(e.v(R, f"J{r}"), 0), pct(e.v(R, f"M{r}") if i < 6 else 1, 0)],
                         "total" if i == 6 else "row"))
        add_table(s, "AUTO_tbl_groups", 0.85, 2.8, widths, 0.307,
                  ["Group", "OI YTD", "Sales YTD", "Board budget YTD", "Gap vs budget", "% Budget", "Backlog",
                   "Backlog margin", "Mix backlog"], body)
        oi, sales, bud = e.v(R, "C11"), e.v(R, "D11"), e.v(R, "T11")
        comp, gap = e.v(R, "V11"), e.v(R, "U11")
        n_ok = sum(1 for r in range(5, 11) if (e.v(R, f"V{r}") or 0) >= 1)
        kp = {"TextBox 7": f"{es_int(oi / 1000)}k m²", "TextBox 8": f"Book-to-bill {es_dec(oi / sales, 1)}x",
              "TextBox 11": f"{es_int(sales / 1000)}k m²", "TextBox 12": f"Gap {es_int(gap / 1000)}k m²",
              "TextBox 14": "BOARD BUDGET YTD", "TextBox 15": f"{es_int(bud / 1000)}k m²",
              "TextBox 16": "Board budget · 6 groups", "TextBox 19": pct(comp, 0),
              "TextBox 20": f"{n_ok} of 6 groups ≥100%", "TextBox 23": f"{es_int(e.v(R, 'I11') / 1000)}k m²",
              "TextBox 24": f"Margin {pct(e.v(R, 'J11'), 0)}"}
        for name, txt in kp.items():
            for sh in shapes_named(s, name):
                set_text(sh, txt)
        for sh in shapes_named(s, "TextBox 12"):
            set_text(sh, f"Gap {es_int(gap / 1000)}k m²", color=TEAL if gap >= 0 else NEG)
        for sh in shapes_named(s, "TextBox 19"):
            set_text(sh, pct(comp, 0), color=sem_color(comp))
        for sh in shapes_named(s, "TextBox 27"):
            set_text(sh, "Board budget YTD = Board total (Presupuesto Junta) allocated with each group's share of "
                         "the prior budget · backlog and margin per ERP channel report.")
        gaps = sorted([(names[r - 5], e.v(R, f"U{r}"), e.v(R, f"V{r}")) for r in range(5, 11)], key=lambda g: g[1])
        neg = [g for g in gaps if g[1] < 0]
        top2 = neg[:2]
        above = [g[0] for g in gaps if g[2] >= 1]
        txt = (f"Groups at {pct(comp, 0)} of the Board YTD budget (gap {es_int(gap / 1000)}k m²). "
               + ("Largest gaps: " + " and ".join(f"{g[0]} ({es_int(g[1] / 1000)}k m², {pct(g[2], 0)})"
                                                    for g in top2) if top2 else "No group below budget")
               + (f"; above budget: {', '.join(above)}." if above else ".")
               + f" Backlog {es_int(e.v(R, 'I11') / 1000)}k m² at {pct(e.v(R, 'J11'), 0)} margin.")
        for sh in shapes_named(s, "TextBox 38"):
            set_text(sh, txt)
        set_notes(s, [txt, "Board budget = Presupuesto Junta YTD × participación de cada grupo en el presupuesto "
                           "anterior (selector en 21_Ppto_Junta B5)."],
                  "RSM_Resumen_Canales (columnas T:W) · 21_Ppto_Junta")
        return s

    def rsm_summary(self, groups_slide):
        if groups_slide is None:
            return
        s = find_slide(self.prs, r"CHANNELS & SEGMENTS")
        if s is None:
            s = self._duplicate_after(groups_slide)
        e, m = self.e, "22_Canales_Segmentos"
        set_text(s.shapes.title, "RSM PERFORMANCE – CHANNELS & SEGMENTS (YTD)")
        for sh in shapes_named(s, "Subtitle"):
            set_text(sh, f"Summary of the RSM groups in the new structure · 3 channels and 2 segments · k m², "
                         f"Jan–{self.mon} {YEAR} · Board budget")
        drop = {sh.name for sh in s.shapes if sh.name.startswith(("KPI ", "TextBox ")) and sh.top < Inches(2.7)}
        clear_auto(s, drop | {"Table 25", "AUTO_tbl_groups"})
        spec = [("B", "k"), ("C", "k"), ("D", "k"), ("E", "gapk"), ("F", "sem"), ("G", "k"), ("H", "pct0")]
        hdr = ["OI YTD", "Sales YTD", "Board budget", "Gap", "% Budget", "Backlog", "Bkl margin"]
        widths = [1.45, 0.75, 0.8, 1.0, 0.7, 0.75, 0.75, 0.8]
        rows_c = self._label_rows(self._rows(m, 67, 4, spec), ["Industrial", "Distribution", "Alternative Sales",
                                                              "Total"])
        rows_s = self._label_rows(self._rows(m, 73, 3, spec), ["Construction", "Refrigeration", "Total"])
        add_textbox(s, "AUTO_lbl_ch", 0.85, 1.5, 4, 0.2, "BY CHANNEL", size=9, color=HEAD, bold=True)
        add_table(s, "AUTO_tbl_channels", 0.85, 1.72, widths, 0.3, ["Channel"] + hdr,
                  [(r, "total" if i == 3 else "row") for i, r in enumerate(rows_c)])
        add_textbox(s, "AUTO_lbl_seg", 0.85, 3.4, 4, 0.2, "BY SEGMENT", size=9, color=HEAD, bold=True)
        add_table(s, "AUTO_tbl_segments", 0.85, 3.62, widths, 0.3, ["Segment"] + hdr,
                  [(r, "total" if i == 2 else "row") for i, r in enumerate(rows_s)])
        cats = ["Industrial", "Distribution", "Alt. Sales", "Construction", "Refrigeration"]
        rws = [67, 68, 69, 73, 74]
        cd = CategoryChartData()
        cd.categories = cats
        cd.add_series("Sales YTD (k m²)", [e.v(m, f"C{r}") / 1000 for r in rws])
        cd.add_series("Board budget YTD (k m²)", [e.v(m, f"D{r}") / 1000 for r in rws])
        gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7.95), Inches(1.5), Inches(4.5),
                                Inches(3.7), cd)
        gf.name = "AUTO_chart_sales_budget"
        ch = gf.chart
        style_chart_text(ch, "0B2246", 8)
        set_chart_title(ch, "Sales vs Board budget YTD (k m²)", "0B2246")
        ch.has_legend = True
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(8)
        ch.legend.font.color.rgb = RGBColor.from_string("0B2246")
        ch.value_axis.visible = False
        ch.value_axis.has_major_gridlines = False
        ch.category_axis.tick_labels.font.size = Pt(7.5)
        plot = ch.plots[0]
        plot.gap_width = 60
        plot.overlap = -10
        for ser, col in zip(plot.series, [TEAL, HEAD]):
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = RGBColor.from_string(col)
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.number_format = "#,##0"
        dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        dl.font.size = Pt(7)
        dl.font.color.rgb = RGBColor.from_string("0B2246")
        for sh in shapes_named(s, "TextBox 27"):
            set_text(sh, f"RSM totals regrouped: Alternative Sales = current channel; Industrial/Distribution "
                         f"prorated {pct(e.v(m, 'J17'), 0)}/{pct(e.v(m, 'J18'), 0)} (sales); segments "
                         f"{pct(e.v(m, 'J25'), 0)}/{pct(e.v(m, 'J26'), 0)} (assigned).")
        c = [e.v(m, f"F{r}") for r in (67, 68, 69, 70)]
        txt = (f"Channels at {pct(c[3], 0)} of the Board budget: Industrial {pct(c[0], 0)}, Distribution "
               f"{pct(c[1], 0)}, Alternative Sales {pct(c[2], 0)}. Backlog {es_int(e.v(m, 'G70') / 1000)}k m² – "
               f"Industrial {pct(e.v(m, 'I67'), 0)} of it; Construction {es_int(e.v(m, 'G73') / 1000)}k / "
               f"Refrigeration {es_int(e.v(m, 'G74') / 1000)}k m².")
        for sh in shapes_named(s, "TextBox 38"):
            set_text(sh, txt)
        set_notes(s, [txt], "22_Canales_Segmentos (tabla resumen filas 66–75)")

    def _duplicate_after(self, src):
        new = self.prs.slides.add_slide(src.slide_layout)
        for sh in list(new.shapes):
            remove(sh)
        for el in src.shapes._spTree:
            if el.tag in (qn("p:sp"), qn("p:graphicFrame"), qn("p:grpSp"), qn("p:cxnSp"), qn("p:pic")):
                new.shapes._spTree.append(copy.deepcopy(el))
        bg = src._element.find(qn("p:cSld")).find(qn("p:bg"))
        if bg is not None:
            new._element.find(qn("p:cSld")).insert(0, copy.deepcopy(bg))
        ids = self.prs.slides._sldIdLst
        new_id = ids[-1]
        ids.remove(new_id)
        idx = [i for i, sl in enumerate(self.prs.slides) if sl.slide_id == src.slide_id][0]
        ids.insert(idx + 1, new_id)
        return new

    # -- 12 · RSM margin evolution (hidden slide) -------------------------------------
    def rsm_margin(self):
        s = find_slide(self.prs, r"RSM PERFORMANCE & MARGIN")
        if s is None:
            return
        e, R = self.e, "RSM_Resumen_Canales"
        ch = [sh for sh in s.shapes if sh.has_chart][0].chart
        names = ["Industrial", "Distribution", "Other Geographies", "Cold Storage", "Site Cali", "Alternative Sales"]
        cd = CategoryChartData()
        cd.categories = names
        cd.add_series("Sales YTD (k m²)", [e.v(R, f"B{r}") for r in range(30, 36)])
        cd.add_series("Board budget YTD (k m²)", [e.v(R, f"C{r}") for r in range(30, 36)])
        cd.add_series(f"Margin {self.py}", [e.v(R, f"D{r}") if isinstance(e.v(R, f"D{r}"), float) else None
                                            for r in range(30, 36)])
        cd.add_series(f"Margin {self.cm}", [e.v(R, f"E{r}") for r in range(30, 36)])
        ch.replace_data(cd)
        set_text(s.shapes.title, "RSM PERFORMANCE & MARGIN EVOLUTION BY GROUP")
        for sh in shapes_named(s, "Subtitle"):
            set_text(sh, f"Sales YTD vs Board budget YTD (k m², bars) · backlog gross margin {self.py} → {self.cm} "
                         f"(lines) · Jan–{self.mon} {YEAR}")
        tbl = [sh for sh in s.shapes if sh.has_table]
        if tbl:
            t = tbl[0].table
            for i in range(7):
                r = 30 + i
                vals = [pct(e.v(R, f"G{r}"), 0), pct(e.v(R, f"D{r}")), pct(e.v(R, f"E{r}")), pp(e.v(R, f"F{r}"))]
                for j, v in enumerate(vals, start=1):
                    cell = t.cell(i + 1, j)
                    runs = cell.text_frame.paragraphs[0].runs
                    if runs:
                        runs[0].text = v
                        for extra in runs[1:]:
                            extra.text = ""
            for j, h in enumerate(["Group", "% Budget", f"GM {self.py}", f"GM {self.cm}", "Δ pp"]):
                set_text(t.cell(0, j), h)
        up = sum(1 for r in range(30, 36) if isinstance(e.v(R, f"F{r}"), (int, float)) and e.v(R, f"F{r}") > 0)
        base = sum(1 for r in range(30, 36) if isinstance(e.v(R, f"F{r}"), (int, float)))
        for sh in s.shapes:
            if sh.has_text_frame and sh.name == "TextBox 10" and sh.top > Inches(5.5):
                set_text(sh, f"Margin improves in {up} of {base} groups with PY base (total {pct(e.v(R, 'D36'))} → "
                             f"{pct(e.v(R, 'E36'))}, {pp(e.v(R, 'F36'))}), while sales are at "
                             f"{pct(e.v(R, 'G36'), 0)} of the Board budget.")

    def run(self):
        self.fix_labels()
        self.overview()
        self.oi_month()
        self.backlog_month()
        self.channel_slide(r"ORDER INTAKE BY CHANNEL", "oi")
        self.channel_slide(r"BACKLOG BY CHANNEL", "backlog")
        self.lost_business()
        g = self.rsm_groups()
        self.rsm_summary(g)
        self.rsm_margin()
        self.renumber()


def main(xlsx, deck, out):
    eng = Engine(xlsx)
    prs = Presentation(deck)
    DeckUpdater(eng, prs).run()
    prs.save(out)
    fix_content_types(out)
    print("OK ->", out)


def fix_content_types(path):
    """python-pptx declares some media as per-part overrides; add the Default entries as well."""
    import zipfile
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        exts = {n.rsplit(".", 1)[-1].lower() for n in zin.namelist() if n.startswith("ppt/media/")}
        types = {"jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png", "svg": "image/svg+xml",
                 "gif": "image/gif", "mp4": "video/mp4", "emf": "image/x-emf", "wmf": "image/x-wmf"}
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                xml = data.decode("utf-8")
                for ext in sorted(exts):
                    if ext in types and f'Extension="{ext}"' not in xml:
                        xml = xml.replace("<Default ", f'<Default Extension="{ext}" ContentType="{types[ext]}"/><Default ', 1)
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    os.replace(tmp, path)


if __name__ == "__main__":
    if len(sys.argv) != 4:
        print(__doc__)
        sys.exit(1)
    main(*sys.argv[1:])
