"""
Genera la plantilla Excel de entrada semanal para el Comite de Operaciones.

Este script se ejecuta UNA sola vez (o cuando se quiera rediseñar la
plantilla). El archivo resultante (plantilla/Plantilla_Semanal_Comite_Operaciones.xlsx)
es el que cada responsable de area llena semana a semana; luego
scripts/generar_ppt_comite.py lo lee y arma automaticamente las diapositivas.

Uso:
    python scripts/crear_plantilla_excel.py
"""

import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT_PATH = "plantilla/Plantilla_Semanal_Comite_Operaciones.xlsx"

FONT_NAME = "Arial"

COLOR_HEADER = "1E2761"   # navy
COLOR_HEADER_TEXT = "FFFFFF"
COLOR_INPUT = "FFF2CC"    # amarillo suave -> celdas para llenar
COLOR_SECTION = "CADCFC"  # azul hielo
COLOR_EXAMPLE = "F2F2F2"

header_font = Font(name=FONT_NAME, size=11, bold=True, color=COLOR_HEADER_TEXT)
header_fill = PatternFill("solid", fgColor=COLOR_HEADER)
section_font = Font(name=FONT_NAME, size=13, bold=True, color=COLOR_HEADER)
section_fill = PatternFill("solid", fgColor=COLOR_SECTION)
input_fill = PatternFill("solid", fgColor=COLOR_INPUT)
example_font = Font(name=FONT_NAME, size=10, italic=True, color="808080")
normal_font = Font(name=FONT_NAME, size=10)
title_font = Font(name=FONT_NAME, size=16, bold=True, color=COLOR_HEADER)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()


def style_header_row(ws, row, n_cols, start_col=1):
    for c in range(start_col, start_col + n_cols):
        cell = ws.cell(row=row, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border


def write_table_headers(ws, row, headers, start_col=1):
    for i, h in enumerate(headers):
        ws.cell(row=row, column=start_col + i, value=h)
    style_header_row(ws, row, len(headers), start_col)


def add_example_row(ws, row, values, start_col=1, input_cols=None):
    """Escribe una fila de ejemplo; marca en amarillo las columnas editables."""
    input_cols = input_cols or []
    for i, v in enumerate(values):
        cell = ws.cell(row=row, column=start_col + i, value=v)
        cell.font = example_font
        cell.border = border
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if i in input_cols:
            cell.fill = input_fill


def autofit(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------------------------------------------------------------------------
# 1. INSTRUCCIONES
# ---------------------------------------------------------------------------
ws = wb.active
ws.title = "Instrucciones"
ws.sheet_view.showGridLines = False
ws["B2"] = "Plantilla Semanal - Comite de Operaciones"
ws["B2"].font = title_font
ws["B4"] = (
    "Como usar esta plantilla"
)
ws["B4"].font = section_font

instrucciones = [
    "1. Cada semana, antes del comite, actualice las hojas: Portada, Resumen Ejecutivo, "
    "KPIs, Proyectos, Incidentes, Riesgos, Plan de Accion y Proximos Pasos.",
    "2. Llene UNICAMENTE las celdas resaltadas en amarillo. Las demas columnas tienen "
    "formulas y se calculan solas (variaciones, % de cumplimiento, nivel de riesgo, etc.).",
    "3. No cambie los encabezados ni el orden de las columnas: el generador de PowerPoint "
    "los identifica por nombre y posicion.",
    "4. Puede agregar filas nuevas dentro de cada tabla (debajo del ejemplo); deje una fila "
    "en blanco despues de la ultima fila con datos para que el script sepa donde parar.",
    "5. Guarde el archivo (mismo nombre) y ejecute:",
    "        python scripts/generar_ppt_comite.py plantilla/Plantilla_Semanal_Comite_Operaciones.xlsx",
    "6. El script genera automaticamente un archivo .pptx nuevo, con fecha, listo para "
    "presentar (ver README.md para el detalle tecnico).",
    "7. Las filas de ejemplo (en gris/italica) son solo guia de formato; reemplacelas o "
    "bórrelas y use sus propios datos.",
]
r = 6
for line in instrucciones:
    ws.cell(row=r, column=2, value=line).font = normal_font
    ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 30
    r += 2

ws.merge_cells("B2:H2")
for rr in range(6, r, 2):
    ws.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=8)
autofit(ws, [3, 14, 14, 14, 14, 14, 14, 14])

# ---------------------------------------------------------------------------
# 2. PORTADA
# ---------------------------------------------------------------------------
ws = wb.create_sheet("Portada")
ws.sheet_view.showGridLines = False
ws["B2"] = "Datos generales de la sesion"
ws["B2"].font = title_font
ws.merge_cells("B2:E2")

campos = [
    ("Comite / Area", "Comite de Operaciones"),
    ("Semana (inicio)", datetime.date(2026, 9, 1)),
    ("Semana (fin)", datetime.date(2026, 9, 7)),
    ("Fecha de presentacion", datetime.date(2026, 9, 8)),
    ("Responsable de la presentacion", "Nombre Apellido"),
    ("Estado general de la semana (Verde/Amarillo/Rojo)", "Amarillo"),
]
r = 4
for etiqueta, ejemplo in campos:
    ws.cell(row=r, column=2, value=etiqueta).font = Font(name=FONT_NAME, bold=True)
    c = ws.cell(row=r, column=4, value=ejemplo)
    c.fill = input_fill
    c.font = normal_font
    c.border = border
    if isinstance(ejemplo, datetime.date):
        c.number_format = "dd/mm/yyyy"
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    r += 1

dv_estado = DataValidation(type="list", formula1='"Verde,Amarillo,Rojo"', allow_blank=True)
ws.add_data_validation(dv_estado)
dv_estado.add(ws.cell(row=9, column=4))

autofit(ws, [3, 30, 3, 30])

# ---------------------------------------------------------------------------
# 3. RESUMEN EJECUTIVO
# ---------------------------------------------------------------------------
ws = wb.create_sheet("Resumen Ejecutivo")
ws.sheet_view.showGridLines = False
ws["B2"] = "Resumen ejecutivo de la semana"
ws["B2"].font = title_font
ws.merge_cells("B2:F2")

ws["B4"] = "Logros / avances destacados (uno por fila)"
ws["B4"].font = section_font
ws.merge_cells("B4:F4")
logros_ejemplo = [
    "Disponibilidad de plataforma se mantuvo en 99.6%, por encima de la meta.",
    "Se cerro el incidente critico INC-014 dentro del SLA.",
    "Proyecto Migracion ERP alcanzo 65% de avance, en linea con el plan.",
]
r = 5
for i, txt in enumerate(logros_ejemplo):
    cell = ws.cell(row=r, column=2, value=txt)
    cell.fill = input_fill
    cell.font = example_font
    cell.border = border
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    ws.row_dimensions[r].height = 28
    r += 1

r += 1
ws.cell(row=r, column=2, value="Puntos de atencion / riesgos a escalar (uno por fila)").font = section_font
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
r += 1
atencion_ejemplo = [
    "Retraso de 2 semanas en entrega de proveedor para el proyecto de Datacenter.",
    "Incremento de 15% en tickets de soporte nivel 2 respecto a la semana anterior.",
]
for txt in atencion_ejemplo:
    cell = ws.cell(row=r, column=2, value=txt)
    cell.fill = input_fill
    cell.font = example_font
    cell.border = border
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    ws.row_dimensions[r].height = 28
    r += 1

autofit(ws, [3, 22, 22, 22, 22, 22])

# ---------------------------------------------------------------------------
# 4. KPIs
# ---------------------------------------------------------------------------
ws = wb.create_sheet("KPIs")
ws.sheet_view.showGridLines = False
ws["B2"] = "Indicadores operacionales (KPIs)"
ws["B2"].font = title_font
ws.merge_cells("B2:J2")

headers = [
    "Indicador", "Unidad", "Meta", "Sentido", "Semana Anterior", "Semana Actual",
    "Variacion", "% Cumplimiento Meta", "Estado",
]
header_row = 4
write_table_headers(ws, header_row, headers, start_col=2)

# Sentido del indicador: de esto depende si "actual > meta" es bueno o malo.
#   "Mayor es mejor"  -> ej. disponibilidad, SLA, CSAT
#   "Menor es mejor"  -> ej. MTTR, tickets abiertos, incidentes
kpis_ejemplo = [
    ("Disponibilidad de servicio", "%", 99.5, "Mayor es mejor", 99.4, 99.6),
    ("SLA cumplido (tickets)", "%", 95, "Mayor es mejor", 93, 96),
    ("Tiempo promedio de resolucion (MTTR)", "horas", 4, "Menor es mejor", 5.2, 4.5),
    ("Tickets abiertos", "#", 50, "Menor es mejor", 62, 48),
    ("Satisfaccion del cliente (CSAT)", "%", 90, "Mayor es mejor", 88, 91),
    ("Cumplimiento de presupuesto operativo", "%", 100, "Mayor es mejor", 97, 98),
]

first_data_row = header_row + 1
for i, (nombre, unidad, meta, sentido, ant, act) in enumerate(kpis_ejemplo):
    row = first_data_row + i
    ws.cell(row=row, column=2, value=nombre)
    ws.cell(row=row, column=3, value=unidad)
    ws.cell(row=row, column=4, value=meta)
    ws.cell(row=row, column=5, value=sentido)
    ws.cell(row=row, column=6, value=ant)
    ws.cell(row=row, column=7, value=act)
    # Variacion = actual - anterior
    ws.cell(row=row, column=8, value=f"=G{row}-F{row}")
    # % cumplimiento: si "Menor es mejor" se invierte la razon (meta/actual),
    # para que un valor mas bajo que la meta siga marcando >=100% de cumplimiento.
    cumplimiento_formula = (
        f'=IF(E{row}="Menor es mejor",IFERROR(D{row}/G{row},0),IFERROR(G{row}/D{row},0))'
    )
    ws.cell(row=row, column=9, value=cumplimiento_formula)
    ws.cell(row=row, column=9).number_format = "0.0%"
    # Estado segun % cumplimiento (ya normalizado por el sentido del indicador)
    estado_formula = (
        f'=IF(I{row}>=1,"Verde",IF(I{row}>=0.9,"Amarillo","Rojo"))'
    )
    ws.cell(row=row, column=10, value=estado_formula)

    for col in (2, 3, 4, 5, 6, 7, 10):
        ws.cell(row=row, column=col).font = example_font
    for col in (4, 5, 6, 7):
        ws.cell(row=row, column=col).fill = input_fill
    ws.cell(row=row, column=2).fill = input_fill
    ws.cell(row=row, column=3).fill = input_fill
    for col in range(2, 11):
        ws.cell(row=row, column=col).border = border
        ws.cell(row=row, column=col).alignment = Alignment(vertical="center", wrap_text=True)

dv_sentido = DataValidation(type="list", formula1='"Mayor es mejor,Menor es mejor"', allow_blank=True)
ws.add_data_validation(dv_sentido)
dv_sentido.add(f"E{first_data_row}:E{first_data_row + 20}")

autofit(ws, [3, 34, 8, 8, 16, 14, 12, 12, 16, 12])

# ---------------------------------------------------------------------------
# 5. PROYECTOS
# ---------------------------------------------------------------------------
ws = wb.create_sheet("Proyectos")
ws.sheet_view.showGridLines = False
ws["B2"] = "Avance de proyectos e iniciativas"
ws["B2"].font = title_font
ws.merge_cells("B2:I2")

headers = ["Proyecto", "Responsable", "% Avance", "Estado (RAG)", "Proximo hito", "Fecha hito", "Comentario"]
header_row = 4
write_table_headers(ws, header_row, headers, start_col=2)

proyectos_ejemplo = [
    ("Migracion ERP", "Maria Gomez", 0.65, "Amarillo", "Pruebas de integracion", datetime.date(2026, 9, 20), "Requiere soporte adicional de TI."),
    ("Renovacion Datacenter", "Carlos Ruiz", 0.30, "Rojo", "Entrega de equipos", datetime.date(2026, 10, 5), "Proveedor con retraso de 2 semanas."),
    ("Automatizacion de reportes", "Ana Torres", 0.80, "Verde", "Puesta en produccion", datetime.date(2026, 9, 15), "Sin novedades."),
]
first_data_row = header_row + 1
for i, (proy, resp, avance, estado, hito, fecha, coment) in enumerate(proyectos_ejemplo):
    row = first_data_row + i
    vals = [proy, resp, avance, estado, hito, fecha, coment]
    for j, v in enumerate(vals):
        cell = ws.cell(row=row, column=2 + j, value=v)
        cell.font = example_font
        cell.fill = input_fill
        cell.border = border
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.cell(row=row, column=4).number_format = "0%"
    ws.cell(row=row, column=7).number_format = "dd/mm/yyyy"

dv_rag = DataValidation(type="list", formula1='"Verde,Amarillo,Rojo"', allow_blank=True)
ws.add_data_validation(dv_rag)
dv_rag.add(f"E{first_data_row}:E{first_data_row + 20}")

autofit(ws, [3, 26, 16, 10, 12, 22, 14, 34])

# ---------------------------------------------------------------------------
# 6. INCIDENTES
# ---------------------------------------------------------------------------
ws = wb.create_sheet("Incidentes")
ws.sheet_view.showGridLines = False
ws["B2"] = "Incidentes y problemas de la semana"
ws["B2"].font = title_font
ws.merge_cells("B2:J2")

headers = ["ID", "Fecha", "Descripcion", "Severidad", "Estado", "Responsable", "Causa raiz", "Accion correctiva"]
header_row = 4
write_table_headers(ws, header_row, headers, start_col=2)

incidentes_ejemplo = [
    ("INC-014", datetime.date(2026, 9, 3), "Caida intermitente del portal de clientes", "Alta", "Cerrado", "Equipo Infraestructura", "Saturacion de memoria en balanceador", "Se aumento capacidad y se agrego alerta temprana"),
    ("INC-015", datetime.date(2026, 9, 5), "Retraso en procesamiento de pagos", "Media", "En seguimiento", "Equipo Pagos", "Error en job nocturno", "Ajuste de job y monitoreo reforzado"),
]
first_data_row = header_row + 1
for i, vals in enumerate(incidentes_ejemplo):
    row = first_data_row + i
    for j, v in enumerate(vals):
        cell = ws.cell(row=row, column=2 + j, value=v)
        cell.font = example_font
        cell.fill = input_fill
        cell.border = border
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.cell(row=row, column=3).number_format = "dd/mm/yyyy"

dv_sev = DataValidation(type="list", formula1='"Alta,Media,Baja"', allow_blank=True)
ws.add_data_validation(dv_sev)
dv_sev.add(f"E{first_data_row}:E{first_data_row + 30}")

autofit(ws, [3, 10, 12, 30, 10, 16, 18, 26, 30])

# ---------------------------------------------------------------------------
# 7. RIESGOS
# ---------------------------------------------------------------------------
ws = wb.create_sheet("Riesgos")
ws.sheet_view.showGridLines = False
ws["B2"] = "Matriz de riesgos operacionales"
ws["B2"].font = title_font
ws.merge_cells("B2:I2")

# Tabla de escala auxiliar (para calcular el nivel de riesgo con MATCH)
ws["L4"] = "Escala"
ws["L4"].font = section_font
ws["L5"] = "Baja"
ws["M5"] = 1
ws["L6"] = "Media"
ws["M6"] = 2
ws["L7"] = "Alta"
ws["M7"] = 3

headers = ["Riesgo", "Probabilidad", "Impacto", "Nivel de riesgo", "Plan de mitigacion", "Responsable"]
header_row = 4
write_table_headers(ws, header_row, headers, start_col=2)

riesgos_ejemplo = [
    ("Falla de proveedor critico de infraestructura", "Media", "Alta", "Plan de mitigacion"),
    ("Rotacion de personal clave en soporte", "Baja", "Media", "Plan de mitigacion"),
]
first_data_row = header_row + 1
for i, (riesgo, prob, impacto, resp) in enumerate(riesgos_ejemplo):
    row = first_data_row + i
    ws.cell(row=row, column=2, value=riesgo)
    ws.cell(row=row, column=3, value=prob)
    ws.cell(row=row, column=4, value=impacto)
    nivel_formula = (
        f'=IF(MATCH(C{row},$L$5:$L$7,0)*MATCH(D{row},$L$5:$L$7,0)>=6,"Alto",'
        f'IF(MATCH(C{row},$L$5:$L$7,0)*MATCH(D{row},$L$5:$L$7,0)>=3,"Medio","Bajo"))'
    )
    ws.cell(row=row, column=5, value=nivel_formula)
    ws.cell(row=row, column=6, value="Definir acciones de mitigacion y responsable")
    ws.cell(row=row, column=7, value="Nombre")

    for col in (2, 3, 4, 5, 6, 7):
        ws.cell(row=row, column=col).font = example_font
        ws.cell(row=row, column=col).border = border
        ws.cell(row=row, column=col).alignment = Alignment(vertical="center", wrap_text=True)
    for col in (2, 3, 4, 6, 7):
        ws.cell(row=row, column=col).fill = input_fill

dv_prob = DataValidation(type="list", formula1='"Baja,Media,Alta"', allow_blank=True)
ws.add_data_validation(dv_prob)
dv_prob.add(f"C{first_data_row}:D{first_data_row + 20}")

autofit(ws, [3, 34, 14, 14, 14, 34, 16])

# ---------------------------------------------------------------------------
# 8. PLAN DE ACCION
# ---------------------------------------------------------------------------
ws = wb.create_sheet("Plan de Accion")
ws.sheet_view.showGridLines = False
ws["B2"] = "Compromisos y plan de accion"
ws["B2"].font = title_font
ws.merge_cells("B2:H2")

headers = ["Accion", "Responsable", "Fecha compromiso", "Prioridad", "Estado"]
header_row = 4
write_table_headers(ws, header_row, headers, start_col=2)

acciones_ejemplo = [
    ("Escalar con proveedor plan de recuperacion de cronograma", "Carlos Ruiz", datetime.date(2026, 9, 10), "Alta", "En curso"),
    ("Reforzar turno de soporte nivel 2 los fines de semana", "Ana Torres", datetime.date(2026, 9, 12), "Media", "Pendiente"),
]
first_data_row = header_row + 1
for i, vals in enumerate(acciones_ejemplo):
    row = first_data_row + i
    for j, v in enumerate(vals):
        cell = ws.cell(row=row, column=2 + j, value=v)
        cell.font = example_font
        cell.fill = input_fill
        cell.border = border
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.cell(row=row, column=4).number_format = "dd/mm/yyyy"

dv_prio = DataValidation(type="list", formula1='"Alta,Media,Baja"', allow_blank=True)
ws.add_data_validation(dv_prio)
dv_prio.add(f"E{first_data_row}:E{first_data_row + 30}")

dv_estado_accion = DataValidation(type="list", formula1='"Pendiente,En curso,Completado,Atrasado"', allow_blank=True)
ws.add_data_validation(dv_estado_accion)
dv_estado_accion.add(f"F{first_data_row}:F{first_data_row + 30}")

autofit(ws, [3, 40, 18, 16, 12, 14])

# ---------------------------------------------------------------------------
# 9. PROXIMOS PASOS
# ---------------------------------------------------------------------------
ws = wb.create_sheet("Proximos Pasos")
ws.sheet_view.showGridLines = False
ws["B2"] = "Agenda / proximos pasos para la siguiente semana"
ws["B2"].font = title_font
ws.merge_cells("B2:F2")

pasos_ejemplo = [
    "Revisar avance del plan de recuperacion con el proveedor de Datacenter.",
    "Presentar resultados de la automatizacion de reportes al comite.",
    "Actualizar matriz de riesgos con nuevos hallazgos de auditoria.",
]
r = 4
for txt in pasos_ejemplo:
    cell = ws.cell(row=r, column=2, value=txt)
    cell.fill = input_fill
    cell.font = example_font
    cell.border = border
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    ws.row_dimensions[r].height = 26
    r += 1

autofit(ws, [3, 24, 24, 24, 24, 24])

wb.save(OUT_PATH)
print(f"Plantilla creada en {OUT_PATH}")
